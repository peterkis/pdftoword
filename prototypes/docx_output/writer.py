"""Real flowing DOCX writer and package-level integrity checks."""

from __future__ import annotations

import hashlib
import io
import posixpath
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree
from PIL import Image

from .common import DemoError, Json, safe_path, save, validate
from .formula import unrendered_math


def fonts() -> Json:
    """Detect available local font families without distributing font files."""
    roots = [Path("/System/Library/Fonts"), Path("/Library/Fonts"), Path.home() / "Library/Fonts"]
    names = [p.name.lower() for r in roots if r.exists() for p in r.rglob("*") if p.is_file()]
    east = next(
        (
            family
            for token, family in [
                ("pingfang", "PingFang SC"),
                ("songti", "Songti SC"),
                ("notosanscjk", "Noto Sans CJK SC"),
            ]
            if any(token in n for n in names)
        ),
        None,
    )
    latin = next(
        (
            family
            for token, family in [("arial", "Arial"), ("times", "Times New Roman")]
            if any(token in n for n in names)
        ),
        None,
    )
    return {
        "east_asia": east,
        "latin": latin,
        "missing": not east or not latin,
        "raster_style": "inferred_demo_11pt",
    }


def text_column_widths(
    rows: list[list[str]], blocks: Json, width: float, parts: Json, font_size: float = 11
) -> list[float]:
    """Allocate display widths for plain-text cells; never manufacture source bboxes."""
    count = max(len(row) for row in rows)
    equal = width / count
    if any(bid in parts for row in rows for bid in row):
        return [equal] * count
    minimum = [0.0] * count
    for row in rows:
        for column, bid in enumerate(row):
            text = blocks[bid]["content"].get("plain_text", "")
            units = sum(1.0 if unicodedata.east_asian_width(c) in {"W", "F"} else 0.5 for c in text)
            minimum[column] = max(minimum[column], units * font_size + 12)
    if sum(minimum) >= width:
        return [equal] * count  # Keep natural wrapping; never shrink the text.
    widths = [max(equal, m) for m in minimum]
    excess = sum(widths) - width
    spare = sum(w - m for w, m in zip(widths, minimum, strict=True))
    return [
        w - excess * (w - m) / spare if spare else w for w, m in zip(widths, minimum, strict=True)
    ]


def build(job: Path, ir: Json, revision: str, *, output_plan: Json | None = None) -> Json:
    """Consume validated IR only; never query a model or overwrite auto output."""
    validate(ir)
    if revision not in {"auto", "reviewed"}:
        raise DemoError("INVALID_REVISION")
    path = safe_path(job, f"{revision}.docx")
    if path.exists() and revision == "auto":
        raise DemoError("AUTO_IMMUTABLE")
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(595.28), Pt(841.89)
    section.left_margin = section.right_margin = Pt(45)
    section.top_margin = section.bottom_margin = Pt(40)
    available = 505.28
    flow = output_plan is not None
    flow_nodes: Json = {}
    flow_sections: Json = {}
    paragraph_cache: Json = {}
    current_section_id = None
    if output_plan is not None:
        for planned in output_plan["sections"]:
            for source_page in planned["source_pages"]:
                flow_sections[source_page] = planned
            for node in planned["nodes"]:
                for leaf in node.get("children", [node]):
                    for source_id in leaf["source_ids"]:
                        flow_nodes[source_id] = leaf

    def configure_section(target: Any, planned: Json) -> None:
        width, height = planned["page_size_pt"]
        left, top, right, bottom = planned["margins_pt"]
        target.page_width, target.page_height = Pt(width), Pt(height)
        target.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
        target.left_margin, target.top_margin = Pt(left), Pt(top)
        target.right_margin, target.bottom_margin = Pt(right), Pt(bottom)

    spatial = ir["metadata"].get("layout_profile") == "pp_geometry_flow"
    content_left = ir["metadata"].get("content_left_pt", 0)
    font_info = fonts()
    if flow:
        assert output_plan is not None
        body_style = ir["output_styles"]["body"]
        font_info = {
            "latin": body_style["latin_font"],
            "east_asia": body_style["east_asia_font"],
            "missing": bool(output_plan["issues"]),
            "raster_style": body_style["basis"],
        }
    for name, size in [("Normal", 11), ("Title", 16), ("Heading 1", 13), ("Caption", 10)]:
        style: Any = doc.styles[name]
        style.font.name = font_info["latin"] or "Arial"
        if flow:
            role = (
                "heading"
                if name in {"Title", "Heading 1"}
                else "caption"
                if name == "Caption"
                else "body"
            )
            size = ir["output_styles"][role]["font_size_pt"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), font_info["east_asia"] or "sans-serif"
        )
        style.paragraph_format.space_after = Pt(5)
        if flow:
            rf = style.element.get_or_add_rPr().get_or_add_rFonts()
            rf.set(qn("w:ascii"), font_info["latin"])
            rf.set(qn("w:hAnsi"), font_info["latin"])
        if spatial and not flow:
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.line_spacing = 1.05
    assets = {a["id"]: a for a in ir["assets"]}
    parts = ir["metadata"].get("inline_parts", {})
    counts = {
        "editable_text_char_count": 0,
        "placed_figure_count": 0,
        "formula_image_count": 0,
        "fallback_region_count": 0,
        "omml_formula_count": 0,
    }

    source_blocks: list[Json] = []
    marker_ends: dict[str, tuple[Any, Any]] = {}

    def asset_source(aid: str) -> Path:
        if aid not in assets:
            raise DemoError("ASSET_MISSING")
        source = safe_path(job, assets[aid]["path"])
        if not source.is_file():
            raise DemoError("ASSET_MISSING")
        return source

    def source_marker(paragraph: Any, b: Json) -> Json:
        # Provenance only: no reference annotations or expected text enter the writer.
        marker = "p2w_" + hashlib.sha256(b["id"].encode()).hexdigest()[:32]
        number = str(len(source_blocks))
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), number)
        start.set(qn("w:name"), marker)
        end.set(qn("w:id"), number)
        paragraph._p.append(start)
        marker_ends[marker] = (paragraph._p, end)
        record = {
            "marker": marker,
            "block_id": b["id"],
            "images": [],
            "page": b["page_index"] + 1,
            "bbox": b["bbox"],
            "kind": b["type"],
            "fallback": b["content"]["kind"] == "image" and b["type"] != "figure",
        }
        source_blocks.append(record)
        return record

    def close_marker(record: Json) -> None:
        """Close the source range after its actual text, math or drawing content."""
        paragraph, end = marker_ends.pop(record["marker"])
        paragraph.append(end)

    def picture(
        paragraph: Any, aid: str, width: float, formula: bool = False, *, owner: Json
    ) -> None:
        source = asset_source(aid)
        a = assets[aid]
        try:
            image_bytes = source.read_bytes()
            with Image.open(io.BytesIO(image_bytes)) as im:
                w, h = im.size
        except OSError as exc:
            raise DemoError("ASSET_UNREADABLE") from exc
        natural = max(1.0, a["source_bbox"][2] - a["source_bbox"][0])
        actual = min(width, natural)
        # Preserve source aspect ratio; bound tall images to printable page height.
        height_limit = flow_nodes[owner["block_id"]]["height_limit_pt"] if flow else 650
        actual = min(actual, height_limit * w / h)
        if flow:
            actual = min(actual, flow_nodes[owner["block_id"]]["width_pt"])
        image_hash = hashlib.sha256(image_bytes).hexdigest()
        paragraph.add_run().add_picture(io.BytesIO(image_bytes), width=Pt(actual))
        owner["images"].append(
            {"sha256": image_hash, "bbox": a["source_bbox"], "fallback": owner["kind"] != "figure"}
        )
        if formula:
            counts["formula_image_count"] += 1

    def set_run_style(run: Any, spec: Json) -> None:
        rf = run._r.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia"):
            rf.set(qn("w:" + key), spec[key])
        run.font.size = Pt(spec["font_size_pt"])
        run.bold, run.italic, run.underline = (
            spec.get("bold", False),
            spec.get("italic", False),
            spec.get("underline", False),
        )
        run.font.superscript, run.font.subscript = (
            spec.get("superscript", False),
            spec.get("subscript", False),
        )
        if spec.get("color"):
            run.font.color.rgb = RGBColor.from_string(spec["color"].lstrip("#"))

    def write_block(parent: Any, b: Json, width: float | None = None, existing: Any = None) -> None:
        if width is None:
            width = available
        style = (
            "Heading 1"
            if b["type"] == "heading"
            else "Caption"
            if b["type"] in {"caption", "footer"}
            else "Normal"
        )
        node = flow_nodes.get(b["id"]) if flow else None
        reused = bool(node and node["id"] in paragraph_cache and existing is None)
        cached = paragraph_cache.get(node["id"]) if node is not None else None
        para = (
            existing
            if existing is not None
            else cached
            if reused
            else parent.add_paragraph(style=style)
        )
        assert para is not None
        if node and existing is None:
            paragraph_cache[node["id"]] = para
        if reused and node is not None and b["id"] in node["joiners"]:
            para.add_run(node["joiners"][b["id"]])
        para.style = style
        source_record = source_marker(para, b)
        if spatial and not flow and b["type"] in {"heading", "footer"}:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if (
            spatial
            and not flow
            and b["type"] == "paragraph"
            and b["geometry_source"] == "pp_structure"
            and b["bbox"][0] > content_left + 60
        ):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = b["type"] in {"heading", "caption", "question"}
        para.paragraph_format.widow_control = True
        if node:
            para.alignment = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }[node["alignment"]]
            fmt = para.paragraph_format
            fmt.left_indent, fmt.right_indent = Pt(node["indent_pt"]), Pt(node["right_indent_pt"])
            fmt.first_line_indent = Pt(node.get("first_line_indent_pt", 0.0))
            fmt.space_before, fmt.space_after = (
                Pt(node["space_before_pt"]),
                Pt(node["space_after_pt"]),
            )
            fmt.line_spacing = node["line_spacing"]
            fmt.keep_with_next, fmt.widow_control = node["keep_with_next"], node["widow_control"]
        content = b["content"]
        if content["kind"] == "image":
            picture(para, content["asset_id"], width, owner=source_record)
            if b["type"] == "figure":
                counts["placed_figure_count"] += 1
            else:
                counts["fallback_region_count"] += 1
            close_marker(source_record)
            return
        if b["id"] in parts:
            for item in parts[b["id"]]:
                if "omml" in item:
                    math = etree.fromstring(
                        item["omml"].encode(),
                        etree.XMLParser(resolve_entities=False, no_network=True),
                    )
                    if (
                        math.tag
                        != "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
                    ):
                        raise DemoError("INVALID_OMML")
                    para._p.append(math)
                    counts["omml_formula_count"] += 1
                elif "asset_id" in item:
                    picture(para, item["asset_id"], width, True, owner=source_record)
                else:
                    if unrendered_math(item["text"]):
                        raise DemoError("UNRENDERED_MATH_REQUIRES_REVIEW")
                    para.add_run(item["text"])
                    counts["editable_text_char_count"] += len(item["text"].strip())
            close_marker(source_record)
            return
        if content["kind"] != "text":
            raise DemoError("UNSUPPORTED_IR_CONTENT")
        text = content["plain_text"]
        if unrendered_math(text):
            raise DemoError("UNRENDERED_MATH_REQUIRES_REVIEW")
        if content["runs"]:
            for run_index, r in enumerate(content["runs"]):
                run = para.add_run(r["text"])
                run.bold, run.italic = r["bold"], r["italic"]
                if r["font_size_pt"]:
                    run.font.size = Pt(r["font_size_pt"])
                if r["font_family"]:
                    run.font.name = r["font_family"]
                if node:
                    set_run_style(run, node["runs"][b["id"]][run_index])
        else:
            run = para.add_run(text)
            if node:
                planned_style = ir["output_styles"][node["style_id"]]
                set_run_style(
                    run,
                    {
                        "ascii": planned_style["latin_font"],
                        "hAnsi": planned_style["latin_font"],
                        "eastAsia": planned_style["east_asia_font"],
                        "font_size_pt": planned_style["font_size_pt"],
                    },
                )
        counts["editable_text_char_count"] += len(text.strip())
        close_marker(source_record)

    for p in ir["pages"]:
        if flow:
            planned = flow_sections[p["page_index"]]
            if planned["id"] != current_section_id:
                if current_section_id is not None:
                    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
                configure_section(section, planned)
                current_section_id = planned["id"]
            available = (
                planned["page_size_pt"][0] - planned["margins_pt"][0] - planned["margins_pt"][2]
            )
        decision = ir["metadata"].get("layout_by_page", {}).get(str(p["page_index"]))
        content_left = (
            decision.get("content_left_pt", 0)
            if decision
            else ir["metadata"].get("content_left_pt", 0)
        )
        if decision:
            spatial = decision["status"] == "APPLIED"
        by_id = {b["id"]: b for b in p["blocks"]}
        group_map: dict[str, Json] = {}
        for g in ir["metadata"].get("figure_groups", []):
            if g["page_index"] == p["page_index"]:
                for pair in g["pairs"]:
                    group_map[pair["label"]] = g
                    group_map[pair["figure"]] = g
        text_map: dict[str, Json] = {}
        for g in ir["metadata"].get("text_groups", []):
            if g["page_index"] == p["page_index"]:
                for row in g["rows"]:
                    for item in row:
                        text_map[item] = g
        done: set[str] = set()
        for bid in p["reading_order"]:
            if bid in done:
                continue
            if bid in text_map:
                g = text_map[bid]
                indent = min(g.get("indent_pt", 0), 36)
                width = available - indent
                table = doc.add_table(rows=len(g["rows"]), cols=g["columns"])
                table.autofit = False
                widths = text_column_widths(
                    g["rows"],
                    by_id,
                    width,
                    parts,
                    ir["output_styles"]["body"]["font_size_pt"] if flow else 11,
                )
                for column, cell_width in zip(table.columns, widths, strict=True):
                    column.width = Pt(cell_width)
                borders = OxmlElement("w:tblBorders")
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    el = OxmlElement("w:" + edge)
                    el.set(qn("w:val"), "nil")
                    borders.append(el)
                table._tbl.tblPr.append(borders)
                ind = OxmlElement("w:tblInd")
                ind.set(qn("w:w"), str(round(indent * 20)))
                ind.set(qn("w:type"), "dxa")
                table._tbl.tblPr.append(ind)
                for ri, row in enumerate(g["rows"]):
                    table.rows[ri]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    for ci, item in enumerate(row):
                        cell = table.cell(ri, ci)
                        cell.width = Pt(widths[ci])
                        write_block(cell, by_id[item], widths[ci] - 10, cell.paragraphs[0])
                        done.add(item)
                continue
            if bid not in group_map:
                write_block(doc, by_id[bid])
                continue
            g = group_map[bid]
            pairs = g["pairs"]
            cols = min(g.get("columns", 2 if g["kind"] == "option_grid" else 3), len(pairs))
            table = doc.add_table(rows=(len(pairs) + cols - 1) // cols, cols=cols)
            table.autofit = False
            for column in table.columns:
                column.width = Pt(available / cols)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:val"), "nil")
                borders.append(e)
            table._tbl.tblPr.append(borders)
            for row in table.rows:
                prop = row._tr.get_or_add_trPr()
                prop.append(OxmlElement("w:cantSplit"))
            for i, pair in enumerate(pairs):
                cell = table.cell(i // cols, i % cols)
                cell.width = Pt(available / cols)
                if g.get("inline_labels"):
                    label = by_id[pair["label"]]
                    para = cell.paragraphs[0]
                    label_source = source_marker(para, label)
                    text = label["content"]["plain_text"]
                    para.add_run(text)
                    close_marker(label_source)
                    para.add_run(" ")
                    figure_source = source_marker(para, by_id[pair["figure"]])
                    counts["editable_text_char_count"] += len(text.strip())
                    picture(
                        para,
                        by_id[pair["figure"]]["content"]["asset_id"],
                        available / cols - 30,
                        owner=figure_source,
                    )
                    close_marker(figure_source)
                    counts["placed_figure_count"] += 1
                    para.paragraph_format.space_after = Pt(6)
                else:
                    first, second = (
                        ("label", "figure") if g["kind"] == "option_grid" else ("figure", "label")
                    )
                    write_block(
                        cell,
                        by_id[pair[first]],
                        available / cols - 16,
                        cell.paragraphs[0] if spatial else None,
                    )
                    write_block(cell, by_id[pair[second]], available / cols - 16)
                    cell.paragraphs[-2].paragraph_format.keep_with_next = True
                    cell.paragraphs[-1].paragraph_format.keep_with_next = False
                done.update(pair.values())
    if marker_ends:
        raise DemoError("SOURCE_RANGE_NOT_CLOSED")
    doc.save(str(path))
    path.chmod(0o600)
    save(
        job / f"source-map.{revision}.json",
        {
            "schema_version": "docx-source-map/1.0",
            "docx_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "pages": [
                {"page": p["page_index"] + 1, "width": p["width_pt"], "height": p["height_pt"]}
                for p in ir["pages"]
            ],
            "blocks": source_blocks,
            "relations": [
                {"from": r["from"], "to": r["to"], "type": r["type"]} for r in ir["relations"]
            ],
        },
    )
    return {**counts, "fonts": font_info, **inspect_package(path)}


def inspect_package(path: Path) -> Json:
    """Validate XML, internal relationship targets, editable runs and image extents."""
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        for name in names:
            if name.endswith((".xml", ".rels")):
                root = etree.fromstring(z.read(name), parser)
                if name.endswith(".rels"):
                    parent = posixpath.dirname(posixpath.dirname(name))
                    for rel in root:
                        target = rel.get("Target", "")
                        if rel.get("TargetMode") == "External":
                            raise DemoError("EXTERNAL_RELATIONSHIP")
                        resolved = posixpath.normpath(posixpath.join(parent, target)).lstrip("/")
                        if resolved not in names:
                            raise DemoError("OOXML_TARGET_MISSING")
            if name.startswith("word/media/"):
                with Image.open(io.BytesIO(z.read(name))) as im:
                    im.verify()
        root = etree.fromstring(z.read("word/document.xml"), parser)
        ns = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        }
        text = "".join(root.xpath("//w:t/text()", namespaces=ns))
        math_text = "".join(root.xpath("//m:oMath//m:t/text()", namespaces=ns))
        for extent in root.xpath("//wp:extent", namespaces=ns):
            if int(extent.get("cx", "0")) <= 0 or int(extent.get("cy", "0")) <= 0:
                raise DemoError("INVALID_IMAGE_EXTENT")
        return {
            "docx_package_valid": True,
            "has_editable_runs": bool(text.strip() or math_text.strip()),
            "package_text_char_count": len(text),
        }
