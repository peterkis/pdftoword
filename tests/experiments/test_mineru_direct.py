"""Contract/negative tests; fake upstream calls are synthetic, not MinerU acceptance."""

from __future__ import annotations

import json
import socket
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from docx import Document
from experiments.mineru_replacement import run_direct as run


def bundle_at(root: Path) -> Path:
    """Create an explicitly synthetic persisted package."""
    result = root / "result"
    result.mkdir(parents=True)
    writer = run.BundleWriter(result)
    writer.write_string(
        "middle_json.json",
        json.dumps({"schema": "docvortex.middle", "schema_version": "2.0", "pages": []}),
    )
    writer.write_string("markdown.md", "Synthetic test only")
    writer.write_string("structured_content.json", "{}")
    writer.write("images/one.png", b"synthetic bytes")
    run.write_json(
        root / "bundle-manifest.json",
        {
            "schema": run.SCHEMA,
            "files": run.inventory(result),
            "source": {"evidence_kind": "synthetic"},
            "parse_configuration": {"tier_requested": "standard"},
        },
    )
    return root


def test_bundle_roundtrip_and_bytes(tmp_path: Path) -> None:
    """Writer/loader preserve exact assets and schema."""
    bundle = bundle_at(tmp_path / "bundle")
    _, middle, resolver = run.load_bundle(bundle)
    assert middle["schema"] == "docvortex.middle"
    assert resolver("images/one.png") == b"synthetic bytes"


@pytest.mark.parametrize(
    "name",
    [
        "../secret",
        "/etc/passwd",
        "https://example.org/image.png",
        "file:///tmp/a",
        "images/../../a",
        "C:\\a",
        "images//a",
        "images/%2e%2e/a",
        "./images/a",
        "images/a?x",
    ],
)
def test_unsafe_paths(tmp_path: Path, name: str) -> None:
    """Reject foreign assets on both read and write."""
    with pytest.raises(run.ExperimentError, match="ASSET_PATH_INVALID"):
        run.BundleWriter(tmp_path).write(name, b"no")
    with pytest.raises(run.ExperimentError, match="ASSET_PATH_INVALID"):
        run.AssetResolver(tmp_path, {})(name)


def test_symlink(tmp_path: Path) -> None:
    """A registered asset cannot escape via a symlink."""
    (tmp_path / "images").symlink_to(tmp_path.parent, target_is_directory=True)
    with pytest.raises(run.ExperimentError, match="ASSET_SYMLINK"):
        run.BundleWriter(tmp_path).write("images/out.png", b"no")


def test_missing_changed_extra_and_unregistered(tmp_path: Path) -> None:
    """Diagnose missing, replaced and unregistered assets distinctly."""
    bundle = bundle_at(tmp_path / "bundle")
    _, _, resolver = run.load_bundle(bundle)
    with pytest.raises(run.ExperimentError, match="ASSET_NOT_REGISTERED"):
        resolver("images/two.png")
    image = bundle / "result/images/one.png"
    image.unlink()
    with pytest.raises(run.ExperimentError, match="ASSET_MISSING"):
        resolver("images/one.png")
    with pytest.raises(run.ExperimentError, match="BUNDLE_FILE_MISSING"):
        run.load_bundle(bundle)
    image.write_bytes(b"changed")
    with pytest.raises(run.ExperimentError, match="BUNDLE_HASH_MISMATCH"):
        run.load_bundle(bundle)
    with pytest.raises(run.ExperimentError, match="ASSET_HASH_MISMATCH"):
        resolver("images/one.png")
    image.write_bytes(b"synthetic bytes")
    (bundle / "result/extra.txt").write_text("extra")
    with pytest.raises(run.ExperimentError, match="BUNDLE_FILE_SET_MISMATCH"):
        run.load_bundle(bundle)


def test_schema_and_paths(tmp_path: Path) -> None:
    """Malformed manifests and MiddleJson never reach a renderer."""
    with pytest.raises(run.ExperimentError, match="BUNDLE_PATH_INVALID"):
        run.load_bundle(tmp_path / "absent")
    bundle = bundle_at(tmp_path / "bundle")
    manifest = bundle / "bundle-manifest.json"
    data = run.read_json(manifest)
    data["schema"] = "wrong"
    run.write_json(manifest, data)
    with pytest.raises(run.ExperimentError, match="BUNDLE_SCHEMA_INVALID"):
        run.load_bundle(bundle)
    data["schema"] = run.SCHEMA
    (bundle / "result/middle_json.json").write_text('{"schema":"wrong"}')
    data["files"] = run.inventory(bundle / "result")
    run.write_json(manifest, data)
    with pytest.raises(run.ExperimentError, match="MIDDLE_SCHEMA_INVALID"):
        run.load_bundle(bundle)


def test_duplicate_json(tmp_path: Path) -> None:
    """Duplicate keys must not silently replace provenance."""
    path = tmp_path / "data.json"
    path.write_text('{"schema":"one","schema":"two"}')
    with pytest.raises(run.ExperimentError, match="JSON_INVALID"):
        run.read_json(path)


def fake_runtime(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> list[str]:
    """Inject public API doubles; importing any parser immediately fails the test."""
    calls: list[str] = []
    doc = Document()
    doc.add_paragraph("Synthetic fixture only")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "cell"
    path = tmp_path / "synthetic.docx"
    doc.save(str(path))

    def render(document: Any, *, asset_resolver: Any) -> bytes:
        calls.append("render_docx")
        assert asset_resolver("images/one.png") == b"synthetic bytes"
        with pytest.raises(run.ExperimentError, match="NETWORK_DISABLED"):
            socket.socket()
        return path.read_bytes()

    modules = {
        "docvortex.schema": SimpleNamespace(MiddleJson=SimpleNamespace(from_dict=lambda x: x)),
        "docvortex.export": SimpleNamespace(validate_materialized_assets=lambda x, y: None),
        "mineru.render": SimpleNamespace(render_docx=render),
    }

    def imports(name: str) -> Any:
        calls.append(name)
        assert name in modules, "export must never import a parser"
        return modules[name]

    monkeypatch.setattr(run.importlib, "import_module", imports)
    monkeypatch.setattr(run, "require_runtime", lambda: {})
    return calls


def test_offline_export_no_parser(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Repeat export from one sealed bundle with network denied and zero parser imports."""
    bundle = bundle_at(tmp_path / "bundle")
    before = run.inventory(bundle)
    calls = fake_runtime(monkeypatch, tmp_path)
    original_socket = socket.socket
    for name in ["export1", "export2"]:
        output = tmp_path / name
        assert run.main(["export", "--bundle", str(bundle), "--output", str(output)]) == 0
        receipt = run.read_json(output / "run-receipt.json")
        assert receipt["parse_calls"] == 0
        assert receipt["docx_inventory"]["tables"] == 1
        assert receipt["word_open"] == "NOT_RUN"
        assert (output / "B.docx").is_file()
    assert calls.count("render_docx") == 2
    assert run.inventory(bundle) == before
    assert socket.socket is original_socket


def test_cli_authorization_and_existing_output(tmp_path: Path) -> None:
    """No authorization never imports runtime; existing evidence cannot be overwritten."""
    output = tmp_path / "run"
    args = [
        "parse",
        "--input",
        "absent.pdf",
        "--page",
        "1",
        "--evidence-kind",
        "real",
        "--output",
        str(output),
    ]
    assert run.main(args) == 1
    receipt = run.read_json(output / "run-receipt.json")
    assert receipt["error"] == "LOCAL_PARSE_AUTHORIZATION_REQUIRED"
    before = run.inventory(output)
    assert run.main(args) == 2
    assert before == run.inventory(output)
    args[args.index("1")] = "0"
    assert run.main(args) == 2
    with pytest.raises(SystemExit) as exc:
        run.cli().parse_args(["export"])
    assert exc.value.code == 2


def test_legacy_runtime_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    """The installed legacy DocVortex runtime cannot become the experiment runtime."""
    repo = Path(run.__file__).resolve().parents[2]
    monkeypatch.setattr(sys, "prefix", str(repo / "tmp/docx-demo/docvortex-runtime/.venv"))
    with pytest.raises(run.ExperimentError, match="ISOLATED_RUNTIME_REQUIRED"):
        run.require_runtime()


def test_renderer_failure_sanitized(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """Source-bearing upstream errors stay out of public diagnostics and receipts."""
    bundle = bundle_at(tmp_path / "bundle")
    fake_runtime(monkeypatch, tmp_path)

    def fail(name: str) -> None:
        raise ValueError("PRIVATE_SOURCE_SENTINEL")

    monkeypatch.setattr(run.importlib, "import_module", fail)
    output = tmp_path / "failed"
    assert run.main(["export", "--bundle", str(bundle), "--output", str(output)]) == 1
    assert "PRIVATE_SOURCE_SENTINEL" not in capsys.readouterr().out
    assert "PRIVATE_SOURCE_SENTINEL" not in (output / "run-receipt.json").read_text()
    assert not (output / "B.docx").exists()


def test_bad_docx() -> None:
    """Arbitrary bytes never count as a Word package."""
    with pytest.raises(run.ExperimentError, match="DOCX_INVALID"):
        run.inspect_docx(b"not a zip")


def test_parse_save_public_export_once(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Exercise the local parse branch with a fake ParseResult.save contract, no models."""
    calls = fake_runtime(monkeypatch, tmp_path)
    original_import = run.importlib.import_module
    input_path = tmp_path / "source.pdf"
    input_path.write_bytes(b"synthetic input only")

    def parse(path: str, **kwargs: Any) -> Any:
        calls.append("parse")
        assert path == str(input_path)
        assert kwargs["page_range"] == "4"
        assert kwargs["tier"] == "standard"

        def save(writer: run.BundleWriter) -> None:
            writer.write_string(
                "middle_json.json",
                json.dumps(
                    {
                        "schema": "docvortex.middle",
                        "schema_version": "2.0",
                        "pages": [],
                        "extensions": {"mineru": {"tier": "standard"}},
                    }
                ),
            )
            writer.write_string("markdown.md", "synthetic")
            writer.write_string("structured_content.json", "{}")
            writer.write("images/one.png", b"synthetic bytes")
            writer.write_string("model_output.json", "{}")

        return SimpleNamespace(save=save)

    def imports(name: str) -> Any:
        if name == "mineru.parser":
            return SimpleNamespace(parse=parse)
        return original_import(name)

    monkeypatch.setattr(run.importlib, "import_module", imports)
    output = tmp_path / "parse-output"
    assert (
        run.main(
            [
                "parse",
                "--input",
                str(input_path),
                "--page",
                "4",
                "--evidence-kind",
                "synthetic",
                "--allow-local-parse",
                "--output",
                str(output),
            ]
        )
        == 0
    )
    assert calls.count("parse") == 1
    assert calls.count("render_docx") == 1
    manifest, _, _ = run.load_bundle(output / "bundle")
    assert manifest["source"]["physical_page"] == 4
    assert manifest["source"]["sha256"] == run.digest(input_path)
    assert "model_output.json" in manifest["files"]
    assert run.read_json(output / "run-receipt.json")["parse_calls"] == 1


def test_failed_parse_no_retry(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """An ambiguous timeout retains a partial run and never resubmits."""
    fake_runtime(monkeypatch, tmp_path)
    calls = []

    def fail(*args: Any, **kwargs: Any) -> Any:
        calls.append(1)
        raise TimeoutError("source-bearing timeout")

    monkeypatch.setattr(run.importlib, "import_module", lambda _: SimpleNamespace(parse=fail))
    source = tmp_path / "source.png"
    source.write_bytes(b"synthetic")
    output = tmp_path / "failed-parse"
    assert (
        run.main(
            [
                "parse",
                "--input",
                str(source),
                "--page",
                "1",
                "--evidence-kind",
                "synthetic",
                "--allow-local-parse",
                "--output",
                str(output),
            ]
        )
        == 1
    )
    receipt = run.read_json(output / "run-receipt.json")
    assert receipt["status"] == "PARTIAL"
    assert receipt["parse_calls"] == 1
    assert receipt["retries"] == 0
    assert len(calls) == 1
    assert not (output / "B.docx").exists()


def test_export_cannot_write_into_original(tmp_path: Path) -> None:
    """A caller cannot accidentally use the saved package as its output directory."""
    bundle = bundle_at(tmp_path / "bundle")
    before = run.inventory(bundle)
    assert run.main(["export", "--bundle", str(bundle), "--output", str(bundle / "new")]) == 2
    assert run.inventory(bundle) == before
