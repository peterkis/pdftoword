"""Source-based option row planning keeps ambiguous geometry conservative."""

from __future__ import annotations

import copy
import socket
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from prototypes.docx_output.common import Json, block
from prototypes.docx_output.planning.flow import plan_flow
from prototypes.docx_output.structure import crop, image_content
from prototypes.docx_output.writer import build
from tests.demo.test_output import setup_ir


def option_case(root: Path) -> tuple[Path, Json]:
    """Build four measured figure regions with stable label/asset ownership."""
    job, ir, page = setup_ir(root)
    pairs = []
    for i, label in enumerate("ABCD"):
        box = [30.0 + i * 110, 80.0, 80.0 + i * 110, 130.0]
        text = block(f"label{i}", 0, list(box), label + ".", "inferred", "option")
        fig = block(f"figure{i}", 0, list(box), "", "ovis_ocr2", "figure")
        aid = crop(job, ir, page, list(box), fig["id"])
        fig.update(content=image_content(aid), render_policy="preserve_image")
        page["blocks"].extend([text, fig])
        page["reading_order"].extend([text["id"], fig["id"]])
        pairs.append({"label": text["id"], "figure": fig["id"]})
    ir["metadata"]["figure_groups"] = [{"page_index": 0, "kind": "option_grid", "pairs": pairs}]
    return job, ir


def test_same_row_uses_four_cells_without_source_mutation(
    private_case: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    job, ir = option_case(private_case)
    original = copy.deepcopy(ir)

    def deny(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("NETWORK_FORBIDDEN")

    monkeypatch.setattr(socket.socket, "connect", deny)
    monkeypatch.setattr(socket.socket, "connect_ex", deny)
    plan = plan_flow(ir, families={"Arial", "Arial Unicode MS"})
    group = plan.output_layout["sections"][0]["nodes"][0]
    assert group["figure_layout"]["columns"] == 4
    assert group["figure_layout"]["basis"] == "source_based"
    assert group["figure_layout"]["figure_ids"] == ["figure0", "figure1", "figure2", "figure3"]
    assert ir == original and plan.document["pages"] == ir["pages"]
    result = build(job, plan.document, "auto", output_plan=plan.output_layout)
    doc = Document(str(job / "auto.docx"))
    assert len(doc.tables[0].rows) == 1 and len(doc.tables[0].columns) == 4
    assert [c.text.strip() for c in doc.tables[0].rows[0].cells] == ["A.", "B.", "C.", "D."]
    assert result["placed_figure_count"] == 4
    assert len(doc.inline_shapes) == 4
    assert len(doc._element.xpath("//w:bookmarkStart")) == 8


@pytest.mark.parametrize(
    "mode",
    [
        "two_rows",
        "overlap",
        "reverse",
        "inferred",
        "asset_mismatch",
        "narrow",
        "explicit",
        "long_label",
        "locked",
        "inline_label_width",
    ],
)
def test_ambiguous_or_explicit_layout_retains_two_columns(private_case: Path, mode: str) -> None:
    job, ir = option_case(private_case)
    blocks = ir["pages"][0]["blocks"]
    if mode == "two_rows":
        for b in blocks[4:]:
            b["bbox"][1] += 100
            b["bbox"][3] += 100
        for a in ir["assets"][2:]:
            a["source_bbox"][1] += 100
            a["source_bbox"][3] += 100
    elif mode == "overlap":
        blocks[3]["bbox"] = list(blocks[1]["bbox"])
    elif mode == "reverse":
        blocks[1]["bbox"], blocks[7]["bbox"] = blocks[7]["bbox"], blocks[1]["bbox"]
    elif mode == "inferred":
        blocks[1]["geometry_source"] = "inferred"
    elif mode == "asset_mismatch":
        blocks[1]["bbox"][2] += 1
    elif mode == "explicit":
        ir["metadata"]["figure_groups"][0]["columns"] = 2
    elif mode == "long_label":
        blocks[0]["content"]["plain_text"] = "Long option label needs room"
    elif mode == "locked":
        blocks[1]["flags"].append("manual_layout_lock")
    if mode in {"overlap", "reverse"}:
        for i, asset in enumerate(ir["assets"]):
            asset["source_bbox"] = list(blocks[2 * i + 1]["bbox"])
    profile = {"margins_pt": [200, 36, 200, 36]} if mode == "narrow" else {}
    if mode == "inline_label_width":
        ir["metadata"]["figure_groups"][0]["inline_labels"] = True
        margin = (ir["pages"][0]["width_pt"] - 280) / 2
        profile = {"margins_pt": [margin, 36, margin, 36]}
    plan = plan_flow(ir, profile, families={"Arial", "Arial Unicode MS"})
    group = plan.output_layout["sections"][0]["nodes"][0]
    assert group["figure_layout"]["columns"] == 2
    assert group["figure_layout"]["basis"] != "source_based"
    build(job, plan.document, "auto", output_plan=plan.output_layout)
    table = Document(str(job / "auto.docx")).tables[0]
    assert len(table.rows) == 2 and len(table.columns) == 2


def test_legacy_keeps_existing_grid(private_case: Path) -> None:
    job, ir = option_case(private_case)
    build(job, ir, "auto")
    table = Document(str(job / "auto.docx")).tables[0]
    assert len(table.rows) == 2 and len(table.columns) == 2
