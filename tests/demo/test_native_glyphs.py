"""Glyph evidence must locate original runs without rewriting their text."""

import copy
from pathlib import Path
from typing import Any

import pytest
from prototypes.docx_output.common import Json
from prototypes.docx_output.native_glyphs import bind_runs


def glyph(index: int, text: str, box: list[float], **extra: Any) -> Json:
    return dict(
        index=index,
        text=text,
        bbox=box,
        generated=False,
        object_id=f"o{index}",
        paint_order=index,
        visibility="painted",
        **extra,
    )


def test_unique_geometry_binding_preserves_spaces_and_font_box() -> None:
    runs = [{"id": "r", "text": "A B", "bbox": [0, 0, 20, 12]}]
    before = copy.deepcopy(runs)
    gs = [
        glyph(0, "A", [1, 3, 5, 10]),
        glyph(1, " ", [5, 3, 6, 10]),
        glyph(2, "B", [7, 3, 11, 10]),
        glyph(3, "A", [40, 3, 45, 10]),
    ]
    evidence = bind_runs(runs, gs)
    assert runs == before
    assert evidence["r"]["status"] == "VERIFIED"
    assert evidence["r"]["char_indices"] == [0, 2]
    assert evidence["r"]["bbox"] == [1, 3, 11, 10]
    assert evidence["r"]["source_whitespace_offsets"] == [1]


def test_reused_glyph_rejects_both_runs() -> None:
    runs = [{"id": x, "text": "A", "bbox": [0, 0, 10, 10]} for x in ["a", "b"]]
    evidence = bind_runs(runs, [glyph(0, "A", [1, 1, 5, 5])])
    assert all(v["status"] == "GLYPH_REUSED" for v in evidence.values())


def test_ambiguous_missing_and_unknown_visibility_reject() -> None:
    run = [{"id": "r", "text": "A", "bbox": [0, 0, 10, 10]}]
    gs = [glyph(0, "A", [1, 1, 5, 5]), glyph(1, "A", [2, 1, 6, 5])]
    assert bind_runs(run, gs)["r"]["status"] == "AMBIGUOUS_GLYPHS"
    assert bind_runs(run, [])["r"]["status"] == "NO_GLYPH_MATCH"
    gs[0]["visibility"] = "unknown"
    assert bind_runs(run, gs[:1])["r"]["status"] != "VERIFIED"


def test_shaded_grid_with_empty_cell_needs_glyph_proof() -> None:
    from prototypes.docx_output.native_grid import ruled_regions
    from tests.demo.test_native_grid_fallback import evidence

    objects, chars = evidence()
    chars.pop()  # A real empty fourth cell.
    bindings = {}
    glyphs = []
    for i, c in enumerate(chars):
        box = c["bbox"][:]
        bindings[c["source_id"]] = {
            "status": "VERIFIED",
            "bbox": box,
            "char_indices": [i],
            "paint_orders": [20 + i],
            "object_ids": [f"t{i}"],
        }
        glyphs.append(glyph(i, "A", box))
        c["bbox"][1] -= 5  # Estimated font box crosses a line; tight ink does not.
    objects.append(
        {
            "id": "shade",
            "type": 2,
            "bbox": [0.2, 0.2, 19.8, 19.8],
            "fill_rect": [0.2, 0.2, 19.8, 19.8],
            "fill_rgba": [224, 224, 224, 255],
            "paint_order": 0,
        }
    )
    ev = {"runs": bindings, "glyphs": glyphs}
    assert ruled_regions(objects, chars) == []
    regions = ruled_regions(objects, chars, glyph_evidence=ev)
    assert len(regions) == 1
    assert regions[0]["cells"][3]["text_status"] == "NO_VISIBLE_TEXT_OBSERVED"
    assert "shade" in regions[0]["path_ids"]
    objects[-1]["paint_order"] = 100
    assert ruled_regions(objects, chars, glyph_evidence=ev) == []


def test_generated_spaces_do_not_rewrite_run() -> None:
    run = {"id": "r", "text": "AB", "bbox": [0, 0, 20, 10]}
    gs = [glyph(0, "A", [1, 1, 5, 8]), glyph(1, " ", [5, 1, 6, 8]), glyph(2, "B", [7, 1, 10, 8])]
    gs[1]["generated"] = True
    assert bind_runs([run], gs)["r"]["ignored_char_indices"] == [1]
    assert run["text"] == "AB"


def test_nonspace_unicode_is_never_normalized() -> None:
    run = {"id": "r", "text": "Ａ", "bbox": [0, 0, 20, 10]}
    assert bind_runs([run], [glyph(0, "A", [1, 1, 5, 8])])["r"]["status"] == "NO_GLYPH_MATCH"




@pytest.mark.parametrize(
    "fault",
    [
        "missing",
        "reuse",
        "crossing",
        "unclaimed",
        "image",
        "bad_fill",
        "unknown_path",
        "broken",
        "merged",
    ],
)
def test_glyph_grid_abstains_atomically(fault: str) -> None:
    from prototypes.docx_output.native_grid import ruled_regions
    from tests.demo.test_native_grid_fallback import evidence

    objects, chars = evidence()
    chars.pop()
    bindings = {
        c["source_id"]: {
            "status": "VERIFIED",
            "bbox": c["bbox"][:],
            "char_indices": [i],
            "paint_orders": [20 + i],
            "object_ids": [f"t{i}"],
        }
        for i, c in enumerate(chars)
    }
    gs = [glyph(i, "X", c["bbox"][:]) for i, c in enumerate(chars)]
    if fault == "missing":
        bindings[chars[0]["source_id"]]["status"] = "NO_GLYPH_MATCH"
    elif fault == "reuse":
        bindings[chars[1]["source_id"]]["char_indices"] = [0]
    elif fault == "crossing":
        bindings[chars[0]["source_id"]]["bbox"][2] = 25
    elif fault == "unclaimed":
        gs.append(glyph(99, "Z", [23, 23, 28, 29]))
    elif fault == "image":
        objects.append({"id": "image", "type": 3, "bbox": [2, 2, 10, 10]})
    elif fault == "bad_fill":
        objects.append(
            {
                "id": "fill",
                "type": 2,
                "bbox": [1, 1, 25, 10],
                "fill_rect": [1, 1, 25, 10],
                "fill_rgba": [224, 224, 224, 255],
                "paint_order": 0,
            }
        )
    elif fault == "unknown_path":
        objects.append({"id": "curve", "type": 2, "bbox": [2, 2, 10, 10]})
    elif fault == "broken":
        objects[1]["bbox"][2] = 15
        objects[1]["solid_line"][1][0] = 15
    else:
        objects = [o for o in objects if o["id"] != "h1"]
    assert ruled_regions(objects, chars, glyph_evidence={"runs": bindings, "glyphs": gs}) == []


@pytest.mark.parametrize("profile", ["legacy", "fidelity-v3.1"])
def test_three_shaded_tables_stay_separate_in_docx(tmp_path: Path, profile: str) -> None:
    import shutil
    import uuid

    from docx import Document
    from prototypes.docx_output.common import PRIVATE, read
    from prototypes.docx_output.pipeline import convert
    from tests.demo.synthetic import make_pdf

    commands = []
    for ti, y in enumerate([100, 220, 340]):
        commands += ["0 0 0 RG 1 w"]
        for yy in [y, y + 30, y + 60]:
            commands.append(f"60 {yy} m 260 {yy} l S")
        for x in [60, 160, 260]:
            commands.append(f"{x} {y} m {x} {y + 60} l S")
        for x in [60, 160]:
            commands.append(f"0.88 g {x + 0.2} {y + 30.2} 99.6 29.6 re f 0 g")
        for ci, ri in [(1, 1), (0, 0), (1, 0)]:
            text = f"G{ti}{ci}{ri}".encode("utf-16-be").hex()
            commands.append(f"BT /F1 9 Tf {70 + ci * 100} {y + 10 + ri * 30} Td <{text}> Tj ET")
    pdf = tmp_path / "tables.pdf"
    make_pdf(pdf, decoration="\n".join(commands).encode())
    root = PRIVATE / ("glyph-grid-test-" + uuid.uuid4().hex)
    try:
        job = convert(pdf, mode="native", output_profile=profile, output_root=root)
        ir = read(job / "layout.auto.json")
        grids = [b for b in ir["pages"][0]["blocks"] if "ruled_grid_image_fallback" in b["flags"]]
        assert len(grids) == 3
        assert all(b["type"] == "table" and b["render_policy"] == "preserve_image" for b in grids)
        doc = Document(str(job / "auto.docx"))
        assert len(doc.inline_shapes) == 4 and len(doc.tables) == 0
        assert not any("G000" in p.text for p in doc.paragraphs)
        assert all(
            len(b["content_candidates"][0]["evidence"]["run_glyph_bindings"]) == 3 for b in grids
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)
