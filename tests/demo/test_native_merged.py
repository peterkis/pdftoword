"""Merged regions require complete interval evidence, not midpoint guesses."""

from __future__ import annotations

from pathlib import Path

import pytest
from prototypes.docx_output.common import Json
from prototypes.docx_output.native_merged import merged_topology


def line(name: str, a: list[float], b: list[float], order: int = 0, white: bool = False) -> Json:
    return {
        "id": name,
        "type": 2,
        "white_line" if white else "solid_line": [a, b],
        "stroke_width_pt": 1.0,
        "line_cap": 0,
        "paint_order": order,
        "bbox": [
            min(a[0], b[0]) - 0.5,
            min(a[1], b[1]) - 0.5,
            max(a[0], b[0]) + 0.5,
            max(a[1], b[1]) + 0.5,
        ],
    }


def table() -> list[Json]:
    return [
        line("top", [0, 0], [40, 0]),
        line("bottom", [0, 40], [40, 40]),
        line("left", [0, 0], [0, 40]),
        line("right", [40, 0], [40, 40]),
        line("midv", [20, 0], [20, 40]),
        line("midh", [20, 20], [40, 20]),
    ]


def test_vertical_merge_with_t_junction() -> None:
    result = merged_topology(table(), [])
    assert len(result) == 1
    assert sorted((c["rowspan"], c["colspan"]) for c in result[0]["cells"]) == [
        (1, 1),
        (1, 1),
        (2, 1),
    ]


@pytest.mark.parametrize("fault", ["partial", "open", "nonrectangular", "white_gap"])
def test_incomplete_boundary_is_not_a_merge(fault: str) -> None:
    objects = table()
    if fault == "partial":
        objects.append(line("stub", [0, 20], [5, 20]))
    elif fault == "open":
        objects = objects[1:]
    elif fault == "nonrectangular":
        objects[-2] = line("midv", [20, 20], [20, 40])
    else:
        objects.append(line("erase", [7, 0], [13, 0], 5, True))
    assert merged_topology(objects, []) == []


def test_white_mask_then_repaint_obeys_order() -> None:
    objects = [
        *table(),
        line("erase", [7, 0], [13, 0], 5, True),
        line("restore", [6, 0], [14, 0], 6),
    ]
    assert len(merged_topology(objects, [])) == 1


@pytest.mark.parametrize("profile", ["legacy", "fidelity-v3.1"])
def test_merged_table_integrates_without_editability_claim(tmp_path: Path, profile: str) -> None:
    import shutil
    import uuid

    from docx import Document
    from prototypes.docx_output.common import PRIVATE, read
    from prototypes.docx_output.pipeline import convert
    from tests.demo.synthetic import make_pdf

    commands = [
        "0 0 0 RG 1 w 60 300 m 260 300 l S 60 360 m 260 360 l S",
        "60 300 m 60 360 l S 160 300 m 160 360 l S 260 300 m 260 360 l S",
        "160 330 m 260 330 l S",
    ]
    for x, y, text in [(70, 325, "Merged"), (170, 340, "Upper"), (170, 310, "Lower")]:
        commands.append(f"BT /F1 9 Tf {x} {y} Td <{text.encode('utf-16-be').hex()}> Tj ET")
    pdf = tmp_path / "merged.pdf"
    make_pdf(pdf, decoration="\n".join(commands).encode())
    root = PRIVATE / ("merged-test-" + uuid.uuid4().hex)
    try:
        job = convert(pdf, mode="native", output_profile=profile, output_root=root)
        ir = read(job / "layout.auto.json")
        grids = [b for b in ir["pages"][0]["blocks"] if "ruled_grid_image_fallback" in b["flags"]]
        assert len(grids) == 1 and grids[0]["type"] == "table"
        assert len(grids[0]["content_candidates"][0]["evidence"]["cells"]) == 3
        doc = Document(str(job / "auto.docx"))
        assert not doc.tables
        assert "Merged" not in "".join(p.text for p in doc.paragraphs)
    finally:
        shutil.rmtree(root, ignore_errors=True)


@pytest.mark.parametrize(
    "fault", ["unknown_object", "missing_glyph", "cross_cell", "foreground_fill", "white_mask"]
)
def test_merged_ownership_rejects_unproven_content(fault: str) -> None:
    from prototypes.docx_output.native_merged import merged_regions

    objects = table()
    chars = [{"source_id": "r", "text": "A", "bbox": [3, 3, 12, 12]}]
    witness = {
        "status": "VERIFIED",
        "bbox": [3, 3, 12, 12],
        "char_indices": [0],
        "paint_orders": [10],
    }
    glyph = {
        "index": 0,
        "text": "A",
        "bbox": [3, 3, 12, 12],
        "visibility": "painted",
        "generated": False,
        "paint_order": 10,
    }
    ev = {"runs": {"r": witness}, "glyphs": [glyph]}
    if fault == "unknown_object":
        objects.append({"id": "image", "type": 3, "bbox": [3, 20, 12, 28]})
    elif fault == "missing_glyph":
        ev["glyphs"] = []
    elif fault == "white_mask":
        objects.append(line("mask", [2, 7], [13, 7], 11, True))
    elif fault == "cross_cell":
        witness["bbox"] = [12, 3, 25, 12]
    else:
        objects.append(
            {
                "id": "fill",
                "type": 2,
                "fill_rect": [1, 1, 15, 15],
                "bbox": [1, 1, 15, 15],
                "paint_order": 11,
                "fill_rgba": [255, 255, 255, 255],
            }
        )
    assert merged_regions(objects, chars, ev, []) == []
