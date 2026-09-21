"""Inline math text must retain exact source style boundaries without network IO."""

from __future__ import annotations

import copy
import socket
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from prototypes.docx_output.common import Json, block, read
from prototypes.docx_output.formula import to_omml
from prototypes.docx_output.planning.flow import plan_flow
from prototypes.docx_output.writer import build
from tests.demo.test_output import setup_ir


def styled_case(root: Path) -> tuple[Path, Json, Json]:
    """Construct mixed script text with two adjacent formulas and measured run boundaries."""
    job, ir, page = setup_ir(root)
    texts = ["中文 ", "Latin $x$$y$ 后", "文。 "]
    b = block("mixed", 0, [10, 10, 300, 40], "".join(texts), "native_pdf")
    b["content"]["runs"] = [
        {
            "text": text,
            "bbox": b["bbox"],
            "font_family": font,
            "font_size_pt": size,
            "bold": i == 0,
            "italic": i == 1,
            "underline": i == 2,
            "superscript": i == 0,
            "subscript": i == 2,
            "color": None,
            "confidence": 0,
            "source_ref": f"source-{i}",
        }
        for i, (text, font, size) in enumerate(
            zip(texts, ["Songti SC", "Arial", "Times New Roman"], [14, 12, 10], strict=True)
        )
    ]
    page["blocks"], page["reading_order"] = [b], [b["id"]]
    ir["metadata"]["inline_parts"] = {
        b["id"]: [
            {"text": "中文 Latin "},
            {"source_text": "$x$", "omml": to_omml("x")},
            {"source_text": "$y$", "omml": to_omml("y")},
            {"text": " 后文。 "},
        ]
    }
    return job, ir, b


@pytest.mark.parametrize("flow", [False, True])
@pytest.mark.parametrize("locked", [False, True])
def test_inline_source_style_boundaries(
    private_case: Path, monkeypatch: pytest.MonkeyPatch, flow: bool, locked: bool
) -> None:
    job, ir, b = styled_case(private_case)
    if locked:
        b["flags"].append("manual_text_lock")
    original = copy.deepcopy(ir)

    def denied(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("NETWORK_FORBIDDEN")

    monkeypatch.setattr(socket.socket, "connect", denied)
    monkeypatch.setattr(socket.socket, "connect_ex", denied)
    plan = (
        plan_flow(ir, {"editable_styles": True}, families={"Songti SC", "Arial", "Times New Roman"})
        if flow
        else None
    )
    build(
        job, plan.document if plan else ir, "auto", output_plan=plan.output_layout if plan else None
    )
    assert ir == original
    doc = Document(str(job / "auto.docx"))
    p = next(p for p in doc.paragraphs if p.text)
    assert [r.text for r in p.runs] == ["中文 ", "Latin ", " 后", "文。 "]
    assert [r.font.size.pt if r.font.size else None for r in p.runs] == [14, 12, 12, 10]
    assert p.runs[0].bold and p.runs[0].font.superscript
    assert p.runs[1].italic and p.runs[2].italic
    assert p.runs[3].underline and p.runs[3].font.subscript
    assert p.runs[0]._r.get_or_add_rPr().get_or_add_rFonts().get(qn("w:eastAsia")) == "Songti SC"
    assert len(p._p.xpath("./m:oMath")) == 2
    children: list[Any] = list(p._p)
    assert children.index(p._p.xpath("./w:bookmarkStart")[0]) < children.index(p.runs[0]._r)
    assert children.index(p._p.xpath("./w:bookmarkEnd")[0]) > children.index(p.runs[-1]._r)
    spans = read(job / "source-map.auto.json")["blocks"][0]["inline_text_spans"]
    assert [s["source_run_index"] for s in spans] == [0, 1, 1, 2]
    assert [s["source_ref"] for s in spans] == ["source-0", "source-1", "source-1", "source-2"]


@pytest.mark.parametrize("mismatch", [False, True])
@pytest.mark.parametrize("locked", [False, True])
def test_inline_unknown_style_is_explicit_inheritance(
    private_case: Path, mismatch: bool, locked: bool
) -> None:
    """Missing or stale source runs cannot lend false emphasis to inline text."""
    job, ir, b = styled_case(private_case)
    if mismatch:
        b["content"]["runs"][0]["text"] = "stale"
    else:
        b["content"]["runs"] = []
    if locked:
        b["flags"].append("manual_text_lock")
    plan = plan_flow(ir, {"editable_styles": True, "body_size_pt": 13}, families={"Arial"})
    build(job, plan.document, "auto", output_plan=plan.output_layout)
    p = next(p for p in Document(str(job / "auto.docx")).paragraphs if p.text)
    assert p.text == "中文 Latin  后文。 "
    assert all(r.bold is not True for r in p.runs)
    assert all(
        (r.font.size is not None and r.font.size.pt == 13) if locked else r.font.size is None
        for r in p.runs
    )
    spans = read(job / "source-map.auto.json")["blocks"][0]["inline_text_spans"]
    assert all(
        s["source_run_index"] is None and s["style_basis"] == "inherited_output_style"
        for s in spans
    )


def test_inline_image_fallback_retains_adjacent_text(private_case: Path) -> None:
    """The existing image fallback keeps its bytes, neighbours and diagnostics."""
    from prototypes.docx_output.structure import crop

    job, ir, b = styled_case(private_case)
    aid = crop(job, ir, ir["pages"][0], b["bbox"], "fallback", "formula_image")
    item = ir["metadata"]["inline_parts"][b["id"]][1]
    del item["omml"]
    item["asset_id"] = aid
    before = copy.deepcopy(ir)
    result = build(job, ir, "auto")
    assert result["formula_image_count"] == 1
    assert ir == before
    doc = Document(str(job / "auto.docx"))
    assert len(doc.inline_shapes) == 1
    assert next(p.text for p in doc.paragraphs if p.text) == "中文 Latin  后文。 "
