"""Visual table reversal must not pass XML-order topology checks."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("where", ["table", "style"])
@pytest.mark.parametrize("enabled", [True, False])
def test_bidi_visual_table(tmp_path: Path, where: str, enabled: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.tables[0]._tbl.tblPr
    if where == "style":
        style = doc.styles.add_style("ReversedTable", WD_STYLE_TYPE.TABLE)
        props = OxmlElement("w:tblPr")
        style._element.append(props)
        doc.tables[0].style = style
    node = OxmlElement("w:bidiVisual")
    node.set(qn("w:val"), "true" if enabled else "false")
    props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_BIDI_OVERRIDE" in result["errors"]) is enabled
    assert result["structure_status"] == ("FAIL" if enabled else "PASS")
