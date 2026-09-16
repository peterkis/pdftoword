"""Independent actual-DOCX acceptance checks, using handwritten references."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document

from acceptance.metrics import evaluate


def test_missing_text_is_a_deletion_not_success(tmp_path: Path) -> None:
    doc: Any = Document()
    doc.add_paragraph("AB")
    output = tmp_path / "actual.docx"
    doc.save(str(output))
    truth = {
        "anchors": [{"id": "a", "page": 1, "bbox": [0, 0, 100, 20]}],
        "units": [
            {
                "unit_id": "u",
                "page": 1,
                "kind": "text",
                "status": "confirmed",
                "reference": {"text": "ABC", "source_anchor_id": "a"},
            }
        ],
    }
    # An unbound paragraph must not be credited to an arbitrary source page.
    result = evaluate(
        output,
        truth,
        {"pages": [{"page": 1, "width": 100, "height": 100}], "blocks": [], "relations": []},
    )
    assert result["metrics"]["text"]["raw"]["deletions"] == 3
    assert result["content_status"] == "FAIL"
    assert result["rendering_status"] == "PENDING"


def test_actual_table_cells_are_not_counted_as_body(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc: Any = Document()
    table = doc.add_table(rows=1, cols=2)
    for cell, text in zip(table.rows[0].cells, ["X", "Y"], strict=True):
        cell.text = text
    marker = OxmlElement("w:bookmarkStart")
    marker.set(qn("w:name"), "table_source")
    marker.set(qn("w:id"), "1")
    table.cell(0, 0).paragraphs[0]._p.insert(0, marker)
    path = tmp_path / "table.docx"
    doc.save(str(path))
    truth = {
        "anchors": [{"id": "t", "page": 1, "bbox": [0, 0, 100, 20]}],
        "units": [
            {
                "unit_id": "t",
                "kind": "table",
                "page": 1,
                "status": "confirmed",
                "reference": {
                    "source_anchor_id": "t",
                    "rows": 1,
                    "cols": 2,
                    "is_data_table": True,
                    "cells": [
                        {"row": 0, "col": 0, "rowspan": 1, "colspan": 1, "text": "X"},
                        {"row": 0, "col": 1, "rowspan": 1, "colspan": 1, "text": "Y"},
                    ],
                },
            }
        ],
    }
    result = evaluate(
        path,
        truth,
        {
            "pages": [{"page": 1, "width": 100, "height": 100}],
            "blocks": [
                {
                    "marker": "table_source",
                    "block_id": "b",
                    "page": 1,
                    "bbox": [0, 0, 100, 20],
                    "kind": "table",
                }
            ],
            "relations": [],
        },
    )
    assert result["metrics"]["tables"]["correct_count"] == 1
    assert result["metrics"]["table_cells"]["correct_count"] == 2
    assert result["metrics"]["text"]["raw"]["reference_chars"] == 0


# Hand-authored fixture: two source pages, one formula, one table, one figure.
# It does not call the production converter to manufacture expected output.
def specimen(root: Path, fault: str = "") -> tuple[Path, dict, dict]:
    import copy
    import hashlib
    import zipfile

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from PIL import Image

    root.mkdir(exist_ok=True)
    doc: Any = Document()
    anchors: list[dict] = []
    blocks: list[dict] = []
    units: list[dict] = []

    def bind(paragraph: object, name: str, page: int, box: list[int], kind: str) -> None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:name"), name)
        start.set(qn("w:id"), str(len(blocks)))
        paragraph._p.insert(0, start)  # type: ignore[attr-defined]
        anchors.append({"id": name, "page": page, "bbox": box})
        blocks.append(
            {
                "marker": name,
                "block_id": name,
                "page": page,
                "bbox": box,
                "kind": kind,
                "fallback": False,
            }
        )

    for name, page, text in [("a", 1, "ABC"), ("b", 2, "DEF")]:
        p = doc.add_paragraph(text)
        bind(p, name, page, [0, 0, 100, 20], "paragraph")
        units.append(
            {
                "unit_id": name,
                "page": page,
                "kind": "text",
                "status": "confirmed",
                "reference": {"text": text, "source_anchor_id": name},
            }
        )
    p = doc.add_paragraph()
    bind(p, "f", 1, [0, 30, 100, 50], "formula")
    math = OxmlElement("m:oMath")
    run, math_text = OxmlElement("m:r"), OxmlElement("m:t")
    math_text.text = "x=1"
    run.append(math_text)
    math.append(run)
    p._p.append(math)
    units.append(
        {
            "unit_id": "f",
            "page": 1,
            "kind": "formula",
            "status": "confirmed",
            "reference": {"text": "x=1", "source_anchor_id": "f"},
        }
    )
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "12", "34"
    bind(table.cell(0, 0).paragraphs[0], "t", 1, [0, 60, 100, 80], "table")
    units.append(
        {
            "unit_id": "t",
            "page": 1,
            "kind": "table",
            "status": "confirmed",
            "reference": {
                "source_anchor_id": "t",
                "rows": 1,
                "cols": 2,
                "is_data_table": True,
                "cells": [
                    {"row": 0, "col": 0, "rowspan": 1, "colspan": 1, "text": "12"},
                    {"row": 0, "col": 1, "rowspan": 1, "colspan": 1, "text": "34"},
                ],
            },
        }
    )
    Image.new("RGB", (20, 20), "blue").save(root / "figure.png")
    picture = doc.add_paragraph()
    picture.add_run().add_picture(str(root / "figure.png"), width=Inches(1))
    bind(picture, "g", 1, [0, 90, 100, 190], "figure")
    blocks[-1]["images"] = [
        {
            "sha256": hashlib.sha256((root / "figure.png").read_bytes()).hexdigest(),
            "bbox": [0, 90, 100, 190],
            "fallback": False,
        }
    ]
    units.append(
        {
            "unit_id": "edge",
            "page": 1,
            "kind": "figure_edge",
            "status": "confirmed",
            "reference": {"from": "g", "to": "a", "relation": "belongs_to"},
        }
    )
    units.append(
        {
            "unit_id": "order",
            "page": 1,
            "kind": "reading_order",
            "status": "confirmed",
            "reference": {"anchors": ["a", "b"]},
        }
    )
    sources: dict[str, Any] = {
        "pages": [{"page": i, "width": 200, "height": 200} for i in [1, 2]],
        "blocks": blocks,
        "relations": [{"from": "g", "to": "a", "type": "belongs_to"}],
    }
    if fault == "missing_char":
        doc.paragraphs[0].runs[0].text = "AB"
    elif fault == "wrong_number":
        table.cell(0, 0).paragraphs[0].runs[0].text = "99"
    elif fault == "duplicate_paragraph":
        doc._element.body.insert(1, copy.deepcopy(doc.paragraphs[0]._p))
    elif fault == "missing_page":
        doc.paragraphs[0].runs[0].text = "ABC" * 100
        doc.paragraphs[1]._p.getparent().remove(doc.paragraphs[1]._p)
    elif fault == "wrong_edge":
        sources["relations"][0]["to"] = "b"
    elif fault == "wrong_source_page":
        blocks[1]["page"] = 3
    elif fault == "missing_formula":
        p._p.remove(math)
    elif fault == "wrong_formula":
        math_text.text = "x=2"
    elif fault == "wrong_span":
        table.cell(0, 0).merge(table.cell(0, 1))
    elif fault == "wrong_order":
        doc._element.body.insert(0, doc.paragraphs[1]._p)
    elif fault == "missing_figure":
        picture._p.getparent().remove(picture._p)
    elif fault == "extra_body":
        doc.add_paragraph("unrelated extra text")
    path = root / "actual.docx"
    doc.save(str(path))
    if fault in {"corrupt_media", "broken_xml", "missing_media"}:
        with zipfile.ZipFile(path) as archive:
            members = {n: archive.read(n) for n in archive.namelist()}
        media = next(n for n in members if n.startswith("word/media/"))
        if fault == "corrupt_media":
            members[media] = b"corrupt"
        elif fault == "missing_media":
            del members[media]
        else:
            members["word/document.xml"] = b"<broken"
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in members.items():
                archive.writestr(name, data)
    return path, {"anchors": anchors, "units": units}, sources


def test_clean_specimen_and_deterministic_scores(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    result = evaluate(path, truth, sources)
    assert result["errors"] == []
    assert result["failures"] == []
    assert result["metrics"]["text"]["raw"]["reference_chars"] == 6
    assert result["metrics"]["formulas"]["correct_count"] == 1
    assert result["semantic_sha256"] == evaluate(path, truth, sources)["semantic_sha256"]


@pytest.mark.parametrize(
    "fault,expected",
    [
        ("missing_char", "TEXT_MISMATCH"),
        ("wrong_number", "TABLE_MISMATCH"),
        ("duplicate_paragraph", "DUPLICATE_SOURCE_BLOCK"),
        ("missing_page", "MISSING_TEXT"),
        ("wrong_edge", "FIGURE_EDGE_MISMATCH"),
        ("wrong_source_page", "MISSING_TEXT"),
        ("missing_formula", "FORMULA_MISMATCH"),
        ("wrong_formula", "FORMULA_MISMATCH"),
        ("wrong_span", "TABLE_MISMATCH"),
        ("wrong_order", "READING_ORDER_MISMATCH"),
        ("missing_figure", "FIGURE_EDGE_MISMATCH"),
        ("extra_body", "UNALIGNED_EDITABLE_TEXT"),
        ("corrupt_media", "CORRUPT_MEDIA"),
        ("broken_xml", "XMLSYNTAXERROR"),
        ("missing_media", "MISSING_RELATIONSHIP_TARGET"),
    ],
)
def test_faulty_real_docx_is_detected(tmp_path: Path, fault: str, expected: str) -> None:
    path, truth, sources = specimen(tmp_path, fault)
    result = evaluate(path, truth, sources)
    assert expected in result["errors"] + [f["code"] for f in result["failures"]]
    assert "FAIL" in (result["content_status"], result["structure_status"])


def test_empty_denominator_and_unsupported_math_are_explicit(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    truth["units"] = [u for u in truth["units"] if u["kind"] == "formula"]
    truth["units"][0]["reference"]["text"] = r"\begin{matrix}x\end{matrix}"
    result = evaluate(path, truth, sources)
    assert result["metrics"]["text"]["raw"]["value"] is None
    assert result["metrics"]["formulas"]["unsupported_count"] == 1
    assert result["metrics"]["formulas"]["scored_count"] == 0


def test_fallback_uses_union_not_sum(tmp_path: Path) -> None:
    from acceptance.metrics import union_area

    assert union_area([[0, 0, 10, 10], [5, 0, 15, 10]]) == 150


def test_paragraph_with_two_reference_lines_counts_output_once(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    truth["units"] = [u for u in truth["units"] if u["unit_id"] == "a"]
    truth["units"][0]["reference"]["text"] = "AB"
    truth["units"].append(
        {
            "unit_id": "a2",
            "page": 1,
            "kind": "text",
            "status": "confirmed",
            "reference": {"text": "C", "source_anchor_id": "a2"},
        }
    )
    truth["anchors"] = [
        {"id": "a", "page": 1, "bbox": [0, 0, 100, 10]},
        {"id": "a2", "page": 1, "bbox": [0, 10, 100, 20]},
    ]
    result = evaluate(path, truth, sources)
    assert result["metrics"]["text"]["raw"]["reference_chars"] == 3
    assert result["metrics"]["text"]["matched"]["distance"] == 0
    assert result["metrics"]["text"]["raw"]["insertions"] == 7


def test_sealed_rescoring_rejects_tampering_and_never_overwrites(tmp_path: Path) -> None:
    from product_acceptance import digest, evaluate_only, write

    path, truth, sources = specimen(tmp_path)
    (tmp_path / "input.pdf").write_bytes(b"synthetic-input-identity")
    source_hash = digest(tmp_path / "input.pdf")
    truth["source_sha256"] = source_hash
    sources["docx_sha256"] = digest(path)
    path.rename(tmp_path / "auto.docx")
    write(tmp_path / "truth.json", truth)
    write(tmp_path / "source-map.json", sources)
    run = {
        "source_sha256": source_hash,
        "selected_pages": [1, 2],
        "annotation_sha256": digest(tmp_path / "truth.json"),
        "sample_id": "synthetic",
        "document_family": "synthetic",
        "source_commit": "test",
        "source_tree_sha256": "test",
        "dirty": True,
        "profile": {"network": "offline"},
        "rendering": {"status": "PENDING"},
    }
    write(tmp_path / "run.json", run)
    write(
        tmp_path / "seal.json",
        {
            n: digest(tmp_path / n)
            for n in ["input.pdf", "auto.docx", "truth.json", "source-map.json", "run.json"]
        },
    )
    first = evaluate_only(tmp_path, tmp_path / "score-1.json")
    second = evaluate_only(tmp_path, tmp_path / "score-2.json")
    assert first["semantic_sha256"] == second["semantic_sha256"]
    with pytest.raises(FileExistsError):
        evaluate_only(tmp_path, tmp_path / "score-1.json")
    (tmp_path / "auto.docx").write_bytes(b"tampered")
    with pytest.raises(ValueError, match="EVIDENCE_HASH_MISMATCH"):
        evaluate_only(tmp_path, tmp_path / "score-3.json")


def test_cli_and_upload_api_run_real_shared_pipeline_once(tmp_path: Path) -> None:
    import json
    import shutil
    import subprocess
    import sys
    import uuid

    import pypdfium2 as pdfium  # type: ignore[import-untyped]

    from acceptance.docx_reader import inspect
    from product_acceptance import ROOT

    # A blank PDF is enough to exercise file bytes, selected page and both entry paths.
    # Content fidelity is tested separately with the handwritten DOCX specimen above.
    source = tmp_path / "中文 空格.pdf"
    pdf = pdfium.PdfDocument.new()
    page = pdf.new_page(200, 200)
    page.close()
    pdf.save(source)
    pdf.close()
    results = []
    for entry in ["cli", "api"]:
        jobs = ROOT / "tmp/docx-demo" / ("acceptance-test-" + uuid.uuid4().hex)
        try:
            process = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "scripts/acceptance/worker.py"),
                    "--entry",
                    entry,
                    "--input",
                    str(source),
                    "--pages",
                    "1",
                    "--jobs",
                    str(jobs),
                ],
                capture_output=True,
                text=True,
                timeout=90,
            )
            assert process.returncode == 0, process.stderr
            result = json.loads(process.stdout)
            assert result["pipeline_calls"] == 1
            assert result["network_attempts"] == 0
            job = Path(result["job"])
            sources = json.loads((job / "source-map.auto.json").read_text())
            assert [p["page"] for p in sources["pages"]] == [1]
            results.append(inspect(job / "auto.docx")["paragraphs"])
        finally:
            if jobs.exists():
                shutil.rmtree(jobs)
    assert results[0] == results[1]


def test_result_schema_and_reviewed_revision_preserve_auto(tmp_path: Path) -> None:
    import json

    from jsonschema import Draft202012Validator

    from product_acceptance import ROOT, digest, import_reviewed, write

    bundle = tmp_path / "auto-bundle"
    path, truth, sources = specimen(bundle)
    path.rename(bundle / "auto.docx")
    (bundle / "input.pdf").write_bytes(b"synthetic-only")
    truth["source_sha256"] = digest(bundle / "input.pdf")
    sources["docx_sha256"] = digest(bundle / "auto.docx")
    write(bundle / "truth.json", truth)
    write(bundle / "source-map.json", sources)
    write(
        bundle / "run.json",
        {
            "source_sha256": truth["source_sha256"],
            "annotation_sha256": digest(bundle / "truth.json"),
            "selected_pages": [1, 2],
            "sample_id": "synthetic",
            "document_family": "synthetic",
            "source_tree_sha256": "test",
            "source_commit": "test",
            "dirty": True,
            "profile": {},
            "rendering": {"status": "PENDING"},
        },
    )
    write(bundle / "seal.json", {p.name: digest(p) for p in bundle.iterdir() if p.is_file()})
    auto_hash = digest(bundle / "auto.docx")
    reviewed = tmp_path / "reviewed.docx"
    doc = Document(str(bundle / "auto.docx"))
    doc.paragraphs[0].runs[0].text = "changed"
    doc.save(str(reviewed))
    result = import_reviewed(bundle, reviewed, tmp_path / "review-bundle")
    assert result["revision"] == "reviewed"
    assert result["auto_sha256"] == auto_hash == digest(bundle / "auto.docx")
    assert result["rendering_status"] == "PENDING"
    assert result["human_acceptance"] == "PENDING"
    assert result["content_status"] == "FAIL"
    schema = json.loads((ROOT / "specs/product-quality/result.schema.json").read_text())
    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema).validate(result)
    result["metrics"]["text"]["eligible_count"] = -1
    assert list(Draft202012Validator(schema).iter_errors(result))


def test_omml_script_is_not_flat_text(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement

    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    math = doc.paragraphs[2]._p[-1]
    math.clear()
    sup = OxmlElement("m:sSup")
    for tag, value in [("m:e", "x"), ("m:sup", "2")]:
        field, run, text = OxmlElement(tag), OxmlElement("m:r"), OxmlElement("m:t")
        text.text = value
        run.append(text)
        field.append(run)
        sup.append(field)
    math.append(sup)
    doc.save(str(path))
    formula = next(u for u in truth["units"] if u["kind"] == "formula")
    formula["reference"]["text"] = "x^2"
    assert evaluate(path, truth, sources)["metrics"]["formulas"]["correct_count"] == 1
    formula["reference"]["text"] = "x2"
    assert evaluate(path, truth, sources)["metrics"]["formulas"]["correct_count"] == 0


def test_extra_wrong_edge_cannot_leave_overall_pass(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    sources["relations"].append({"from": "g", "to": "b", "type": "belongs_to"})
    result = evaluate(path, truth, sources)
    assert result["metrics"]["relation_precision"]["value"] == 0.5
    assert result["structure_status"] == "FAIL"


def test_unexpected_text_in_formula_paragraph_is_not_hidden(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    doc.paragraphs[2].add_run("UNEXPECTED BODY PARAGRAPH")
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNALIGNED_EDITABLE_TEXT" in result["errors"]
    assert result["structure_status"] == "FAIL"


def test_barless_fraction_is_not_normal_fraction(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    math = doc.paragraphs[2]._p[-1]
    math.clear()
    fraction = OxmlElement("m:f")
    prop, kind = OxmlElement("m:fPr"), OxmlElement("m:type")
    kind.set(qn("m:val"), "noBar")
    prop.append(kind)
    fraction.append(prop)
    for tag, value in [("m:num", "1"), ("m:den", "2")]:
        field, run, text = OxmlElement(tag), OxmlElement("m:r"), OxmlElement("m:t")
        text.text = value
        run.append(text)
        field.append(run)
        fraction.append(field)
    math.append(fraction)
    doc.save(str(path))
    next(u for u in truth["units"] if u["kind"] == "formula")["reference"]["text"] = r"\frac{1}{2}"
    assert evaluate(path, truth, sources)["metrics"]["formulas"]["correct_count"] == 0
    fraction.remove(prop)
    doc.save(str(path))
    assert evaluate(path, truth, sources)["metrics"]["formulas"]["correct_count"] == 1


@pytest.mark.parametrize("missing", ["registry", "file"])
def test_writer_keeps_safe_asset_missing_code(tmp_path: Path, missing: str) -> None:
    import json
    import shutil
    import uuid

    from prototypes.docx_output.common import DemoError
    from prototypes.docx_output.pipeline import convert
    from prototypes.docx_output.writer import build
    from tests.demo.synthetic import make_pdf

    from product_acceptance import ROOT

    source = tmp_path / "synthetic.pdf"
    make_pdf(source)
    jobs = ROOT / "tmp/docx-demo" / ("acceptance-assets-" + uuid.uuid4().hex)
    try:
        job = convert(source, mode="native", output_root=jobs)
        ir = json.loads((job / "layout.auto.json").read_text())
        image_block = next(
            b for p in ir["pages"] for b in p["blocks"] if b["content"]["kind"] == "image"
        )
        asset = next(a for a in ir["assets"] if a["id"] == image_block["content"]["asset_id"])
        if missing == "registry":
            ir["assets"].remove(asset)
        else:
            (job / asset["path"]).unlink()
        with pytest.raises(DemoError, match=r"^ASSET_MISSING$"):
            build(job, ir, "reviewed")
    finally:
        if jobs.exists():
            shutil.rmtree(jobs)


def test_missing_office_keeps_real_structure_checks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import shutil
    import uuid

    from prototypes.docx_output.common import new_job, save
    from prototypes.docx_output.render import render

    from acceptance.docx_reader import inspect
    from product_acceptance import ROOT

    jobs = ROOT / "tmp/docx-demo" / ("acceptance-render-" + uuid.uuid4().hex)
    try:
        job = new_job(jobs)
        path, _, _ = specimen(tmp_path)
        shutil.copyfile(path, job / "auto.docx")
        save(job / "qa.json", {})
        monkeypatch.setattr(shutil, "which", lambda _: None)
        original = Path.is_file
        monkeypatch.setattr(
            Path, "is_file", lambda p: False if p.name == "soffice" else original(p)
        )
        qa = render(job)
        assert qa["render_status"] == "DOCX_VISUAL_REVIEW_PENDING"
        assert qa["renderer"] is None
        assert inspect(job / "auto.docx")["errors"] == []
        assert not (job / "rendered").exists()
    finally:
        if jobs.exists():
            shutil.rmtree(jobs)


def test_duplicate_formula_in_body_is_detected(tmp_path: Path) -> None:
    import copy

    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    doc.paragraphs[0]._p.append(copy.deepcopy(doc.paragraphs[2]._p[-1]))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNALIGNED_FORMULA" in result["errors"]
    assert result["structure_status"] == "FAIL"
