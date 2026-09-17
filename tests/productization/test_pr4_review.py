"""Regressions for the external Codex review of PR 4."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("wrong", [False, True])
def test_partially_supported_formulas_distinguish_unscored_from_wrong(
    tmp_path: Path, wrong: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    truth["anchors"].append({"id": "unsupported", "page": 2, "bbox": [0, 30, 100, 50]})
    truth["units"].append(
        {
            "unit_id": "unsupported",
            "page": 2,
            "kind": "formula",
            "status": "confirmed",
            "reference": {
                "text": r"\begin{matrix}x\end{matrix}",
                "source_anchor_id": "unsupported",
            },
        }
    )
    # Unsupported semantics remain reviewable only when the expression is retained.
    doc: Any = Document(str(path))
    paragraph = doc.add_paragraph()
    marker = OxmlElement("w:bookmarkStart")
    marker.set(qn("w:name"), "unsupported")
    marker.set(qn("w:id"), "99")
    paragraph._p.append(marker)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "99")
    paragraph._p.append(end)
    math = OxmlElement("m:oMath")
    parent = math
    for tag in ["m:m", "m:mr", "m:e", "m:r", "m:t"]:
        child = OxmlElement(tag)
        parent.append(child)
        parent = child
    parent.text = "x"
    paragraph._p.append(math)
    doc.save(str(path))
    sources["blocks"].append(
        {
            "block_id": "unsupported",
            "marker": "unsupported",
            "page": 2,
            "bbox": [0, 30, 100, 50],
            "kind": "formula",
        }
    )
    if wrong:
        next(u for u in truth["units"] if u["unit_id"] == "f")["reference"]["text"] = "x=2"
    result = evaluate(path, truth, sources)
    metric = result["metrics"]["formulas"]
    assert (metric["eligible_count"], metric["scored_count"], metric["not_scored_count"]) == (
        2,
        1,
        1,
    )
    assert metric["value"] == (0 if wrong else 0.5)
    assert metric["status"] == ("FAIL" if wrong else "REVIEW_REQUIRED")
    assert result["structure_status"] == ("FAIL" if wrong else "REVIEW_REQUIRED")
    assert result["content_status"] == ("FAIL" if wrong else "REVIEW_REQUIRED")


@pytest.mark.parametrize(
    "kind,property_value,correct",
    [
        ("f", "bar", True),
        ("f", None, True),
        ("f", "noBar", False),
        ("f", "skw", False),
        ("sSup", None, True),
        ("sSub", None, True),
        ("sSubSup", None, True),
    ],
)
def test_office_default_math_properties_are_accepted_without_erasing_semantics(
    tmp_path: Path, kind: str, property_value: str | None, correct: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    math = doc.paragraphs[2]._p[-1]
    math.clear()
    expression = OxmlElement("m:" + kind)
    prop = OxmlElement("m:" + kind + "Pr")
    control = OxmlElement("m:ctrlPr")
    control.append(OxmlElement("w:rPr"))
    prop.append(control)
    if property_value is not None:
        type_node = OxmlElement("m:type")
        type_node.set(qn("m:val"), property_value)
        prop.insert(0, type_node)
    expression.append(prop)
    fields, expected = {
        "f": ([("num", "1"), ("den", "2")], r"\frac{1}{2}"),
        "sSup": ([("e", "x"), ("sup", "2")], "x^2"),
        "sSub": ([("e", "x"), ("sub", "2")], "x_2"),
        "sSubSup": ([("e", "x"), ("sub", "1"), ("sup", "2")], "x_1^2"),
    }[kind]
    for tag, value in fields:
        field, run, text = OxmlElement("m:" + tag), OxmlElement("m:r"), OxmlElement("m:t")
        text.text = value
        run.append(text)
        field.append(run)
        expression.append(field)
    math.append(expression)
    doc.save(str(path))
    next(u for u in truth["units"] if u["kind"] == "formula")["reference"]["text"] = expected
    result = evaluate(path, truth, sources)
    assert result["metrics"]["formulas"]["correct_count"] == int(correct)
