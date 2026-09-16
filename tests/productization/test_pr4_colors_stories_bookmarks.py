"""Validate colors, dynamic content, story roots and source markers."""
import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("case", ["text", "math", "style", "default", "black", "auto"])
def test_foreground_visibility(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if case == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif case == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif case == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    color = OxmlElement("w:color")
    color.set(qn("w:val"), {"black": "000000", "auto": "auto"}.get(case, "FFFFFF"))
    props.append(color)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_VISIBILITY_STYLE" in result["errors"]) is (case not in {"black", "auto"})


@pytest.mark.parametrize("where", ["body", "header", "footer"])
def test_dynamic_page_number(tmp_path: Path, where: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = (
        doc.paragraphs[0] if where == "body" else getattr(doc.sections[0], where).paragraphs[0]
    )
    paragraph.add_run()._r.append(OxmlElement("w:pgNum"))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    code = "UNSUPPORTED_BODY_CONTENT" if where == "body" else "UNSUPPORTED_VISIBLE_STORY"
    assert code in result["errors"]


@pytest.mark.parametrize("kind", ["header", "footer"])
@pytest.mark.parametrize("fault", ["wrong_root", "wrong_namespace", "valid"])
def test_story_root(tmp_path: Path, kind: str, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    getattr(doc.sections[0], kind).paragraphs[0].text = ""
    doc.save(str(path))
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    name = "word/" + kind + "1.xml"
    root = etree.fromstring(files[name])
    if fault == "wrong_root":
        root.tag = qn("w:styles")
    elif fault == "wrong_namespace":
        root.tag = "{urn:invalid}" + etree.QName(root).localname
    files[name] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for n, data in files.items():
            archive.writestr(n, data)
    result = evaluate(path, truth, sources)
    assert ("OPC_STORY_ROOT_INVALID" in result["errors"]) is (fault != "valid")


@pytest.mark.parametrize(
    "fault", ["missing", "duplicate_name", "duplicate_id", "orphan", "id", "order", "valid"]
)
def test_bookmark_pairs(tmp_path: Path, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0]._p
    start = paragraph.xpath("./w:bookmarkStart")[0]
    end = paragraph.xpath("./w:bookmarkEnd")[0]
    if fault == "missing":
        paragraph.remove(end)
    elif fault in {"duplicate_name", "duplicate_id"}:
        duplicate = copy.deepcopy(start)
        duplicate.set(qn("w:id"), "99" if fault == "duplicate_name" else start.get(qn("w:id")))
        duplicate.set(
            qn("w:name"), start.get(qn("w:name")) if fault == "duplicate_name" else "other"
        )
        finish = OxmlElement("w:bookmarkEnd")
        finish.set(qn("w:id"), duplicate.get(qn("w:id")))
        paragraph.append(duplicate)
        paragraph.append(finish)
    elif fault == "orphan":
        paragraph.remove(start)
    elif fault == "id":
        start.set(qn("w:id"), "invalid")
    elif fault == "order":
        paragraph.remove(end)
        paragraph.insert(0, end)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("INVALID_BOOKMARK" in result["errors"]) is (fault != "valid")


@pytest.mark.parametrize("color", ["000000", "FFFFFF"])
def test_page_background_cannot_hide_black_text(tmp_path: Path, color: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    background = OxmlElement("w:background")
    background.set(qn("w:color"), color)
    doc._element.insert(0, background)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_VISIBILITY_STYLE" in result["errors"]) is (color != "FFFFFF")
