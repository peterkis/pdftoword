"""Reviewed imports cannot launder incomplete seals or invisible text controls."""

import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate
from product_acceptance import digest, import_reviewed, write


def base_bundle(root: Path) -> Path:
    """Create a complete synthetic baseline with independently recorded identities."""
    path, truth, sources = specimen(root)
    path.rename(root / "auto.docx")
    (root / "input.pdf").write_bytes(b"synthetic source")
    truth["source_sha256"] = digest(root / "input.pdf")
    sources["docx_sha256"] = digest(root / "auto.docx")
    write(root / "truth.json", truth)
    write(root / "source-map.json", sources)
    write(
        root / "run.json",
        {
            "source_sha256": truth["source_sha256"],
            "annotation_sha256": digest(root / "truth.json"),
            "selected_pages": [1, 2],
            "sample_id": "synthetic",
            "document_family": "synthetic",
            "source_tree_sha256": "test",
            "source_commit": "test",
            "dirty": True,
            "profile": {},
            "rendering": {"status": "PENDING"},
        },
    )
    write(
        root / "seal.json",
        {
            name: digest(root / name)
            for name in ["auto.docx", "input.pdf", "truth.json", "source-map.json", "run.json"]
        },
    )
    return root


@pytest.mark.parametrize(
    "omitted", ["all", "auto.docx", "input.pdf", "truth.json", "source-map.json", "run.json"]
)
def test_reviewed_import_requires_all_base_members(tmp_path: Path, omitted: str) -> None:
    root = base_bundle(tmp_path / "base")
    seal = json.loads((root / "seal.json").read_text())
    seal = {} if omitted == "all" else {k: v for k, v in seal.items() if k != omitted}
    (root / "seal.json").write_text(json.dumps(seal))
    with pytest.raises(ValueError, match=r"^INCOMPLETE_EVIDENCE_SEAL$"):
        import_reviewed(root, root / "auto.docx", tmp_path / "reviewed")
    assert not (tmp_path / "reviewed").exists()


def test_reviewed_import_rejects_inconsistent_sealed_identity(tmp_path: Path) -> None:
    root = base_bundle(tmp_path / "base")
    sources = json.loads((root / "source-map.json").read_text())
    sources["docx_sha256"] = "0" * 64
    (root / "source-map.json").write_text(json.dumps(sources))
    seal = json.loads((root / "seal.json").read_text())
    seal["source-map.json"] = digest(root / "source-map.json")
    (root / "seal.json").write_text(json.dumps(seal))
    with pytest.raises(ValueError, match=r"^SOURCE_MAP_OUTPUT_MISMATCH$"):
        import_reviewed(root, root / "auto.docx", tmp_path / "reviewed")
    assert not (tmp_path / "reviewed").exists()


@pytest.mark.parametrize("tag", ["tab", "br", "cr"])
@pytest.mark.parametrize("scope", ["body", "table"])
def test_word_text_controls_participate_in_raw_comparison(
    tmp_path: Path, tag: str, scope: str
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    p = doc.paragraphs[0] if scope == "body" else doc.tables[0].cell(0, 0).paragraphs[0]
    p.runs[0].text = "A" if scope == "body" else "1"
    p.runs[0]._r.append(OxmlElement("w:" + tag))
    p.add_run("BC" if scope == "body" else "2")
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["content_status"] == "FAIL"
    if scope == "body":
        assert result["metrics"]["text"]["raw"]["insertions"] == 1
        assert result["metrics"]["text"]["light"]["distance"] == 0
    else:
        assert result["metrics"]["table_cells"]["correct_count"] == 1


def test_tab_stop_formatting_is_not_a_text_tab(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    doc.paragraphs[0].paragraph_format.tab_stops.add_tab_stop(Pt(72))
    doc.save(str(path))
    assert evaluate(path, truth, sources)["metrics"]["text"]["raw"]["distance"] == 0


@pytest.mark.parametrize("new_paragraph", [True, False])
def test_internal_linked_blip_is_explicitly_rejected(tmp_path: Path, new_paragraph: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    p = doc.add_paragraph() if new_paragraph else doc.paragraphs[-1]
    if new_paragraph:
        p.add_run().add_picture(str(tmp_path / "figure.png"))
    blip = p._p.xpath(".//a:blip")[0]
    blip.set(qn("r:link"), blip.get(qn("r:embed")))
    del blip.attrib[qn("r:embed")]
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_LINKED_IMAGE" in result["errors"]
    assert result["structure_status"] == "FAIL"


def test_table_paragraph_boundary_is_not_concatenated_away(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    cell = doc.tables[0].cell(0, 0)
    cell.paragraphs[0].runs[0].text = "1"
    cell.add_paragraph("2")
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["metrics"]["table_cells"]["correct_count"] == 1
    assert result["content_status"] == "FAIL"
