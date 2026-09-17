"""Section body regions and custom tab stops must stay on the page."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("case", ["top", "bottom", "height", "negative", "valid"])
def test_vertical_body_region(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    section = doc.sections[0]._sectPr
    margins = section.find(qn("w:pgMar"))
    if case in {"top", "bottom"}:
        margins.set(qn("w:" + case), "31680")
    elif case == "height":
        section.find(qn("w:pgSz")).set(qn("w:h"), "100")
    elif case == "negative":
        margins.set(qn("w:top"), "-100")
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is (case != "valid")
    assert result["structure_status"] == ("PASS" if case == "valid" else "FAIL")


@pytest.mark.parametrize("where", ["direct", "style", "default"])
@pytest.mark.parametrize("position", ["1440", "31680"])
def test_tab_stop_position(tmp_path: Path, where: str, position: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    doc.paragraphs[0].runs[0].text = "A\tB"
    next(u for u in truth["units"] if u["unit_id"] == "a")["reference"]["text"] = "A\tB"
    props = doc.paragraphs[0]._p.get_or_add_pPr()
    if where == "style":
        props = doc.styles["Normal"]._element.get_or_add_pPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:pPrDefault/w:pPr")[0]
    tabs, tab = OxmlElement("w:tabs"), OxmlElement("w:tab")
    tab.set(qn("w:val"), "left")
    tab.set(qn("w:pos"), position)
    tabs.append(tab)
    props.append(tabs)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is (position == "31680")
    assert result["structure_status"] == ("PASS" if position == "1440" else "FAIL")
