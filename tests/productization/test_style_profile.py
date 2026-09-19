"""Actual style plans, public heading calls and DOCX inheritance; no OCR mocks."""

from __future__ import annotations

import copy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from prototypes.docx_output.common import DemoError, new_job
from prototypes.docx_output.planning.flow import plan_flow
from prototypes.docx_output.planning.style_profile import (
    font_family,
    heading_candidates,
    weighted_sizes,
)
from prototypes.docx_output.renderers.flow import FlowRenderer
from prototypes.docx_output.structure_processors.docvortex import DocVortexStructureProcessor
from tests.productization.test_native_reflow import case, native_lines

__all__ = ["case"]


def test_cluster_is_native_character_weighted_and_ignores_header(case: Path) -> None:
    ir = native_lines(
        case,
        [
            ("x" * 80, [10, 10, 300, 20]),
            ("y" * 60, [10, 25, 300, 35]),
            ("TITLE", [10, 40, 300, 50]),
        ],
    )
    blocks = ir["pages"][0]["blocks"]
    blocks[0]["content"]["runs"][0]["font_size_pt"] = 10.01
    blocks[1]["content"]["runs"][0]["font_size_pt"] = 10.19
    blocks[2]["content"]["runs"][0]["font_size_pt"] = 48
    blocks[2]["type"] = "heading"
    c = weighted_sizes(ir, {"paragraph"})
    assert len(c) == 1 and c[0]["characters"] == 140 and 10 < c[0]["size_pt"] < 10.2
    blocks[0]["geometry_source"] = "pp_structure"
    assert weighted_sizes(ir, {"paragraph"})[0]["characters"] == 60


@pytest.mark.parametrize(
    "name,expected",
    [
        ("ABCDEF+Arial", "Arial"),
        ("abcdef+Arial", "Fallback"),
        (None, "Fallback"),
        ("C2_3", "Fallback"),
        ("SimSun", "Songti SC"),
    ],
)
def test_font_mapping_is_verified_not_resource_name(name: str | None, expected: str) -> None:
    assert font_family(name, "Fallback", {"Arial", "Songti SC"})[0] == expected


@pytest.mark.parametrize(
    "fault", ["none", "jump", "repeat", "partial", "question", "empty", "new_document"]
)
def test_heading_numbering_unknown_and_source_hashes(case: Path, fault: str) -> None:
    ir = native_lines(case, [("1 Overview", [10, 10, 300, 20]), ("1.1 Detail", [10, 30, 300, 40])])
    for b in ir["pages"][0]["blocks"]:
        b["type"] = "heading"
    second = ir["pages"][0]["blocks"][1]
    if fault == "jump":
        second["content"]["plain_text"] = "1.1.1 Too deep"
    if fault == "repeat":
        second["content"]["plain_text"] = "1 Again"
    if fault == "empty":
        second["content"]["plain_text"] = ""
    if fault == "question":
        second["type"] = "question"
    if fault == "partial":
        ir["pages"][0]["page_index"] = 3
    if fault == "new_document":
        ir["metadata"]["document_start_pages"] = [0]
    before = copy.deepcopy(ir)
    c = heading_candidates(ir)
    assert ir == before
    if fault in {"none", "new_document"}:
        assert [h["level"] for h in c] == [1, 2]
    elif fault == "question":
        assert len(c) == 1
    else:
        assert c[-1]["level"] is None


def test_actual_public_title_default_is_separate_from_numbered_candidate(case: Path) -> None:
    ir = native_lines(case, [("Document title", [10, 10, 300, 20]), ("1 Scope", [10, 30, 300, 40])])
    for b in ir["pages"][0]["blocks"]:
        b["type"] = "heading"
    ir["metadata"]["structure_evidence"] = {"atom-0": {"heading_role": "document_title"}}
    processor = DocVortexStructureProcessor()
    candidate = processor.process(ir)
    actual = processor.last_execution["middle"]["pages"][0]["blocks"]
    assert actual[0]["type"] == "doc_title" and actual[0]["level"] == 1
    profile = plan_flow(candidate.document, families={"Arial"}).document["metadata"][
        "style_profile"
    ]
    headings = profile["heading_candidates"]
    assert headings[0]["level"] is None
    assert headings[1]["level"] == 1 and headings[1]["upstream_default_level"] == 2


def test_global_style_inheritance_and_source_overrides(case: Path) -> None:
    ir = native_lines(
        case, [("Editable body", [10, 10, 300, 20]), ("Footnote2", [10, 30, 300, 40])]
    )
    r = ir["pages"][0]["blocks"][1]["content"]["runs"][0]
    r.update(superscript=True, underline=True, color="112233")
    original = copy.deepcopy(ir)
    plan = plan_flow(ir, {"editable_styles": True}, families={"Arial", "Arial Unicode MS"})
    job = new_job(case / "out")
    FlowRenderer().render(job, plan, "auto")
    doc = Document(str(job / "auto.docx"))
    runs = [r for p in doc.paragraphs for r in p.runs if r.text]
    assert runs[0]._r.rPr is not None
    assert runs[0].font.size is None and runs[0]._r.rPr.rFonts is None
    assert (
        runs[1].font.superscript and runs[1].underline and str(runs[1].font.color.rgb) == "112233"
    )
    rf = doc.styles["Normal"].element.rPr.rFonts
    assert all(rf.get(qn("w:" + k)) for k in ["ascii", "hAnsi", "eastAsia", "cs"])
    assert ir == original and plan.document["pages"] == ir["pages"]


def test_scan_unknown_canvas_and_mixed_sections(case: Path) -> None:
    ir = native_lines(case, [("scan", [10, 10, 200, 20])])
    ir["pages"][0]["blocks"][0]["geometry_source"] = "pp_structure"
    second = copy.deepcopy(ir["pages"][0])
    second.update(page_index=1, width_pt=800, height_pt=400, blocks=[], reading_order=[])
    ir["pages"].append(second)
    ir["source"]["page_count"] = 2
    ir["provenance"].setdefault("pages", {}).setdefault("0", {})["size_basis"] = (
        "virtual_width_595.28pt_not_DPI"
    )
    p = plan_flow(ir, families=set())
    assert len(p.output_layout["sections"]) == 2
    assert p.document["output_styles"]["body"]["basis"] == "inferred_scan_profile"
    assert "virtual" in p.document["metadata"]["style_profile"]["source_pages"][0]["size_basis"]
    assert (
        p.document["metadata"]["style_profile"]["font_inventory"]["other_operating_systems"]
        == "NOT_VERIFIED"
    )
    with pytest.raises(DemoError):
        plan_flow(ir, {"page_size_pt": [0, 0]})


def test_style_only_export_seal_and_content_preservation(case: Path) -> None:
    from prototypes.docx_output.common import digest, save
    from prototypes.docx_output.pipeline import finish
    from prototypes.docx_output.style_replay import export_style

    ir = native_lines(case, [("Body", [10, 10, 200, 20])])
    source = new_job(case / "source")
    import shutil

    shutil.copyfile(next(case.rglob("source-0.png")), source / "assets/source-0.png")
    finish(source, ir)
    before = digest(source / "layout.auto.json")
    output = export_style(
        source,
        case / "styled",
        {"body_size_pt": 14, "editable_styles": True, "page_size_pt": [800, 400]},
        before,
    )
    from prototypes.docx_output.common import read

    assert read(output / "layout.auto.json")["pages"] == ir["pages"]
    assert read(output / "style-replay.json")["model_calls"] == 0
    assert digest(source / "layout.auto.json") == before
    doc = Document(str(output / "auto.docx"))
    assert doc.styles["Normal"].font.size is not None
    assert doc.styles["Normal"].font.size.pt == 14
    width = doc.sections[0].page_width
    assert width is not None and width.pt == 800
    with pytest.raises(DemoError, match="SEAL_MISMATCH"):
        export_style(source, case / "bad", {}, "0" * 64)
    save(output / "test-receipt.json", {"source_unchanged": True})


def test_cjk_subset_uses_verified_east_asia_family(case: Path) -> None:
    ir = native_lines(case, [("中文", [10, 10, 200, 20])])
    ir["pages"][0]["blocks"][0]["content"]["runs"][0]["font_family"] = "ABCDEF+SimSun"
    plan = plan_flow(
        ir, {"editable_styles": True}, families={"Arial", "Arial Unicode MS", "SimSun"}
    )
    spec = plan.output_layout["sections"][0]["nodes"][0]["runs"]["atom-0"][0]
    assert spec["eastAsia"] == "SimSun" and not spec["inherit_fonts"]


def test_manual_run_lock_does_not_inherit_changed_global_size(case: Path) -> None:
    ir = native_lines(case, [("locked", [10, 10, 200, 20])])
    ir["pages"][0]["blocks"][0]["flags"].append("manual_lock")
    p = plan_flow(ir, {"editable_styles": True, "body_size_pt": 18}, families={"Arial"})
    spec = p.output_layout["sections"][0]["nodes"][0]["runs"]["atom-0"][0]
    assert spec["font_size_pt"] == 10 and not spec["inherit_base"]


def test_raster_without_dpi_is_not_claimed_as_measured_paper(case: Path) -> None:
    ir = native_lines(case, [("scan", [10, 10, 200, 20])])
    ir["source"]["filename"] = "scan.png"
    ir["provenance"]["pages"] = {}
    p = plan_flow(ir, families=set())
    assert (
        p.document["metadata"]["style_profile"]["source_pages"][0]["size_basis"]
        == "virtual_canvas_physical_size_unknown"
    )


def test_native_title_only_and_blank_body_are_unknown(case: Path) -> None:
    ir = native_lines(case, [("Only title", [10, 10, 200, 20])])
    ir["pages"][0]["blocks"][0]["type"] = "heading"
    p = plan_flow(ir, families=set())
    assert p.document["output_styles"]["body"]["basis"] == "unmeasured_native_body_default"
    ir["pages"][0].update(blocks=[], reading_order=[])
    p = plan_flow(ir, families=set())
    assert not p.document["metadata"]["style_profile"]["native_size_clusters"]
