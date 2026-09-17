"""Package validity and visible content remain prerequisites to acceptance."""

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "fault", ["content_types", "root_rels", "main_relation", "wrong_target", "wrong_type"]
)
def test_opc_required_parts_and_main_relationship(tmp_path: Path, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    if fault == "content_types":
        del files["[Content_Types].xml"]
    elif fault == "root_rels":
        del files["_rels/.rels"]
    elif fault in {"main_relation", "wrong_target"}:
        root = etree.fromstring(files["_rels/.rels"])
        relation = next(r for r in root if r.get("Type", "").endswith("/officeDocument"))
        if fault == "main_relation":
            root.remove(relation)
        else:
            relation.set("Target", "docProps/core.xml")
        files["_rels/.rels"] = etree.tostring(root)
    else:
        root = etree.fromstring(files["[Content_Types].xml"])
        next(n for n in root if n.get("PartName") == "/word/document.xml").set(
            "ContentType", "text/plain"
        )
        files["[Content_Types].xml"] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    assert result["structure_status"] == "FAIL"
    assert any(code.startswith("OPC_") for code in result["errors"])


@pytest.mark.parametrize(
    "mode,hidden",
    [
        ("direct", True),
        ("character_style", True),
        ("paragraph_style", True),
        ("override", False),
        ("unused", False),
    ],
)
def test_hidden_runs_and_inherited_styles_are_not_accepted(
    tmp_path: Path, mode: str, hidden: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    if mode == "direct":
        doc.paragraphs[0].runs[0].font.hidden = True
    else:
        kind = WD_STYLE_TYPE.PARAGRAPH if mode == "paragraph_style" else WD_STYLE_TYPE.CHARACTER
        base = doc.styles.add_style("HiddenBase", kind)
        base.font.hidden = True
        derived = doc.styles.add_style("DerivedHidden", kind)
        derived.base_style = base
        if mode == "paragraph_style":
            doc.paragraphs[0].style = derived
        elif mode != "unused":
            doc.paragraphs[0].runs[0].style = derived
            if mode == "override":
                doc.paragraphs[0].runs[0].font.hidden = False
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("HIDDEN_CONTENT" in result["errors"]) is hidden
    assert result["structure_status"] == ("FAIL" if hidden else "PASS")


@pytest.mark.parametrize("wrong", [False, True])
def test_layout_table_is_scored_as_text_not_data_topology(tmp_path: Path, wrong: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    next(u for u in truth["units"] if u["kind"] == "table")["reference"]["is_data_table"] = False
    if wrong:
        doc = Document(str(path))
        doc.tables[0].cell(0, 1).paragraphs[0].runs[0].text = "99"
        doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["metrics"]["tables"]["eligible_count"] == 0
    assert result["metrics"]["text"]["raw"]["reference_chars"] == 10
    assert result["content_status"] == ("FAIL" if wrong else "PASS")
    assert "UNALIGNED_EDITABLE_TEXT" not in result["errors"]


@pytest.mark.parametrize("story", ["header", "footer"])
@pytest.mark.parametrize("tag", ["sym", "noBreakHyphen"])
def test_nonbody_symbols_share_visible_node_policy(tmp_path: Path, story: str, tag: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    node = OxmlElement("w:" + tag)
    if tag == "sym":
        node.set(qn("w:font"), "Wingdings")
        node.set(qn("w:char"), "F041")
    getattr(doc.sections[0], story).paragraphs[0].add_run()._r.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_VISIBLE_STORY" in result["errors"]
    assert result["structure_status"] == "FAIL"


def test_hidden_style_is_loaded_from_actual_relationship_target(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    style = doc.styles.add_style("HiddenElsewhere", WD_STYLE_TYPE.CHARACTER)
    style.font.hidden = True
    doc.paragraphs[0].runs[0].style = style
    doc.save(str(path))
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    files["word/custom-styles.xml"] = files.pop("word/styles.xml")
    rels = etree.fromstring(files["word/_rels/document.xml.rels"])
    next(r for r in rels if r.get("Type", "").endswith("/styles")).set(
        "Target", "custom-styles.xml"
    )
    files["word/_rels/document.xml.rels"] = etree.tostring(rels)
    types = etree.fromstring(files["[Content_Types].xml"])
    next(t for t in types if t.get("PartName") == "/word/styles.xml").set(
        "PartName", "/word/custom-styles.xml"
    )
    files["[Content_Types].xml"] = etree.tostring(types)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    assert "HIDDEN_CONTENT" in evaluate(path, truth, sources)["errors"]
