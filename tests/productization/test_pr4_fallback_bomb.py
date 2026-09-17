"""Ordinary figures cannot replace missing content; image limits fail safely."""

import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("kind", ["text", "formula", "table"])
@pytest.mark.parametrize("fallback", [True, False])
def test_only_explicit_fallback_preserves_missing_content(
    tmp_path: Path, kind: str, fallback: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    bid = {"text": "a", "formula": "f", "table": "t"}[kind]
    block = next(b for b in sources["blocks"] if b["block_id"] == bid)
    paragraph = doc.paragraphs[0] if kind == "text" else doc.paragraphs[2]
    if kind == "table":
        paragraph = doc.tables[0].cell(0, 0).paragraphs[0]
        for row in doc.tables[0].rows:
            for cell in row.cells:
                for run in cell.paragraphs[0].runs:
                    run.text = ""
    elif kind == "text":
        paragraph.runs[0].text = ""
    else:
        paragraph._p.remove(paragraph._p.xpath(".//m:oMath")[0])
    paragraph.add_run().add_picture(str(tmp_path / "figure.png"), width=Inches(1))
    image = copy.deepcopy(sources["blocks"][-1]["images"][0])
    image.update(bbox=block["bbox"], fallback=fallback)
    block["images"] = [image]
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["content_status"] == ("REVIEW_REQUIRED" if fallback else "FAIL")
    assert result["metrics"]["necessary_content"]["source_image_retained_count"] == int(fallback)


@pytest.mark.parametrize("outside_media", [False, True])
def test_decompression_bomb_is_an_auditable_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, outside_media: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    if outside_media:
        with zipfile.ZipFile(path) as archive:
            files = {n: archive.read(n) for n in archive.namelist()}
        root = etree.fromstring(files["word/_rels/document.xml.rels"])
        rel = next(r for r in root if r.get("Type", "").endswith("/image"))
        target = rel.get("Target")
        assert target is not None
        old = "word/" + target
        files["assets/figure.png"] = files.pop(old)
        rel.set("Target", "../assets/figure.png")
        files["word/_rels/document.xml.rels"] = etree.tostring(root)
        with zipfile.ZipFile(path, "w") as archive:
            for n, data in files.items():
                archive.writestr(n, data)
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 50)
    result = evaluate(path, truth, sources)
    assert "IMAGE_PIXEL_LIMIT" in result["errors"]
    assert result["structure_status"] == "FAIL"
