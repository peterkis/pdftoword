"""Compatibility alternatives and foreign formula nodes must not become valid text."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


@pytest.mark.parametrize("scope", ["body", "header"])
@pytest.mark.parametrize("same", [False, True])
def test_alternate_content_is_rejected_instead_of_combining_branches(
    tmp_path: Path, scope: str, same: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0] if scope == "body" else doc.sections[0].header.paragraphs[0]
    for run in list(paragraph._p.findall(qn("w:r"))):
        paragraph._p.remove(run)
    alternate = etree.Element(f"{{{MC}}}AlternateContent", nsmap={"mc": MC, "w14": W14})
    for tag, value in [("Choice", "ABC" if same else "A"), ("Fallback", "ABC" if same else "BC")]:
        branch = etree.SubElement(alternate, f"{{{MC}}}{tag}")
        if tag == "Choice":
            branch.set("Requires", "w14")
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = value
        run.append(text)
        branch.append(run)
    paragraph._p.append(alternate)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_ALTERNATE_CONTENT" in result["errors"]
    assert result["structure_status"] == "FAIL"


def test_alternate_style_properties_are_not_combined(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    alternate = etree.Element(f"{{{MC}}}AlternateContent", nsmap={"mc": MC, "w14": W14})
    choice = etree.SubElement(alternate, f"{{{MC}}}Choice", Requires="w14")
    props = OxmlElement("w:rPr")
    props.append(OxmlElement("w:vanish"))
    choice.append(props)
    etree.SubElement(alternate, f"{{{MC}}}Fallback").append(OxmlElement("w:rPr"))
    doc.styles["Normal"]._element.append(alternate)
    doc.save(str(path))
    assert "UNSUPPORTED_ALTERNATE_CONTENT" in evaluate(path, truth, sources)["errors"]


@pytest.mark.parametrize(
    "case",
    [
        "word_run",
        "word_text",
        "unqualified",
        "foreign_style",
        "word_properties",
        "comments",
        "word_comments",
        "nested_math_text",
    ],
)
def test_math_nodes_keep_namespace_identity(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    math = doc.paragraphs[2]._p[-1]
    run, text = math[0], math[0][-1]
    if case == "word_run":
        run.tag = qn("w:r")
    elif case == "word_text":
        text.tag = qn("w:t")
    elif case == "unqualified":
        run.tag, text.tag = "r", "t"
    elif case == "foreign_style":
        props = OxmlElement("m:rPr")
        etree.SubElement(props, "{urn:foreign}sty", val="p")
        run.insert(0, props)
    elif case == "word_properties":
        props = OxmlElement("w:rPr")
        props.append(OxmlElement("w:b"))
        run.insert(0, props)
    elif case == "comments":
        run.insert(0, etree.Comment("ordinary XML comment"))
        text.text = "x"
        comment = etree.Comment("split text")
        comment.tail = "=1"
        text.append(comment)
    elif case == "word_comments":
        word_text = doc.paragraphs[0]._p.xpath(".//w:t")[0]
        word_text.text = "A"
        comment = etree.Comment("split Word text")
        comment.tail = "BC"
        word_text.append(comment)
    else:
        etree.SubElement(text, qn("m:r"))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    if case in {"word_properties", "comments", "word_comments"}:
        assert result["errors"] == []
        assert result["content_status"] == "PASS"
        assert result["structure_status"] == "PASS"
    else:
        assert (
            "INVALID_OMML_STRUCTURE" if case == "nested_math_text" else "INVALID_OMML_NAMESPACE"
        ) in result["errors"]
        assert result["structure_status"] == "FAIL"
