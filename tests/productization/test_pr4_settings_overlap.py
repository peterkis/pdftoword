"""Compatibility branches and overlapping truth must not distort acceptance."""

import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.docx_reader import NS
from acceptance.metrics import evaluate


@pytest.mark.parametrize("part", ["settings", "fontTable", "theme/theme1"])
def test_metadata_alternate_content_is_rejected(tmp_path: Path, part: str) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    name = "word/" + part + ".xml"
    root = etree.fromstring(files[name])
    alternate = etree.SubElement(root, "{" + NS["mc"] + "}AlternateContent")
    fallback = etree.SubElement(alternate, "{" + NS["mc"] + "}Fallback")
    if part == "settings":
        protection = OxmlElement("w:documentProtection")
        protection.set(qn("w:enforcement"), "true")
        fallback.append(protection)
    files[name] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for n, data in files.items():
            archive.writestr(n, data)
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_ALTERNATE_CONTENT" in result["errors"]
    assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize("kind", ["table", "formula"])
@pytest.mark.parametrize("broken", [True, False])
def test_structural_text_overlap_is_not_counted_twice(
    tmp_path: Path, kind: str, broken: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    if broken:
        doc: Any = Document(str(path))
        if kind == "formula":
            doc.paragraphs[2]._p.xpath(".//m:t")[0].text = "wrong"
        else:
            doc.tables[0].cell(0, 1).text = "wrong"
        doc.save(str(path))
    before = evaluate(path, truth, sources)
    unit = next(u for u in truth["units"] if u["kind"] == kind)
    truth["units"].append(
        {
            "unit_id": "overlap",
            "page": unit["page"],
            "kind": "text",
            "status": "confirmed",
            "reference": {
                "source_anchor_id": unit["reference"]["source_anchor_id"],
                "text": "x=1" if kind == "formula" else "1234",
            },
        }
    )
    result = evaluate(path, truth, sources)
    assert (
        result["metrics"]["necessary_content"]["eligible_count"]
        == before["metrics"]["necessary_content"]["eligible_count"]
    )
    assert result["content_status"] == ("FAIL" if broken else "PASS")
    assert result["metrics"]["necessary_content"]["excluded_structural_overlap_text_count"] == 1
