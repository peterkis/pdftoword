"""Frozen inputs, image fallback and unsupported body markup stay explicit."""

import copy
import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate
from product_acceptance import digest, run_sample


@pytest.mark.parametrize(
    "omitted", ["all", "manifest.json", "corpus/source.pdf", "annotations/ref.json"]
)
def test_dataset_requires_sealed_manifest_and_selected_members(
    tmp_path: Path, omitted: str
) -> None:
    dataset = tmp_path / "dataset"
    (dataset / "corpus").mkdir(parents=True)
    (dataset / "annotations").mkdir()
    (dataset / "corpus/source.pdf").write_bytes(b"invalid synthetic PDF never to be converted")
    (dataset / "annotations/ref.json").write_text("{}")
    sample = {
        "sample_id": "s",
        "document_family": "f",
        "local_read_authorized": True,
        "usage": "initial",
        "split": "development",
        "category": "N",
        "selected_pages": [1],
        "private_path": "source.pdf",
        "annotation_path": "ref.json",
        "source_sha256": digest(dataset / "corpus/source.pdf"),
    }
    (dataset / "manifest.json").write_text(json.dumps({"status": "FROZEN", "items": [sample]}))
    hashes = {
        name: digest(dataset / name)
        for name in ["manifest.json", "corpus/source.pdf", "annotations/ref.json"]
        if omitted != "all" and name != omitted
    }
    (dataset / "seal.json").write_text(json.dumps(hashes))
    output = tmp_path / "output"
    with pytest.raises(ValueError, match=r"^INCOMPLETE_DATASET_SEAL$"):
        run_sample(dataset, "s", "cli", output)
    assert not output.exists()


@pytest.mark.parametrize("case", ["missing", "mismatch", "other_missing"])
def test_verified_source_image_preserves_content_but_not_editability(
    tmp_path: Path, case: str
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    for run in doc.paragraphs[0].runs:
        run.text = "XYZ" if case == "mismatch" else ""
    doc.paragraphs[0].add_run().add_picture(str(tmp_path / "figure.png"))
    image = copy.deepcopy(sources["blocks"][-1]["images"][0])
    image.update(bbox=[0, 0, 100, 20], fallback=True)
    sources["blocks"][0]["images"] = [image]
    if case == "other_missing":
        for run in doc.paragraphs[1].runs:
            run.text = ""
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["content_status"] == ("FAIL" if case == "other_missing" else "REVIEW_REQUIRED")
    assert result["editability_status"] == "FAIL"
    assert result["metrics"]["necessary_content"]["visual_crop_acceptance"] == "PENDING"


@pytest.mark.parametrize(
    "tag",
    ["sym", "altChunk", "noBreakHyphen", "fldSimple", "pict", "object", "contentPart", "chart"],
)
def test_unsupported_visible_body_nodes_cannot_pass(tmp_path: Path, tag: str) -> None:
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    if tag == "altChunk":
        part = Part(PackURI("/word/extra.html"), "text/html", b"<html><p>extra</p></html>")
        rid = doc.part.relate_to(
            part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk"
        )
        element = OxmlElement("w:altChunk")
        element.set(qn("r:id"), rid)
        doc._element.body.insert(0, element)
    elif tag == "chart":
        doc.paragraphs[0].add_run()._r.append(OxmlElement("w:drawing"))
    else:
        element = OxmlElement("w:" + tag)
        if tag == "sym":
            element.set(qn("w:font"), "Wingdings")
            element.set(qn("w:char"), "F041")
        if tag == "fldSimple":
            element.set(qn("w:instr"), "PAGE")
        doc.paragraphs[0].add_run()._r.append(element)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_BODY_CONTENT" in result["errors"]
    assert result["structure_status"] == "FAIL"
