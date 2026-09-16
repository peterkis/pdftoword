"""Write protection, dynamic dates and partial relation truth remain explicit."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("password", [True, False])
def test_write_protection(tmp_path: Path, password: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    node = OxmlElement("w:writeProtection")
    if password:
        node.set(qn("w:hash"), "c3ludGhldGlj")
        node.set(qn("w:salt"), "c2FsdA==")
    doc.settings._element.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_DOCUMENT_PROTECTION" in result["errors"]


@pytest.mark.parametrize(
    "tag", ["dayLong", "dayShort", "monthLong", "monthShort", "yearLong", "yearShort"]
)
@pytest.mark.parametrize("where", ["body", "footer"])
def test_dynamic_date(tmp_path: Path, tag: str, where: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0] if where == "body" else doc.sections[0].footer.paragraphs[0]
    paragraph.add_run()._r.append(OxmlElement("w:" + tag))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    code = "UNSUPPORTED_BODY_CONTENT" if where == "body" else "UNSUPPORTED_VISIBLE_STORY"
    assert code in result["errors"]


@pytest.mark.parametrize("complete", [False, True])
@pytest.mark.parametrize("extra", ["unknown", "duplicate"])
def test_partial_relation_predictions(tmp_path: Path, complete: bool, extra: str) -> None:
    path, truth, sources = specimen(tmp_path)
    truth["coverage_complete"] = complete
    edge = dict(sources["relations"][0])
    if extra == "unknown":
        edge["to"] = "f"
    sources["relations"].append(edge)
    result = evaluate(path, truth, sources)
    precision = result["metrics"]["relation_precision"]
    unknown = not complete and extra == "unknown"
    assert precision["scored_count"] == (1 if unknown else 2)
    assert precision["not_scored_count"] == int(unknown)
    assert precision["correct_count"] == 1
    assert (result["structure_status"] == "FAIL") is (not unknown)
