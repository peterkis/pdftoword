"""Uncertain references own their output without granting unrelated content a pass."""

import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("kind", ["text", "formula", "table"])
@pytest.mark.parametrize("extra", [False, True])
def test_uncertain_content_is_owned_but_unrelated_extras_fail(
    tmp_path: Path, kind: str, extra: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    next(u for u in truth["units"] if u["kind"] == kind)["status"] = "uncertain"
    if extra:
        doc: Any = Document(str(path))
        if kind == "text":
            doc.add_paragraph("Unrelated text")
        elif kind == "formula":
            doc.add_paragraph()._p.append(copy.deepcopy(doc.paragraphs[2]._p[-1]))
        else:
            doc.add_table(rows=1, cols=1).style = "Table Grid"
        doc.save(str(path))
    result = evaluate(path, truth, sources)
    error = {
        "text": "UNALIGNED_EDITABLE_TEXT",
        "formula": "UNALIGNED_FORMULA",
        "table": "UNALIGNED_TABLE",
    }[kind]
    assert (error in result["errors"]) is extra
    if not extra:
        assert result["errors"] == []
        assert result["content_status"] == "REVIEW_REQUIRED"
        assert result["metrics"]["necessary_content"]["uncertain_count"] == 1
    else:
        assert result["structure_status"] == "FAIL"


def test_additional_empty_bordered_table_is_not_invisible_to_qa(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    doc.add_table(rows=2, cols=2).style = "Table Grid"
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNALIGNED_TABLE" in result["errors"]
    assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize(
    "fault", ["relationship_type", "non_image_type", "wrong_image_type", "missing_type"]
)
def test_blip_requires_image_relationship_and_correct_content_type(
    tmp_path: Path, fault: str
) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    if fault == "relationship_type":
        root = etree.fromstring(files["word/_rels/document.xml.rels"])
        next(r for r in root if r.get("Type", "").endswith("/image")).set("Type", "urn:not-image")
        files["word/_rels/document.xml.rels"] = etree.tostring(root)
    else:
        root = etree.fromstring(files["[Content_Types].xml"])
        entry = next(c for c in root if c.get("Extension") == "png")
        if fault == "missing_type":
            root.remove(entry)
        else:
            entry.set("ContentType", "text/plain" if fault == "non_image_type" else "image/jpeg")
        files["[Content_Types].xml"] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    assert any(
        code in result["errors"]
        for code in ["INVALID_IMAGE_RELATIONSHIP", "IMAGE_CONTENT_TYPE_INVALID"]
    )
    assert result["structure_status"] == "FAIL"
