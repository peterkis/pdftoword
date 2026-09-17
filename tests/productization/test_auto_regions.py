"""Real mixed PDF/DOCX tests with only HTTP replaced at the transport boundary."""

from __future__ import annotations

import base64
import io
import json
import shutil
import uuid
from collections.abc import Iterator
from pathlib import Path

import httpx
import pytest
from docx import Document
from PIL import Image, ImageDraw
from prototypes.docx_output.common import PRIVATE, DemoError, digest, read
from prototypes.docx_output.pipeline import convert
from prototypes.docx_output.route_plan import execute_routes
from tests.productization.synthetic.generate import pdf_bytes


@pytest.fixture
def case() -> Iterator[Path]:
    root = PRIVATE / ("region-test-" + uuid.uuid4().hex)
    root.mkdir(parents=True)
    try:
        yield root
    finally:
        shutil.rmtree(root)


def mixed_pdf(source: Path) -> None:
    """Native header, separate raster text and separate vector illustration."""
    image = Image.new("RGB", (360, 120), "white")
    ImageDraw.Draw(image).text((12, 40), "Scanned content 12 < 34", fill="black")
    encoded = image.tobytes().hex().encode()
    stream = b"BT /F1 14 Tf 30 340 Td (Native content stays editable) Tj ET\n"
    stream += (
        b"q 180 0 0 60 30 220 cm BI /W 360 /H 120 /CS /RGB /BPC 8 /F /AHx ID "
        + encoded
        + b"> EI Q\n"
    )
    stream += b"1 0 0 RG 2 w 30 100 80 60 re S\n"
    source.write_bytes(pdf_bytes(stream))


def transport(
    monkeypatch: pytest.MonkeyPatch, *, purpose: str = "text", fail: bool = False
) -> list[str]:
    calls: list[str] = []
    original = httpx.Client

    def handle(request: httpx.Request) -> httpx.Response:
        provider = "pp" if request.url.port == 8080 else "ovis"
        calls.append(provider)
        if fail:
            return httpx.Response(503)
        data = json.loads(request.content)
        if provider == "ovis":
            return httpx.Response(
                200,
                json={
                    "model": "ovis-ocr2",
                    "choices": [
                        {"finish_reason": "stop", "message": {"content": "Scanned content 12 < 34"}}
                    ],
                },
            )
        with Image.open(io.BytesIO(base64.b64decode(data["file"]))) as im:
            width, height = im.size
        bbox = [10, 10, width - 10, height - 10]
        return httpx.Response(
            200,
            json={
                "errorCode": 0,
                "errorMsg": "Success",
                "logId": "test",
                "result": {
                    "dataInfo": {"width": width, "height": height},
                    "layoutParsingResults": [
                        {
                            "markdown": {"text": "", "isStart": True, "isEnd": True},
                            "prunedResult": {
                                "width": width,
                                "height": height,
                                "parsing_res_list": [
                                    {
                                        "block_label": purpose,
                                        "block_bbox": bbox,
                                        "block_content": "POISON PP TEXT",
                                    }
                                ],
                                "overall_ocr_res": {"rec_boxes": [bbox]},
                                "formula_res_list": [],
                            },
                        }
                    ],
                },
            },
        )

    def client(*args: object, **kwargs: object) -> httpx.Client:
        return original(transport=httpx.MockTransport(handle))

    monkeypatch.setattr(httpx, "Client", client)
    return calls


def prepared(case: Path) -> tuple[Path, dict]:
    source = case / "mixed.pdf"
    mixed_pdf(source)
    job = convert(source, mode="auto", output_root=case / "jobs")
    return job, read(job / "route-plan.json")


def test_mixed_preparation_zero_http_and_preserves_native(case: Path) -> None:
    job, plan = prepared(case)
    assert plan["request_budget"] == 2
    text = "\n".join(p.text for p in Document(str(job / "auto.docx")).paragraphs)
    assert "Native content stays editable" in text
    assert read(job / "qa.json")["model_call_count"] == 0
    assert len(plan["pages"][0]["regions"]) == 2


def test_auto_preparation_keeps_native_review_issues(case: Path) -> None:
    """Route preparation must not discard findings raised by native extraction."""
    source = case / "formula.pdf"
    source.write_bytes(pdf_bytes(b"BT /F1 14 Tf 30 340 Td (x$^2$) Tj ET"))
    job = convert(source, mode="auto", output_root=case / "jobs")
    issues = read(job / "layout.auto.json")["issues"]
    assert any(issue["type"] == "NATIVE_FORMULA_DELIMITER_REVIEW" for issue in issues)


def test_region_execution_produces_new_editable_docx(
    case: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    job, plan = prepared(case)
    before = digest(job / "auto.docx")
    calls = transport(monkeypatch)
    result = execute_routes(job, plan["plan_hash"], 2, True)
    assert calls == ["pp", "ovis"]
    text = "\n".join(p.text for p in Document(str(result / "auto.docx")).paragraphs)
    assert text.count("Native content stays editable") == 1
    assert text.count("Scanned content 12 < 34") == 1
    assert "POISON" not in text
    assert result != job and digest(job / "auto.docx") == before
    assert read(result / "qa.json")["placed_figure_count"] == 1
    with pytest.raises(DemoError, match="ALREADY_ATTEMPTED"):
        execute_routes(job, plan["plan_hash"], 2, True)
    assert calls == ["pp", "ovis"]


def test_ambiguous_staged_inputs_are_rejected_before_model_calls(
    case: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    job, plan = prepared(case)
    (job / "input.extra").write_bytes(b"unexpected")
    calls = transport(monkeypatch)
    with pytest.raises(DemoError, match="ROUTE_INPUT_AMBIGUOUS"):
        execute_routes(job, plan["plan_hash"], 2, True)
    assert calls == []


@pytest.mark.parametrize("fault", ["hash", "budget", "source", "crop", "expiry"])
def test_changed_plan_sends_nothing(
    case: Path, monkeypatch: pytest.MonkeyPatch, fault: str
) -> None:
    from prototypes.docx_output.common import save
    from prototypes.docx_output.route_plan import semantic_hash

    job, plan = prepared(case)
    calls = transport(monkeypatch)
    approved_hash, budget = plan["plan_hash"], 2
    if fault == "hash":
        approved_hash = "wrong"
    elif fault == "budget":
        budget = 3
    elif fault == "source":
        (job / "input.pdf").write_bytes(b"changed")
    elif fault == "crop":
        ir = read(job / "layout.prepared.json")
        aid = plan["pages"][0]["regions"][0]["asset_id"]
        path = next(a["path"] for a in ir["assets"] if a["id"] == aid)
        (job / path).write_bytes(b"changed")
    else:
        plan["expires_at"] = 0
        plan["plan_hash"] = semantic_hash({k: v for k, v in plan.items() if k != "plan_hash"})
        approved_hash = plan["plan_hash"]
        save(job / "route-plan.json", plan)
    with pytest.raises(DemoError):
        execute_routes(job, approved_hash, budget, True)
    assert calls == []


@pytest.mark.parametrize(
    "purpose,fail,expected",
    [
        ("image", False, "PRESERVED_FIGURE"),
        ("unknown", False, "REGION_PURPOSE_REVIEW_REQUIRED"),
        ("text", True, "PROVIDER_HTTP_ERROR"),
    ],
)
def test_region_failures_never_remove_native(
    case: Path, monkeypatch: pytest.MonkeyPatch, purpose: str, fail: bool, expected: str
) -> None:
    job, plan = prepared(case)
    calls = transport(monkeypatch, purpose=purpose, fail=fail)
    result = execute_routes(job, plan["plan_hash"], 2, True)
    assert calls == ["pp"]
    assert read(result / "region-results.json")["pages"][0]["regions"][0]["status"] == expected
    assert "Native content stays editable" in "\n".join(
        p.text for p in Document(str(result / "auto.docx")).paragraphs
    )


def test_cancel_before_execution_sends_nothing(case: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from prototypes.docx_output.common import save

    job, plan = prepared(case)
    calls = transport(monkeypatch)
    save(job / "route-cancelled.json", {"cancelled": True})
    child = execute_routes(job, plan["plan_hash"], 2, True)
    assert calls == []
    assert read(child / "region-results.json")["pages"][0]["regions"][0]["status"] == "CANCELLED"
    assert read(child / "qa.json")["execution_status"] == "PARTIAL"


def test_prepared_layout_is_bound_to_approval(case: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from prototypes.docx_output.common import save

    job, plan = prepared(case)
    calls = transport(monkeypatch)
    ir = read(job / "layout.prepared.json")
    ir["pages"][0]["blocks"][0]["content"]["plain_text"] = "changed after approval"
    save(job / "layout.prepared.json", ir)
    with pytest.raises(DemoError, match="PREPARED_CHANGED"):
        execute_routes(job, plan["plan_hash"], 2, True)
    assert calls == []


def test_api_uses_same_prepared_plan(case: Path) -> None:
    from fastapi.testclient import TestClient
    from prototypes.docx_output.server import create_app

    job, plan = prepared(case)
    client = TestClient(create_app(output_root=case / "jobs"), base_url="http://127.0.0.1:8765")
    token = client.get("/api/session").json()["token"]
    response = client.get("/api/route/" + job.name)
    assert response.status_code == 200
    assert response.json()["plan_hash"] == plan["plan_hash"]
    headers = {"Origin": "http://127.0.0.1:8765", "X-Demo-Session": token}
    invalid = client.post(
        "/api/route/" + job.name + "/execute",
        headers=headers,
        json={"plan_hash": "wrong", "budget": 2, "confirm_no_auth": True},
    )
    assert invalid.status_code == 400
    assert not (job / "route-execution.json").exists()
    malformed = client.post(
        "/api/route/" + job.name + "/execute",
        headers=headers,
        json=[],
    )
    assert malformed.status_code == 400
    assert malformed.json()["detail"] == "ROUTE_REQUEST_OBJECT_REQUIRED"


def test_region_table_keeps_surrounding_prose_editable(case: Path) -> None:
    from prototypes.docx_output.common import read
    from prototypes.docx_output.mixed_fusion import reconstruct_region

    job, plan = prepared(case)
    region = plan["pages"][0]["regions"][0]
    width, height = region["pixel_size"]
    ir = read(job / "layout.prepared.json")
    responses = {
        "ovis": {
            "choices": [
                {
                    "finish_reason": "stop",
                    "message": {
                        "content": (
                            "Before table\n\n<table><tr><td>Cell</td></tr></table>\n\nAfter table"
                        )
                    },
                }
            ]
        },
        "pp": {
            "result": {
                "layoutParsingResults": [
                    {
                        "prunedResult": {
                            "width": width,
                            "height": height,
                            "parsing_res_list": [
                                {"block_label": "text", "block_bbox": [10, 0, width - 10, 20]},
                                {
                                    "block_label": "table",
                                    "block_bbox": [10, 30, width - 10, height - 30],
                                },
                                {
                                    "block_label": "text",
                                    "block_bbox": [10, height - 25, width - 10, height - 5],
                                },
                            ],
                            "overall_ocr_res": {"rec_boxes": []},
                            "formula_res_list": [],
                        }
                    }
                ]
            }
        },
    }
    reconstruct_region(job, ir, region, responses, {"requests": []})
    blocks = ir["pages"][0]["blocks"]
    text = [b["content"].get("plain_text") for b in blocks]
    assert "Before table" in text and "After table" in text
    table = next(b for b in blocks if "REGION_TABLE_FALLBACK" in b["flags"])
    assert blocks.index(table) > text.index("Before table")
    assert blocks.index(table) < text.index("After table")
    assert table["bbox"][3] - table["bbox"][1] < region["bbox"][3] - region["bbox"][1]


def test_explicit_rebuild_reuses_successes_without_http(
    case: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    job, plan = prepared(case)
    calls = transport(monkeypatch)
    first = execute_routes(job, plan["plan_hash"], 2, True)
    second_base = convert(job / "input.pdf", mode="auto", output_root=case / "jobs")
    second_plan = read(second_base / "route-plan.json")
    second = execute_routes(second_base, second_plan["plan_hash"], 2, True, reuse_from=first)
    assert calls == ["pp", "ovis"]
    assert read(second / "request-manifest.json")["model_call_count"] == 0
    assert len(read(second / "request-manifest.json")["requests"]) == 2
    assert "Scanned content 12 < 34" in "\n".join(
        p.text for p in Document(str(second / "auto.docx")).paragraphs
    )


def test_upload_auto_is_offline_and_builds_real_docx(case: Path) -> None:
    import time

    from fastapi.testclient import TestClient
    from prototypes.docx_output.server import create_app

    source = case / "upload.pdf"
    mixed_pdf(source)
    with TestClient(
        create_app(output_root=case / "jobs"), base_url="http://127.0.0.1:8765"
    ) as client:
        token = client.get("/api/session").json()["token"]
        response = client.post(
            "/api/upload",
            headers={"Origin": "http://127.0.0.1:8765", "X-Demo-Session": token},
            files={"file": ("source.pdf", source.read_bytes(), "application/pdf")},
            data={"mode": "auto", "pages": "1"},
        )
        assert response.status_code == 200
        deadline = time.monotonic() + 5
        status = client.get("/api/status").json()
        while status["busy"] and time.monotonic() < deadline:
            time.sleep(0.01)
            status = client.get("/api/status").json()
        assert not status["busy"] and status["state"] != "失败"
        job = case / "jobs" / status["job_id"]
        assert read(job / "qa.json")["model_call_count"] == 0
        assert (job / "auto.docx").exists()
        assert client.get("/api/route/" + job.name).json()["request_budget"] == 2


def test_region_labels_mixed_is_not_ordinary_figure() -> None:
    from prototypes.docx_output.region_bridge import layout_purpose

    assert (
        layout_purpose(
            {
                "result": {
                    "layoutParsingResults": [
                        {
                            "prunedResult": {
                                "width": 100,
                                "height": 100,
                                "parsing_res_list": [
                                    {"block_label": "image"},
                                    {"block_label": "text"},
                                ],
                            }
                        }
                    ]
                }
            },
            [100, 100],
        )
        == "mixed"
    )


def test_human_review_cannot_be_discarded_by_old_plan(
    case: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from prototypes.docx_output.common import save

    job, plan = prepared(case)
    save(job / "layout.reviewed.json", {"human_selection": "must be retained"})
    calls = transport(monkeypatch)
    with pytest.raises(DemoError, match="REVIEWED_SOURCE_REQUIRES_NEW_PLAN"):
        execute_routes(job, plan["plan_hash"], 2, True)
    assert calls == []
    assert read(job / "layout.reviewed.json") == {"human_selection": "must be retained"}


def test_fill_in_blanks_remain_literal_editable_text(case: Path) -> None:
    from prototypes.docx_output.ovis_replay import recover_ovis
    from prototypes.docx_output.pipeline import finish
    from tests.demo.test_output import setup_ir

    job, ir, p = setup_ir(case)
    text = "(1) I like ___ flowers.\n\n(2) ___ stars are bright."
    recover_ovis(
        job,
        ir,
        p,
        {"choices": [{"finish_reason": "stop", "message": {"content": text}}]},
        "synthetic-blank",
    )
    finish(job, ir)
    actual = "\n".join(p.text for p in Document(str(job / "auto.docx")).paragraphs)
    assert "(1) I like ___ flowers." in actual
    assert "(2) ___ stars are bright." in actual
    assert not any(b["content"]["kind"] == "image" for b in p["blocks"])
