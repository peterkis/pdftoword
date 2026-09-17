"""Character spacing and displayed case cannot silently diverge from strict text."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("where", ["direct", "math", "style", "default"])
@pytest.mark.parametrize(
    "tag,value,invalid",
    [
        ("spacing", "-31680", True),
        ("spacing", "31680", True),
        ("spacing", "0", False),
        ("caps", "true", True),
        ("caps", "false", False),
        ("smallCaps", "true", True),
        ("smallCaps", "false", False),
    ],
)
def test_spacing_and_case_effects(
    tmp_path: Path, where: str, tag: str, value: str, invalid: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if where == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif where == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    node = OxmlElement("w:" + tag)
    node.set(qn("w:val"), value)
    props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    code = "UNSUPPORTED_TEXT_POSITION" if tag == "spacing" else "UNSUPPORTED_TEXT_CASE"
    assert (code in result["errors"]) is invalid
    assert result["structure_status"] == ("FAIL" if invalid else "PASS")
