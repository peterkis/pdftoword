"""Nonbody text controls are content, while paragraph tab stops are formatting."""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("story", ["header", "footer"])
@pytest.mark.parametrize("tag", ["tab", "br", "cr", "formatting_only"])
def test_story_controls_are_not_discarded(tmp_path: Path, story: str, tag: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    paragraph = getattr(doc.sections[0], story).paragraphs[0]
    if tag == "formatting_only":
        paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(72))
    else:
        paragraph.add_run()._r.append(OxmlElement("w:" + tag))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_VISIBLE_STORY" in result["errors"]) is (tag != "formatting_only")
    assert result["structure_status"] == ("PASS" if tag == "formatting_only" else "FAIL")
