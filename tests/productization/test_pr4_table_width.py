"""Fixed table widths must not push necessary cells outside the page."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("case", ["grid", "table", "cell", "percent", "row_sum", "valid"])
def test_fixed_table_width_bounds(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    table = doc.tables[0]
    table.autofit = False
    if case == "grid":
        table._tbl.tblGrid[0].set(qn("w:w"), "31680")
    elif case in {"table", "percent"}:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        width.set(qn("w:type"), "pct" if case == "percent" else "dxa")
        width.set(qn("w:w"), "10000" if case == "percent" else "31680")
    elif case == "cell":
        table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"), "31680")
    elif case == "row_sum":
        for cell in table.rows[0].cells:
            cell._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"), "6000")
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TABLE_WIDTH" in result["errors"]) is (case != "valid")
    assert result["structure_status"] == ("PASS" if case == "valid" else "FAIL")
