"""Unscored cross-page table semantics must be visible in the aggregate status."""
from pathlib import Path

import pytest
from docx import Document
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("segments", [1, 2])
@pytest.mark.parametrize("broken", [False, True])
def test_continuation_pending_propagates(tmp_path: Path, segments: int, broken: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    reference = next(u for u in truth["units"] if u["kind"] == "table")["reference"]
    reference.update(logical_table_id="logical-table", segment_count=segments, segment_index=0)
    if broken:
        doc = Document(str(path))
        doc.tables[0].cell(0, 1).text = "wrong"
        doc.save(str(path))
    result = evaluate(path, truth, sources)
    expected = "FAIL" if broken else "REVIEW_REQUIRED" if segments > 1 else "PASS"
    assert result["structure_status"] == expected
    assert result["metrics"]["tables"]["continuation"]["not_scored_count"] == int(segments > 1)
