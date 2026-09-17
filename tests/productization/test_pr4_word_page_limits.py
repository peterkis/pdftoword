"""Declared page regions must fit Word's documented page size limits."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("axis", ["w", "h"])
@pytest.mark.parametrize("size", ["31680", "31681"])
def test_word_page_size_boundary(tmp_path: Path, axis: str, size: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    doc.sections[0]._sectPr.find(qn("w:pgSz")).set(qn("w:" + axis), size)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_PAGE_SIZE" in result["errors"]) is (size == "31681")
    assert result["structure_status"] == ("PASS" if size == "31680" else "FAIL")


def test_fake_giant_page_cannot_authorize_giant_fallback(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    page = doc.sections[0]._sectPr.find(qn("w:pgSz"))
    page.set(qn("w:w"), "100000")
    page.set(qn("w:h"), "100000")
    for extent in doc.paragraphs[-1]._p.xpath(".//wp:extent | .//a:xfrm/a:ext"):
        extent.set("cx", "40000000")
        extent.set("cy", "40000000")
    sources["blocks"][-1]["images"][0]["fallback"] = True
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_PAGE_SIZE" in result["errors"]
    assert result["structure_status"] == "FAIL"
