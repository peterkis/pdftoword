"""A matching image blob cannot substitute for a renderable DrawingML picture."""

import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "fault",
    [
        "bare",
        "no_extent",
        "no_docpr",
        "wrong_uri",
        "no_transform",
        "zero_transform",
        "orphan",
        "cropped",
        "transparent",
        "anchor",
        "hidden",
        "local_dpi",
    ],
)
def test_picture_requires_supported_complete_container(tmp_path: Path, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[-1]
    drawing = paragraph._p.xpath(".//w:drawing")[0]
    inline = drawing[0]
    blip = drawing.xpath(".//a:blip")[0]
    if fault == "bare":
        saved = copy.deepcopy(blip)
        drawing.clear()
        drawing.append(saved)
    elif fault == "no_extent":
        inline.remove(inline.find(qn("wp:extent")))
    elif fault == "no_docpr":
        inline.remove(inline.find(qn("wp:docPr")))
    elif fault == "wrong_uri":
        drawing.xpath(".//a:graphicData")[0].set("uri", "urn:not-picture")
    elif fault == "no_transform":
        transform = drawing.xpath(".//a:xfrm")[0]
        transform.getparent().remove(transform)
    elif fault == "zero_transform":
        drawing.xpath(".//a:xfrm/a:ext")[0].set("cx", "0")
    elif fault == "orphan":
        saved = copy.deepcopy(blip)
        parent = drawing.getparent()
        parent.remove(drawing)
        parent.append(saved)
    elif fault == "cropped":
        crop = OxmlElement("a:srcRect")
        crop.set("l", "50000")
        blip.getparent().append(crop)
    elif fault == "transparent":
        alpha = OxmlElement("a:alphaModFix")
        alpha.set("amt", "0")
        blip.append(alpha)
    elif fault == "anchor":
        inline.tag = qn("wp:anchor")
    elif fault == "hidden":
        inline.find(qn("wp:docPr")).set("hidden", "1")
    else:
        extensions, extension = OxmlElement("a:extLst"), OxmlElement("a:ext")
        extension.set("uri", "{28A0092B-C50C-407E-A947-70E740481C1C}")
        dpi = etree.SubElement(
            extension, "{http://schemas.microsoft.com/office/drawing/2010/main}useLocalDpi"
        )
        dpi.set("val", "0")
        extensions.append(extension)
        blip.append(extensions)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    if fault == "local_dpi":
        assert result["structure_status"] == "PASS"
        assert result["errors"] == []
    else:
        assert any(
            c in result["errors"]
            for c in ["INVALID_DRAWING_CONTAINER", "UNSUPPORTED_IMAGE_RENDERING", "HIDDEN_CONTENT"]
        )
        assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize(
    "case",
    [
        "missing",
        "empty",
        "relative",
        "whitespace",
        "bad_percent",
        "urn",
        "missing_target",
        "bad_mode",
    ],
)
def test_relationship_required_attributes_are_validated(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    root = etree.fromstring(files["word/_rels/document.xml.rels"])
    relation = etree.SubElement(
        root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
    )
    relation.set("Id", "rExtra")
    relation.set("Target", "styles.xml")
    if case != "missing":
        relation.set(
            "Type",
            {
                "empty": "",
                "relative": "not-absolute",
                "whitespace": "urn:bad type",
                "bad_percent": "urn:bad%xy",
            }.get(case, "urn:custom:type"),
        )
    if case == "missing_target":
        del relation.attrib["Target"]
    if case == "bad_mode":
        relation.set("TargetMode", "Unknown")
    files["word/_rels/document.xml.rels"] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    if case == "urn":
        assert result["structure_status"] == "PASS"
    else:
        code = (
            "OPC_RELATIONSHIP_TARGET_INVALID"
            if case == "missing_target"
            else "OPC_RELATIONSHIP_MODE_INVALID"
            if case == "bad_mode"
            else "OPC_RELATIONSHIP_TYPE_INVALID"
        )
        assert code in result["errors"]
        assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize("hidden", [None, "0", "false", "1", "true"])
def test_picture_nonvisual_properties_control_visibility(
    tmp_path: Path, hidden: str | None
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    properties = doc.paragraphs[-1]._p.xpath(".//pic:cNvPr")[0]
    if hidden is not None:
        properties.set("hidden", hidden)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    invisible = hidden in {"1", "true"}
    assert ("HIDDEN_CONTENT" in result["errors"]) is invisible
    assert result["structure_status"] == ("FAIL" if invisible else "PASS")
