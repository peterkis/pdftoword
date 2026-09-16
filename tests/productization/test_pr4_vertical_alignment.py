"""Unannotated run-level superscript/subscript semantics cannot pass text matching."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("where", ["direct", "math", "style", "default"])
@pytest.mark.parametrize("alignment", ["superscript", "subscript", "baseline"])
def test_run_vertical_alignment(tmp_path: Path, where: str, alignment: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if where == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif where == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    node = OxmlElement("w:vertAlign")
    node.set(qn("w:val"), alignment)
    props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_VERTICAL_ALIGNMENT" in result["errors"]) is (alignment != "baseline")
    assert result["structure_status"] == ("PASS" if alignment == "baseline" else "FAIL")
