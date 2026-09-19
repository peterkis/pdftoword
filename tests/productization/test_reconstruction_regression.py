"""Actual package regression checks reject missing payload and silent replacements."""

from __future__ import annotations

import shutil
import uuid
from collections.abc import Iterator
from pathlib import Path

import pytest
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from prototypes.docx_output.acceptance import audit_pair
from prototypes.docx_output.common import PRIVATE, DemoError, Json, new_job
from prototypes.docx_output.pipeline import finish
from prototypes.docx_output.planning.flow import plan_flow
from tests.productization.test_shared_structure_actual import text_ir


@pytest.fixture
def case() -> Iterator[Path]:
    root = PRIVATE / ("acceptance-regression-" + uuid.uuid4().hex)
    root.mkdir(parents=True)
    try:
        yield root
    finally:
        shutil.rmtree(root)


def pair(root: Path) -> tuple[Path, Path]:
    ir: Json = text_ir(root, "Source inventory 中文")
    before, after = new_job(root / "before"), new_job(root / "after")
    finish(before, ir)
    plan = plan_flow(ir)
    finish(after, plan.document, render_plan=plan)
    return before, after


def test_actual_before_after_inventory_is_not_visual_acceptance(case: Path) -> None:
    before, after = pair(case)
    report = audit_pair(before, after)
    assert report["content_inventory_unchanged"] and report["source_ranges_nonempty"]
    assert report["word_acceptance"] == "NOT_EVALUATED"
    assert report["visual_acceptance"] == "NOT_EVALUATED"


@pytest.mark.parametrize("fault", ["empty_range", "changed_text", "missing_marker"])
def test_package_tampering_is_rejected(case: Path, fault: str) -> None:
    before, after = pair(case)
    doc = Document(str(after / "auto.docx"))
    body = doc.element.body
    start = body.xpath(".//w:bookmarkStart")[0]
    if fault == "empty_range":
        end = body.xpath(".//w:bookmarkEnd")[0]
        parent = start.getparent()
        parent.remove(end)
        parent.insert(list(parent).index(start) + 1, end)
    elif fault == "missing_marker":
        start.set(qn("w:name"), "wrong")
    else:
        body.xpath(".//w:t")[0].text = "Silent replacement"
    doc.save(str(after / "auto.docx"))
    with pytest.raises(
        DemoError,
        match=r"SOURCE_RANGE_EMPTY|SOURCE_MARKER_NOT_UNIQUE|CONTENT_INVENTORY_MISMATCH|SOURCE_RANGE_TEXT_MISMATCH",
    ):
        audit_pair(before, after)


def test_missing_docx_is_not_reported_as_accepted(case: Path) -> None:
    before, after = pair(case)
    shutil.move(after / "auto.docx", after / "withheld.docx")
    with pytest.raises(PackageNotFoundError):
        audit_pair(before, after)


@pytest.mark.parametrize("confirmed", [False, True])
def test_only_confirmed_caption_ranges_keep_figure_with_caption(
    case: Path, confirmed: bool
) -> None:
    from PIL import Image
    from prototypes.docx_output.common import block, digest, relation
    from prototypes.docx_output.structure import image_content

    ir = text_ir(case)
    job = new_job(case / "caption")
    path = job / "assets/figure.png"
    Image.new("RGB", (100, 100), "white").save(path)
    figure = block("figure", 0, [10, 10, 110, 110], "", "native_pdf", "figure")
    figure["content"] = image_content("figure-asset")
    caption1 = block("caption1", 0, [10, 115, 250, 125], "Figure 1: first line", "native_pdf")
    caption2 = block("caption2", 0, [10, 130, 250, 140], "second line", "native_pdf")
    ir["pages"][0]["blocks"] = [figure, caption1, caption2]
    ir["pages"][0]["reading_order"] = [b["id"] for b in ir["pages"][0]["blocks"]]
    ir["assets"] = [
        {
            "id": "figure-asset",
            "type": "image",
            "path": "assets/figure.png",
            "sha256": digest(path),
            "mime_type": "image/png",
            "source_page_index": 0,
            "source_bbox": figure["bbox"],
            "extraction_method": "synthetic",
            "pixel_width": 100,
            "pixel_height": 100,
        }
    ]
    for b in (caption1, caption2):
        relation(ir, "caption_of", b["id"], figure["id"], {"manual": confirmed})
    plan = plan_flow(ir)
    finish(job, plan.document, render_plan=plan)
    paragraphs = Document(str(job / "auto.docx")).paragraphs
    assert bool(paragraphs[0].paragraph_format.keep_with_next) is confirmed
    assert bool(paragraphs[1].paragraph_format.keep_with_next) is confirmed
    assert not paragraphs[2].paragraph_format.keep_with_next
    assert plan.document["pages"] == ir["pages"]


def test_swapped_nonempty_bookmarks_do_not_prove_correct_binding(case: Path) -> None:
    from prototypes.docx_output.acceptance import package_inventory
    from prototypes.docx_output.common import block

    ir = text_ir(case, "First source")
    second = block("second", 0, [10, 70, 200, 90], "Second source", "native_pdf")
    ir["pages"][0]["blocks"].append(second)
    ir["pages"][0]["reading_order"].append("second")
    job = new_job(case / "swapped")
    finish(job, ir)
    assert package_inventory(job)["source_payloads_verified"]
    document = Document(str(job / "auto.docx"))
    starts = document.element.body.xpath(".//w:bookmarkStart")
    first_name, second_name = [start.get(qn("w:name")) for start in starts]
    starts[0].set(qn("w:name"), second_name)
    starts[1].set(qn("w:name"), first_name)
    document.save(str(job / "auto.docx"))
    with pytest.raises(DemoError, match="SOURCE_RANGE_TEXT_MISMATCH"):
        package_inventory(job)
