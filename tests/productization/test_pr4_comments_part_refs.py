"""Comments and invalid part references cannot silently pass acceptance."""
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("tag", ["commentReference", "commentRangeStart", "commentRangeEnd"])
def test_comment_markers_are_unsupported(tmp_path: Path, tag: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    marker = OxmlElement("w:" + tag)
    marker.set(qn("w:id"), "0")
    doc.paragraphs[0].runs[0]._r.append(marker)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_BODY_CONTENT" in result["errors"]


@pytest.mark.parametrize("content", [True, False])
def test_comments_part_content(tmp_path: Path, content: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    root = OxmlElement("w:comments")
    if content:
        comment = OxmlElement("w:comment")
        comment.set(qn("w:id"), "0")
        paragraph, run, text = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "Synthetic comment"
        run.append(text)
        paragraph.append(run)
        comment.append(paragraph)
        root.append(comment)
    files["word/comments.xml"] = etree.tostring(root)
    types = etree.fromstring(files["[Content_Types].xml"])
    etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                     PartName="/word/comments.xml",
                     ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
    files["[Content_Types].xml"] = etree.tostring(types)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_COMMENTS" in result["errors"]) is content


@pytest.mark.parametrize("kind", ["header", "footer"])
@pytest.mark.parametrize("fault", ["missing", "image", "opposite", "valid"])
def test_story_reference_type(tmp_path: Path, kind: str, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    getattr(doc.sections[0], kind).paragraphs[0].text = ""
    reference = doc.sections[0]._sectPr.find(qn("w:" + kind + "Reference"))
    if fault == "missing":
        reference.set(qn("r:id"), "missing")
    elif fault == "image":
        reference.set(qn("r:id"), doc.paragraphs[-1]._p.xpath(".//a:blip")[0].get(qn("r:embed")))
    elif fault == "opposite":
        reference.tag = qn("w:" + ("footer" if kind == "header" else "header") + "Reference")
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("OPC_STORY_REFERENCE_INVALID" in result["errors"]) is (fault != "valid")


@pytest.mark.parametrize("part", ["styles", "numbering"])
@pytest.mark.parametrize("fault", ["wrong_root", "namespace", "valid"])
def test_style_numbering_roots(tmp_path: Path, part: str, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    name = "word/" + part + ".xml"
    root = etree.fromstring(files[name])
    if fault == "wrong_root":
        root.tag = qn("w:p")
    elif fault == "namespace":
        root.tag = "{urn:invalid}" + part
    files[name] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for n, data in files.items():
            archive.writestr(n, data)
    result = evaluate(path, truth, sources)
    assert ("OPC_WORD_ROOT_INVALID" in result["errors"]) is (fault != "valid")
