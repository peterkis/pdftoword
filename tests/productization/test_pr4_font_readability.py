"""Tiny source formats stay unchanged but cannot pass readable-output acceptance."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("tag", ["sz", "szCs"])
@pytest.mark.parametrize("where", ["text", "math", "style", "inherited", "default", "readable"])
def test_minimum_font_format(tmp_path: Path, tag: str, where: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if where == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif where in {"style", "inherited"}:
        style = doc.styles.add_style("TinyStyle", WD_STYLE_TYPE.PARAGRAPH)
        props = style._element.get_or_add_rPr()
        if where == "inherited":
            derived = doc.styles.add_style("DerivedTiny", WD_STYLE_TYPE.PARAGRAPH)
            derived.base_style = style
            style = derived
        doc.paragraphs[0].style = style
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    for existing in props.findall(qn("w:" + tag)):
        props.remove(existing)
    size = OxmlElement("w:" + tag)
    size.set(qn("w:val"), "12" if where == "readable" else "1")
    props.append(size)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_FONT_SCALE" in result["errors"]) is (where != "readable")
    assert result["structure_status"] == ("PASS" if where == "readable" else "FAIL")
