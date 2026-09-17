"""Default table numbering and legacy merges must not yield false passes."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("mode", ["default", "inherited", "conditional", "override", "disabled"])
def test_default_table_numbering(tmp_path: Path, mode: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.TABLE:
            style._element.attrib.pop(qn("w:default"), None)
    style = doc.styles.add_style("DefaultNumberedTable", WD_STYLE_TYPE.TABLE)
    style._element.set(qn("w:default"), "1")
    target = style._element
    if mode == "inherited":
        parent = doc.styles.add_style("NumberedParent", WD_STYLE_TYPE.TABLE)
        style.base_style = parent
        target = parent._element
    elif mode == "conditional":
        condition = OxmlElement("w:tblStylePr")
        condition.set(qn("w:type"), "firstRow")
        target.append(condition)
        target = condition
    props = OxmlElement("w:pPr")
    numbering, number = OxmlElement("w:numPr"), OxmlElement("w:numId")
    number.set(qn("w:val"), "1")
    numbering.append(number)
    props.append(numbering)
    target.append(props)
    table = doc.tables[0]
    table.style = None
    if mode == "override":
        table.style = doc.styles.add_style("PlainTable", WD_STYLE_TYPE.TABLE)
    elif mode == "disabled":
        for cell in table.rows[0].cells:
            prop = OxmlElement("w:numPr")
            number = OxmlElement("w:numId")
            number.set(qn("w:val"), "0")
            prop.append(number)
            cell.paragraphs[0]._p.get_or_add_pPr().append(prop)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    unsupported = mode not in {"override", "disabled"}
    assert ("UNSUPPORTED_NUMBERING" in result["errors"]) is unsupported
    assert result["structure_status"] == ("FAIL" if unsupported else "PASS")


@pytest.mark.parametrize("explicit_continue", [True, False])
def test_legacy_horizontal_merge_is_rejected(tmp_path: Path, explicit_continue: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    cells = list(doc.tables[0].rows[0].cells)
    for index, cell in enumerate(cells):
        merge = OxmlElement("w:hMerge")
        if index == 0 or explicit_continue:
            merge.set(qn("w:val"), "restart" if index == 0 else "continue")
        cell._tc.get_or_add_tcPr().append(merge)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_HORIZONTAL_MERGE" in result["errors"]
    assert result["structure_status"] == "FAIL"
