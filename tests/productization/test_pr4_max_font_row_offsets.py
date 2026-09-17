"""Large fonts and row-leading/trailing grid widths must fit the page."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("where", ["direct", "math", "style", "default"])
@pytest.mark.parametrize("tag", ["sz", "szCs"])
@pytest.mark.parametrize("value", ["3276", "144"])
def test_max_font_scale(tmp_path: Path, where: str, tag: str, value: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if where == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif where == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    for old in props.findall(qn("w:" + tag)):
        props.remove(old)
    node = OxmlElement("w:" + tag)
    node.set(qn("w:val"), value)
    props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_FONT_SCALE" in result["errors"]) is (value == "3276")
    assert result["structure_status"] == ("FAIL" if value == "3276" else "PASS")


@pytest.mark.parametrize("side", ["Before", "After"])
@pytest.mark.parametrize("overflow", [True, False])
def test_sparse_row_reserved_width(tmp_path: Path, side: str, overflow: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    table = doc.tables[0]
    table.autofit = False
    row = table.rows[0]._tr
    cells = row.findall(qn("w:tc"))
    removed, kept = (cells[0], cells[1]) if side == "Before" else (cells[1], cells[0])
    if side == "Before":
        for marker in removed.xpath(".//w:bookmarkStart | .//w:bookmarkEnd"):
            kept.find(qn("w:p")).append(marker)
    row.remove(removed)
    props = row.get_or_add_trPr()
    grid, width = OxmlElement("w:grid" + side), OxmlElement("w:w" + side)
    grid.set(qn("w:val"), "1")
    width.set(qn("w:w"), "4320")
    width.set(qn("w:type"), "dxa")
    props.append(grid)
    props.append(width)
    kept.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), "8640" if overflow else "4320")
    reference = next(u for u in truth["units"] if u["kind"] == "table")["reference"]
    reference["cells"] = [reference["cells"][1 if side == "Before" else 0]]
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TABLE_WIDTH" in result["errors"]) is overflow
    assert result["structure_status"] == ("FAIL" if overflow else "PASS")
