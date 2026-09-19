"""Actual Flow plan, OOXML and shared-continuation checks with no model calls."""

from __future__ import annotations

import copy
import shutil
import uuid
from collections.abc import Iterator
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from prototypes.docx_output.common import PRIVATE, DemoError, Json, digest, new_job, read, validate
from prototypes.docx_output.planning.flow import plan_flow
from prototypes.docx_output.renderers.flow import FlowRenderer
from prototypes.docx_output.structure_processors.docvortex import DocVortexStructureProcessor
from tests.productization.test_shared_structure_actual import continuation_ir, text_ir

FONTS = {"Arial", "Arial Unicode MS"}


@pytest.fixture
def case() -> Iterator[Path]:
    path = PRIVATE / ("flow-test-" + uuid.uuid4().hex)
    path.mkdir(parents=True)
    try:
        yield path
    finally:
        shutil.rmtree(path)


@pytest.mark.parametrize("size", [(300, 500), (700, 350)])
@pytest.mark.parametrize("body", [11, 15])
def test_paper_style_and_source_coordinates(case: Path, size: tuple[int, int], body: int) -> None:
    ir = text_ir(case, "Native text")
    p = ir["pages"][0]
    p.update(width_pt=size[0], height_pt=size[1])
    b = p["blocks"][0]
    b["type"] = "paragraph"
    b["bbox"] = [20, 20, size[0] - 20, 50]
    b["content"]["runs"] = [
        {
            "text": "Native text",
            "bbox": b["bbox"],
            "font_family": "Arial",
            "font_size_pt": body,
            "bold": True,
            "italic": True,
            "underline": True,
            "superscript": True,
            "subscript": False,
            "color": "112233",
            "confidence": 1,
            "source_ref": b["id"],
        }
    ]
    original = copy.deepcopy(ir)
    plan = plan_flow(ir, families=FONTS)
    assert plan.document["pages"] == original["pages"] and ir == original
    assert plan.document["schema_version"] == "layout-ir/1.2"
    assert plan.document["output_styles"]["body"]["font_size_pt"] == body
    assert plan.output_layout["sections"][0]["page_size_pt"] == list(size)
    validate(plan.document)
    job = new_job(case / "output")
    FlowRenderer().render(job, plan, "auto")
    doc = Document(str(job / "auto.docx"))
    section = doc.sections[0]
    assert section.page_width is not None and section.page_height is not None
    assert section.page_width.pt == pytest.approx(size[0], abs=0.05)
    assert section.page_height.pt == pytest.approx(size[1], abs=0.05)
    run = next(r for p in doc.paragraphs for r in p.runs if r.text)
    assert run._r.rPr is not None
    rf = run._r.rPr.rFonts
    assert rf is not None
    assert [rf.get(qn("w:" + key)) for key in ["ascii", "hAnsi", "eastAsia"]] == [
        "Arial",
        "Arial",
        "Arial Unicode MS",
    ]
    assert run.bold and run.italic and run.underline and run.font.superscript
    assert run.font.size is not None
    assert str(run.font.color.rgb) == "112233" and run.font.size.pt == body
    assert read(job / "source-map.auto.json")["blocks"][0]["bbox"] == b["bbox"]


def test_missing_cjk_is_reported_without_claiming_font_acceptance(case: Path) -> None:
    ir = text_ir(case, "中文")
    plan = plan_flow(ir, families=set())
    assert any(i["family"] == "Arial Unicode MS" for i in plan.output_layout["issues"])
    job = new_job(case / "missing-font")
    stats = FlowRenderer().render(job, plan, "auto")
    assert stats["fonts"]["missing"]


def test_no_automatic_centering_or_unknown_source_font_claim(case: Path) -> None:
    ir = text_ir(case, "Offset paragraph")
    p = ir["pages"][0]
    b = p["blocks"][0]
    b.update(type="paragraph", bbox=[150, 100, 400, 130], geometry_source="pp_structure")
    ir["metadata"].update(layout_profile="pp_geometry_flow", content_left_pt=0)
    plan = plan_flow(ir, families=FONTS)
    node = plan.output_layout["sections"][0]["nodes"][0]
    assert node["alignment"] == "left"
    assert plan.document["output_styles"]["body"]["basis"] == "inferred_scan_profile"
    job = new_job(case / "offset")
    FlowRenderer().render(job, plan, "auto")
    assert Document(str(job / "auto.docx")).paragraphs[0].alignment == 0


def test_only_shared_verified_continuation_joins(case: Path) -> None:
    ir = continuation_ir(case)
    unchanged = plan_flow(ir, families=FONTS)
    assert sum(len(s["nodes"]) for s in unchanged.output_layout["sections"]) == 2
    processed = DocVortexStructureProcessor().process(ir)
    plan = plan_flow(processed.document, families=FONTS)
    nodes = [n for s in plan.output_layout["sections"] for n in s["nodes"]]
    assert len(nodes) == 1 and len(nodes[0]["source_ids"]) == 2
    assert plan.output_layout["generated_join_spaces"] == 1
    job = new_job(case / "joined")
    FlowRenderer().render(job, plan, "auto")
    paragraphs = [p for p in Document(str(job / "auto.docx")).paragraphs if p.text]
    assert len(paragraphs) == 1
    assert paragraphs[0].text == "a paragraph continues with further words"
    assert len(read(job / "source-map.auto.json")["blocks"]) == 2


def test_tall_figure_uses_printable_height_and_retains_aspect(case: Path) -> None:
    ir = text_ir(case, "replace")
    p = ir["pages"][0]
    p.update(width_pt=300, height_pt=620)
    b = p["blocks"][0]
    b.update(type="figure", bbox=[10, 10, 20, 610])
    b["content"] = {
        "kind": "image",
        "asset_id": "tall",
        "content_mode": "figure",
        "render_mode": "preserve_as_image",
        "ocr_mode": "none",
        "placement_hint": "inline",
        "internal_text": [],
    }
    job = new_job(case / "tall")
    path = job / "assets/tall.png"
    Image.new("RGB", (50, 3000), "white").save(path)
    ir["assets"] = [
        {
            "id": "tall",
            "type": "image",
            "path": "assets/tall.png",
            "sha256": digest(path),
            "mime_type": "image/png",
            "source_page_index": 0,
            "source_bbox": b["bbox"],
            "extraction_method": "synthetic",
            "pixel_width": 50,
            "pixel_height": 3000,
        }
    ]
    plan = plan_flow(ir, families=FONTS)
    node = plan.output_layout["sections"][0]["nodes"][0]
    FlowRenderer().render(job, plan, "auto")
    shape = Document(str(job / "auto.docx")).inline_shapes[0]
    assert shape.height.pt <= node["height_limit_pt"] + 0.05
    assert shape.width / shape.height == pytest.approx(1 / 60, rel=1e-4)
    assert plan.document["assets"] == ir["assets"]


def test_edit_long_paragraph_replans_same_source_frame(case: Path) -> None:
    ir = text_ir(case, "short")
    plan = plan_flow(ir, {"body_size_pt": 15}, families=FONTS)
    edited = copy.deepcopy(plan.document)
    b = edited["pages"][0]["blocks"][0]
    bbox = list(b["bbox"])
    b["type"] = "paragraph"
    b["content"]["plain_text"] = "long paragraph " * 300
    b["content"]["runs"] = []
    revised = plan_flow(edited, families=FONTS)
    job = new_job(case / "edited")
    FlowRenderer().render(job, revised, "auto")
    assert revised.document["pages"][0]["blocks"][0]["bbox"] == bbox
    assert revised.document["output_styles"]["body"]["font_size_pt"] == 15
    assert Document(str(job / "auto.docx")).paragraphs[0].text == b["content"]["plain_text"]


@pytest.mark.parametrize(
    "profile",
    [
        {"body_size_pt": True},
        {"body_size_pt": 0},
        {"page_size_pt": [float("nan"), 300]},
        {"margins_pt": [1000, 0, 1000, 0]},
        {"east_asia_font": []},
        {"unknown": 1},
    ],
)
def test_invalid_profile_is_rejected(case: Path, profile: Json) -> None:
    with pytest.raises(DemoError):
        plan_flow(text_ir(case), profile, families=FONTS)


def test_plan_cannot_drop_source_nodes(case: Path) -> None:
    plan = plan_flow(text_ir(case), families=FONTS)
    plan.output_layout["sections"][0]["nodes"] = []
    with pytest.raises(DemoError, match="SOURCE_COVERAGE"):
        plan.as_dict()


def test_flow_defaults_do_not_change_legacy_paper(case: Path) -> None:
    from prototypes.docx_output.pipeline import finish
    from prototypes.docx_output.planning.render_plan import RenderPlan
    from prototypes.docx_output.structure_processors.base import StructureCandidate

    ir = text_ir(case)
    plan = RenderPlan.from_ir(ir)
    job = new_job(case / "legacy")
    finish(job, ir, render_plan=plan, structure_candidate=StructureCandidate(ir, {}))
    section = Document(str(job / "auto.docx")).sections[0]
    assert section.page_width is not None
    assert section.page_width.pt == pytest.approx(595.28, abs=0.05)
    assert read(job / "render-manifest.auto.json")["renderer"]["name"] == "legacy"


def test_render_plan_schema_alone_validates_embedded_ir(case: Path) -> None:
    from jsonschema import Draft202012Validator, ValidationError
    from prototypes.docx_output.common import ROOT

    value = plan_flow(text_ir(case), families=FONTS).as_dict()
    del value["document"]["pages"]
    with pytest.raises(ValidationError):
        Draft202012Validator(read(ROOT / "specs/render-plan.schema.json")).validate(value)


@pytest.mark.parametrize("override", [False, True])
def test_known_native_cjk_family_is_retained_unless_profile_overrides(
    case: Path, override: bool
) -> None:
    ir = text_ir(case, "中文")
    b = ir["pages"][0]["blocks"][0]
    b["content"]["runs"] = [
        {
            "text": "中文",
            "bbox": b["bbox"],
            "font_family": "Source CJK",
            "font_size_pt": 15,
            "bold": False,
            "italic": False,
            "underline": False,
            "superscript": False,
            "subscript": False,
            "color": None,
            "confidence": 1,
            "source_ref": b["id"],
        }
    ]
    profile = {"east_asia_font": "Arial Unicode MS"} if override else {}
    plan = plan_flow(ir, profile, families=FONTS | {"Source CJK"})
    node = plan.output_layout["sections"][0]["nodes"][0]
    assert node["runs"][b["id"]][0]["eastAsia"] == (
        "Arial Unicode MS" if override else "Source CJK"
    )
    assert (
        plan.document["pages"][0]["blocks"][0]["content"]["runs"][0]["font_family"] == "Source CJK"
    )


def test_cli_failure_does_not_print_private_path(
    case: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    from prototypes.docx_output.flow_replay import main

    private = case / "sensitive-source-name"
    private.mkdir()
    monkeypatch.setattr(
        "sys.argv",
        ["flow_replay", "--source-job", str(private), "--source-seal", str(case / "missing-seal")],
    )
    assert main() == 1
    error = capsys.readouterr().err
    assert "sensitive-source-name" not in error and str(case) not in error
