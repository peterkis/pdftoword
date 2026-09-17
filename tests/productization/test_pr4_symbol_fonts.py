"""Symbol fonts and theme indirection cannot preserve strict Unicode display claims."""

import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.docx_reader import NS
from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "case", ["direct", "math", "style", "default", "theme", "charset", "normal"]
)
def test_symbol_font_mapping(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    doc.paragraphs[0].runs[0].text = "abc"
    next(u for u in truth["units"] if u["unit_id"] == "a")["reference"]["text"] = "abc"
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if case == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif case == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif case == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    for existing in props.findall(qn("w:rFonts")):
        props.remove(existing)
    fonts = OxmlElement("w:rFonts")
    if case == "theme":
        fonts.set(qn("w:asciiTheme"), "minorHAnsi")
    else:
        fonts.set(
            qn("w:ascii"),
            "Arial" if case == "normal" else "CustomSymbols" if case == "charset" else "Wingdings",
        )
    props.append(fonts)
    doc.save(str(path))
    if case in {"theme", "charset"}:
        with zipfile.ZipFile(path) as archive:
            files = {n: archive.read(n) for n in archive.namelist()}
        if case == "theme":
            root = etree.fromstring(files["word/theme/theme1.xml"])
            latin = root.find("a:themeElements/a:fontScheme/a:minorFont/a:latin", NS)
            assert latin is not None
            latin.set("typeface", "Wingdings")
            files["word/theme/theme1.xml"] = etree.tostring(root)
        else:
            root = etree.fromstring(files["word/fontTable.xml"])
            font, charset = OxmlElement("w:font"), OxmlElement("w:charset")
            font.set(qn("w:name"), "CustomSymbols")
            charset.set(qn("w:val"), "02")
            font.append(charset)
            root.append(font)
            files["word/fontTable.xml"] = etree.tostring(root)
        with zipfile.ZipFile(path, "w") as archive:
            for n, data in files.items():
                archive.writestr(n, data)
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_FONT_MAPPING" in result["errors"]) is (case != "normal")
    assert result["structure_status"] == ("PASS" if case == "normal" else "FAIL")
