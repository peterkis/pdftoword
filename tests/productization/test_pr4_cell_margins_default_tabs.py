"""Cell margins and default tab stops cannot move required content off-page."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("where", ["cell", "table", "style"])
@pytest.mark.parametrize("width", ["108", "31680"])
def test_cell_margins(tmp_path: Path, where: str, width: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    if where == "cell":
        props = doc.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
        name = "tcMar"
    else:
        props = doc.tables[0]._tbl.tblPr
        name = "tblCellMar"
        if where == "style":
            style = doc.styles.add_style("WideCellMargins", WD_STYLE_TYPE.TABLE)
            props = OxmlElement("w:tblPr")
            style._element.append(props)
            doc.tables[0].style = style
    margins, left = OxmlElement("w:" + name), OxmlElement("w:left")
    left.set(qn("w:w"), width)
    left.set(qn("w:type"), "dxa")
    margins.append(left)
    props.append(margins)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_CELL_MARGINS" in result["errors"]) is (width != "108")
    assert result["structure_status"] == ("PASS" if width == "108" else "FAIL")


@pytest.mark.parametrize("contains_tab", [True, False])
@pytest.mark.parametrize("distance", ["720", "31680"])
def test_default_tab_stop(tmp_path: Path, contains_tab: bool, distance: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    if contains_tab:
        doc.paragraphs[0].runs[0].text = "A\tB"
        next(u for u in truth["units"] if u["unit_id"] == "a")["reference"]["text"] = "A\tB"
    for old in doc.settings._element.findall(qn("w:defaultTabStop")):
        doc.settings._element.remove(old)
    node = OxmlElement("w:defaultTabStop")
    node.set(qn("w:val"), distance)
    doc.settings._element.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    invalid = contains_tab and distance == "31680"
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is invalid
    assert result["structure_status"] == ("FAIL" if invalid else "PASS")
