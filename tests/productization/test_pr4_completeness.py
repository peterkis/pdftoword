"""Missing unsupported formulas and unexamined visible content cannot pass."""

import copy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("retained", ["math", "image", "missing", "empty"])
def test_unsupported_formula_must_still_be_present(tmp_path: Path, retained: str) -> None:
    path, truth, sources = specimen(tmp_path)
    formula = next(u for u in truth["units"] if u["kind"] == "formula")
    formula["reference"]["text"] = r"\begin{matrix}x\end{matrix}"
    # An unrelated preserved source image must not mask the missing formula either.
    truth["units"].append(
        {
            "unit_id": "image-label",
            "page": 1,
            "kind": "text",
            "status": "confirmed",
            "reference": {"text": "label", "source_anchor_id": "g", "content_scope": "figure_text"},
        }
    )
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[2]
    if retained == "empty":
        paragraph._p[-1].clear()
    elif retained != "math":
        paragraph._p.remove(paragraph._p[-1])
    if retained == "image":
        paragraph.add_run().add_picture(str(tmp_path / "figure.png"))
        image = copy.deepcopy(sources["blocks"][-1]["images"][0])
        image["bbox"] = [0, 30, 100, 50]
        image["fallback"] = True
        sources["blocks"][2]["images"] = [image]
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["content_status"] == (
        "FAIL" if retained in {"missing", "empty"} else "REVIEW_REQUIRED"
    )
    assert ("MISSING_FORMULA" in [f["code"] for f in result["failures"]]) is (
        retained in {"missing", "empty"}
    )
    if retained in {"missing", "empty"}:
        assert result["metrics"]["necessary_content"]["status"] == "FAIL"


@pytest.mark.parametrize("story", ["header", "footer"])
@pytest.mark.parametrize("content", ["text", "formula", "image", "empty"])
def test_referenced_nonbody_story_is_checked(tmp_path: Path, story: str, content: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = getattr(doc.sections[0], story).paragraphs[0]
    if content == "text":
        paragraph.add_run("Unexpected story text")
    elif content == "formula":
        paragraph._p.append(copy.deepcopy(doc.paragraphs[2]._p[-1]))
    elif content == "image":
        paragraph.add_run().add_picture(str(tmp_path / "figure.png"))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_VISIBLE_STORY" in result["errors"]) is (content != "empty")
    assert result["structure_status"] == ("PASS" if content == "empty" else "FAIL")


@pytest.mark.parametrize(
    "complete,uncertain,expected",
    [
        (True, False, "FAIL"),
        (False, False, "NOT_SCORED"),
        (True, True, "NOT_SCORED"),
    ],
)
def test_zero_reference_edges_distinguish_complete_and_unknown_truth(
    tmp_path: Path, complete: bool, uncertain: bool, expected: str
) -> None:
    path, truth, sources = specimen(tmp_path)
    truth["coverage_complete"] = complete
    if uncertain:
        next(u for u in truth["units"] if u["kind"] == "figure_edge")["status"] = "uncertain"
    else:
        truth["units"] = [u for u in truth["units"] if u["kind"] != "figure_edge"]
    result = evaluate(path, truth, sources)
    assert result["metrics"]["relation_precision"]["eligible_count"] == 1
    assert result["metrics"]["relation_precision"]["status"] == expected
    if expected == "FAIL":
        assert result["structure_status"] == "FAIL"
