"""Display math and row property exceptions must not escape inspection."""
import copy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("place", ["body", "cell", "paragraph"])
def test_display_math_outside_word_paragraph(tmp_path: Path, place: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    math = doc.paragraphs[2]._p.xpath(".//m:oMath")[0]
    display = OxmlElement("m:oMathPara")
    if place == "paragraph":
        math.getparent().remove(math)
        display.append(math)
        doc.paragraphs[2]._p.append(display)
    else:
        display.append(copy.deepcopy(math))
        parent = doc._element.body if place == "body" else doc.tables[0].cell(0, 0)._tc
        parent.append(display)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_BLOCK_MATH" in result["errors"]) is (place != "paragraph")
    assert result["structure_status"] == ("PASS" if place == "paragraph" else "FAIL")


@pytest.mark.parametrize("fill", ["000000", "FFFFFF", "nil"])
def test_row_property_exception_shading(tmp_path: Path, fill: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    exceptions, shading = OxmlElement("w:tblPrEx"), OxmlElement("w:shd")
    shading.set(qn("w:val"), "nil" if fill == "nil" else "clear")
    shading.set(qn("w:fill"), fill if fill != "nil" else "auto")
    exceptions.append(shading)
    doc.tables[0].rows[0]._tr.insert(0, exceptions)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_VISIBILITY_STYLE" in result["errors"]) is (fill != "nil")
    assert result["structure_status"] == ("PASS" if fill == "nil" else "FAIL")
