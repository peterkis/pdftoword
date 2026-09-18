"""Real public renderer, registered assets, content range mapping and explicit fallback."""

from __future__ import annotations

import copy
import shutil
import uuid
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from lxml import etree
from prototypes.docx_output.common import PRIVATE, DemoError, digest, new_job, read, save
from prototypes.docx_output.docvortex_runtime import call_worker
from prototypes.docx_output.planning.render_plan import RenderPlan
from prototypes.docx_output.render_replay import compare_renderers, inventory
from prototypes.docx_output.renderers.docvortex import DocVortexRenderer
from prototypes.docx_output.structure_processors.bridge import bridge, direct_middle
from tests.productization.test_renderer_boundary import source_job


@pytest.fixture
def case() -> Iterator[Path]:
    """Use real private files and isolated worker, never mock the public renderer."""
    root = PRIVATE / ("docvortex-render-test-" + uuid.uuid4().hex)
    root.mkdir(parents=True)
    try:
        yield root
    finally:
        shutil.rmtree(root)


def test_real_renderer_same_plan_and_source_ranges(case: Path) -> None:
    source, ir = source_job(case)
    before = inventory(source)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    outputs = read(comparison / "comparison.json")["outputs"]
    a, b = [case / "jobs" / row["job_id"] for row in outputs]
    assert digest(a / "render-plan.auto.json") == digest(b / "render-plan.auto.json")
    audit = read(b / "docvortex-render-audit.auto.json")
    assert audit["public_call"] == "COMPLETE" and not audit["fallback"]
    assert not audit["worker"]["pdfium_loaded"]
    mapping = read(b / "source-map.auto.json")
    assert mapping["schema_version"] == "docx-source-map/1.0"
    assert mapping["docx_sha256"] == digest(b / "auto.docx")
    assert [x["block_id"] for x in mapping["blocks"]] == ["question", "figure"]
    with zipfile.ZipFile(b / "auto.docx") as package:
        tree = etree.fromstring(package.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        starts = tree.xpath('//w:bookmarkStart[starts-with(@w:name,"p2w_")]', namespaces=ns)
        assert len(starts) == 2
        assert starts[0].getnext().tag.endswith("}r")
    assert Document(str(b / "auto.docx")).styles["Normal"].font.size.pt == 11
    assert inventory(source) == before
    assert read(b / "layout.auto.json")["pages"] == ir["pages"]


def test_unsupported_geometry_falls_back_without_losing_content(case: Path) -> None:
    source, ir = source_job(case)
    ir["metadata"]["layout_profile"] = "pp_geometry_flow"
    save(source / "layout.auto.json", ir)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    b = case / "jobs" / read(comparison / "comparison.json")["outputs"][1]["job_id"]
    audit = read(b / "docvortex-render-audit.auto.json")
    assert audit["fallback"] and audit["public_call"] == "NOT_RUN"
    assert read(b / "qa.json")["effective_renderer"] == "legacy"
    assert Document(str(b / "auto.docx")).paragraphs[0].text == "1. Keep < exact > content"


@pytest.mark.parametrize("fault", ["missing", "changed", "remote", "traversal"])
def test_worker_asset_resolver_rejects_unregistered_inputs(case: Path, fault: str) -> None:
    source, ir = source_job(case)
    value = bridge(ir, for_renderer=True)
    middle = direct_middle(value)
    asset = ir["assets"][0]
    registry = {asset["path"]: asset["sha256"]}
    if fault == "missing":
        registry.clear()
    elif fault == "changed":
        registry[asset["path"]] = "0" * 64
    else:
        body = middle["pages"][0]["blocks"][1]["content"][0]
        if fault == "remote":
            body["image_url"] = "https://example.invalid/no-network.png"
        else:
            body["image_path"] = "../outside.png"
    with pytest.raises(DemoError, match="DOCVORTEX_PUBLIC_CALL_FAILED"):
        call_worker(
            {
                "action": "render",
                "middle": middle,
                "assets": registry,
                "job": str(source),
                "output": str(case / "forbidden.docx"),
            }
        )
    assert not (case / "forbidden.docx").exists()


def test_auxiliary_text_is_preserved_by_adapter(case: Path) -> None:
    source, ir = source_job(case)
    ir["pages"][0]["blocks"][0]["type"] = "footer"
    save(source / "layout.auto.json", ir)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    b = case / "jobs" / read(comparison / "comparison.json")["outputs"][1]["job_id"]
    assert not read(b / "docvortex-render-audit.auto.json")["fallback"]
    assert Document(str(b / "auto.docx")).paragraphs[0].text == "1. Keep < exact > content"


def test_formula_image_fallback_and_raw_latex_rejection(case: Path) -> None:
    source, ir = source_job(case)
    b = ir["pages"][0]["blocks"][1]
    b["type"] = "formula"
    b["content"] = {
        "kind": "formula",
        "latex": r"\frac{",
        "mathml": None,
        "source_asset_id": ir["assets"][0]["id"],
        "render_mode": "omml_with_image_fallback",
        "confidence": 0,
        "omml_status": "pending",
    }
    target = new_job(case / "jobs")
    for asset in ir["assets"]:
        dest = target / asset["path"]
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source / asset["path"], dest)
    stats = DocVortexRenderer().render(target, RenderPlan.from_ir(ir), "auto")
    assert stats["formula_image_count"] == 1 and stats["omml_formula_count"] == 0
    assert (
        read(target / "docvortex-render-audit.auto.json")["losses"][0]["code"]
        == "FORMULA_IMAGE_FALLBACK"
    )
    ir2 = copy.deepcopy(ir)
    ir2["pages"][0]["blocks"][1]["content"]["source_asset_id"] = "unavailable"
    target2 = new_job(case / "jobs")
    for asset in ir2["assets"]:
        dest = target2 / asset["path"]
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source / asset["path"], dest)
    with pytest.raises(DemoError):
        DocVortexRenderer().render(target2, RenderPlan.from_ir(ir2), "auto")
    assert not (target2 / "auto.docx").exists()


def test_integrity_failure_never_downgrades_to_legacy(case: Path) -> None:
    source, ir = source_job(case)
    target = new_job(case / "jobs")
    for asset in ir["assets"]:
        dest = target / asset["path"]
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source / asset["path"], dest)
    ir["assets"][0]["sha256"] = "0" * 64
    with pytest.raises(DemoError, match="ASSET_HASH"):
        DocVortexRenderer().render(target, RenderPlan.from_ir(ir), "auto")
    assert not (target / "auto.docx").exists()


def test_poc_cannot_allocate_inside_original_job(case: Path) -> None:
    from prototypes.docx_output.reuse_poc import run_poc

    source, _ = source_job(case)
    before = inventory(source)
    with pytest.raises(DemoError, match="OUTPUT_INSIDE_SOURCE"):
        run_poc(source, output_root=source / "nested")
    assert inventory(source) == before


def test_inline_formula_is_written_as_actual_omml(case: Path) -> None:
    source, ir = source_job(case)
    ir["pages"][0]["blocks"][0]["content"]["plain_text"] = "Value x^2"
    ir["metadata"]["structure_evidence"] = {
        "question": {
            "inline_spans": [
                {"type": "text", "content": "Value "},
                {"type": "equation_inline", "content": "x^2"},
            ]
        }
    }
    save(source / "layout.auto.json", ir)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    b = case / "jobs" / read(comparison / "comparison.json")["outputs"][1]["job_id"]
    audit = read(b / "docvortex-render-audit.auto.json")
    assert audit["public_call"] == "COMPLETE" and not audit["fallback"]
    assert read(b / "qa.json")["omml_formula_count"] == 1


@pytest.mark.parametrize("source_kind", ["data", "remote"])
def test_table_html_cannot_bypass_local_asset_registry(case: Path, source_kind: str) -> None:
    import base64

    source, ir = source_job(case)
    middle = direct_middle(bridge(ir, for_renderer=True))
    asset = ir["assets"][0]
    url = "https://example.invalid/unregistered.png"
    if source_kind == "data":
        encoded = base64.b64encode((source / asset["path"]).read_bytes()).decode()
        url = "data:image/png;base64," + encoded
    image = middle["pages"][0]["blocks"][1]
    image["type"] = "table"
    image["content"][0] = {
        "type": "table_body",
        "index": image["index"],
        "bbox": image["bbox"],
        "content": f'<table><tr><td><img src="{url}"/></td></tr></table>',
    }
    with pytest.raises(DemoError, match="DOCVORTEX_PUBLIC_CALL_FAILED"):
        call_worker(
            {
                "action": "render",
                "middle": middle,
                "assets": {asset["path"]: asset["sha256"]},
                "job": str(source),
                "output": str(case / "forbidden.docx"),
            }
        )
    assert not (case / "forbidden.docx").exists()


def test_editable_table_has_one_verified_source_range(case: Path) -> None:
    source, ir = source_job(case)
    ir["pages"][0]["blocks"][0]["type"] = "table"
    ir["pages"][0]["blocks"][0]["content"]["plain_text"] = "ABC"
    ir["metadata"]["structure_evidence"] = {
        "question": {
            "selected_html": (
                '<table><tr><td rowspan="2">A</td><td>B</td></tr><tr><td>C</td></tr></table>'
            )
        }
    }
    save(source / "layout.auto.json", ir)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    b = case / "jobs" / read(comparison / "comparison.json")["outputs"][1]["job_id"]
    assert not read(b / "docvortex-render-audit.auto.json")["fallback"]
    assert len(Document(str(b / "auto.docx")).tables) == 1
    mapping = read(b / "source-map.auto.json")
    assert mapping["blocks"][0]["kind"] == "table"
    with zipfile.ZipFile(b / "auto.docx") as package:
        xml = etree.fromstring(package.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        marker = xml.xpath(
            "//w:bookmarkStart[@w:name=$name]", namespaces=ns, name=mapping["blocks"][0]["marker"]
        )[0]
        assert marker.getnext().tag.endswith("}tbl")
        assert marker.getnext().getnext().tag.endswith("}bookmarkEnd")


def test_deduplicated_image_paths_keep_distinct_source_boxes(case: Path) -> None:
    source, ir = source_job(case)
    first = ir["assets"][0]
    duplicate = copy.deepcopy(first)
    duplicate.update(id="second-asset", source_bbox=[110, 150, 200, 200])
    ir["assets"].append(duplicate)
    second = copy.deepcopy(ir["pages"][0]["blocks"][1])
    second["id"] = "second-figure"
    second["bbox"] = duplicate["source_bbox"]
    second["content"]["asset_id"] = duplicate["id"]
    ir["pages"][0]["blocks"].append(second)
    ir["pages"][0]["reading_order"].append(second["id"])
    save(source / "layout.auto.json", ir)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    target = case / "jobs" / read(comparison / "comparison.json")["outputs"][1]["job_id"]
    records = read(target / "source-map.auto.json")["blocks"]
    assert records[2]["images"][0]["bbox"] == duplicate["source_bbox"]
    assert records[1]["images"][0]["bbox"] == first["source_bbox"]


def test_ordinary_html_table_with_void_tag_and_entity(case: Path) -> None:
    source, ir = source_job(case)
    ir["pages"][0]["blocks"][0]["type"] = "table"
    ir["pages"][0]["blocks"][0]["content"]["plain_text"] = "A\nB\u00a0C"
    ir["metadata"]["structure_evidence"] = {
        "question": {"selected_html": "<table><tr><td>A<br>B&nbsp;C</td></tr></table>"}
    }
    save(source / "layout.auto.json", ir)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    target = case / "jobs" / read(comparison / "comparison.json")["outputs"][1]["job_id"]
    assert not read(target / "docvortex-render-audit.auto.json")["fallback"]
    assert Document(str(target / "auto.docx")).tables[0].cell(0, 0).text == "A\nB\u00a0C"


@pytest.mark.parametrize("change", ["boundaries", "merge"])
def test_table_topology_change_rejected_before_source_binding(case: Path, change: str) -> None:
    from prototypes.docx_output.renderers.docvortex import bind_source_ranges

    _, ir = source_job(case)
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    if change == "boundaries":
        table.cell(0, 0).text, table.cell(0, 1).text = "AB", "C"
    else:
        table.cell(0, 0).merge(table.cell(0, 1)).text = "ABC"
    raw = case / "altered-table.docx"
    doc.save(str(raw))
    entry = {
        "block": ir["pages"][0]["blocks"][0],
        "raw": {"type": "table", "content": "<table><tr><td>A</td><td>BC</td></tr></table>"},
    }
    with pytest.raises(DemoError, match="DOCVORTEX_TABLE_TOPOLOGY"):
        bind_source_ranges(raw, case / "rejected.docx", ir, [entry])
    assert not (case / "rejected.docx").exists()


@pytest.mark.parametrize("inline", [False, True])
def test_changed_formula_rejected_before_source_binding(case: Path, inline: bool) -> None:
    from prototypes.docx_output.formula import to_omml
    from prototypes.docx_output.renderers.docvortex import bind_source_ranges

    _, ir = source_job(case)
    doc = Document()
    paragraph: Any = doc.add_paragraph()._p
    paragraph.append(etree.fromstring(to_omml("x^3").encode()))
    raw = case / "altered-formula.docx"
    doc.save(str(raw))
    source = (
        {"type": "text", "content": [{"type": "equation_inline", "content": "x^2"}]}
        if inline
        else {"type": "equation", "content": "x^2"}
    )
    with pytest.raises(DemoError, match="DOCVORTEX_FORMULA_UNVERIFIED"):
        bind_source_ranges(
            raw, case / "rejected.docx", ir, [{"block": ir["pages"][0]["blocks"][0], "raw": source}]
        )
    assert not (case / "rejected.docx").exists()


@pytest.mark.parametrize("markup", [None, "", "   "])
def test_table_without_structure_evidence_falls_back(case: Path, markup: str | None) -> None:
    source, ir = source_job(case)
    ir["pages"][0]["blocks"][0]["type"] = "table"
    ir["metadata"]["structure_evidence"] = {"question": {"selected_html": markup}}
    save(source / "layout.auto.json", ir)
    comparison = compare_renderers(
        source, output_root=case / "jobs", renderer_b=DocVortexRenderer()
    )
    target = case / "jobs" / read(comparison / "comparison.json")["outputs"][1]["job_id"]
    audit = read(target / "docvortex-render-audit.auto.json")
    assert audit["fallback"]
    assert any(loss["code"] == "TABLE_STRUCTURE_EVIDENCE_MISSING" for loss in audit["losses"])
