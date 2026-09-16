"""No revision view is guessed when reading reviewed Word documents."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


def revision(tag: str) -> Any:
    """Create a synthetic, traceable Word revision marker."""
    node = OxmlElement("w:" + tag)
    node.set(qn("w:id"), "1")
    node.set(qn("w:author"), "Synthetic reviewer")
    node.set(qn("w:date"), "2026-09-16T00:00:00Z")
    return node


@pytest.mark.parametrize(
    "case",
    [
        "move_text",
        "delete_formula",
        "delete_image",
        "insert_text",
        "format_change",
        "cell_delete",
        "header_move",
        "move_range",
        "enabled_only",
    ],
)
def test_pending_revisions_cannot_supply_acceptance_content(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0]
    if case in {"move_text", "insert_text"}:
        wrapper = revision("moveFrom" if case == "move_text" else "ins")
        wrapper.append(paragraph.runs[0]._r)
        paragraph._p.append(wrapper)
        if case == "move_text":
            paragraph._p.append(revision("moveTo"))
    elif case == "delete_formula":
        paragraph = doc.paragraphs[2]
        wrapper = revision("del")
        wrapper.append(paragraph._p[-1])
        paragraph._p.append(wrapper)
    elif case == "delete_image":
        paragraph = doc.paragraphs[-1]
        wrapper = revision("del")
        wrapper.append(paragraph.runs[0]._r)
        paragraph._p.append(wrapper)
    elif case == "format_change":
        change = revision("rPrChange")
        change.append(OxmlElement("w:rPr"))
        paragraph.runs[0]._r.get_or_add_rPr().append(change)
    elif case == "cell_delete":
        doc.tables[0].cell(0, 0)._tc.get_or_add_tcPr().append(revision("cellDel"))
    elif case == "header_move":
        wrapper = revision("moveFrom")
        doc.sections[0].header.paragraphs[0]._p.append(wrapper)
    elif case == "move_range":
        paragraph._p.insert(0, revision("moveFromRangeStart"))
        paragraph._p.append(revision("moveFromRangeEnd"))
    else:
        doc.settings._element.append(OxmlElement("w:trackRevisions"))
        paragraph._p.set(qn("w:rsidR"), "00000001")
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_REVISIONS" in result["errors"]) is (case != "enabled_only")
    assert result["structure_status"] == ("PASS" if case == "enabled_only" else "FAIL")
