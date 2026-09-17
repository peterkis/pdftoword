"""Inline pictures must fit their paragraph or cell, not just a whole page."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from tests.productization.test_acceptance_harness import specimen

from acceptance.docx_reader import inspect


@pytest.mark.parametrize(
    "where", ["plain", "direct_indent", "style_indent", "first_cell", "second_cell"]
)
@pytest.mark.parametrize("large", [True, False])
def test_inline_image_container_width(tmp_path: Path, where: str, large: bool) -> None:
    path, _, _ = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[-1]
    if where == "direct_indent":
        paragraph.paragraph_format.left_indent = Pt(90)
    elif where == "style_indent":
        style = doc.styles.add_style("IndentedPicture", WD_STYLE_TYPE.PARAGRAPH)
        style.paragraph_format.left_indent = Pt(90)
        paragraph.style = style
    elif where in {"first_cell", "second_cell"}:
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0 if where == "first_cell" else 1)._tc.append(paragraph._p)
    if large:
        section = doc.sections[0]
        extent = section.page_width - section.left_margin - section.right_margin
        for node in paragraph._p.xpath(".//wp:extent | .//a:xfrm/a:ext"):
            node.set("cx", str(extent))
            node.set("cy", str(extent))
    doc.save(str(path))
    result = inspect(path)
    invalid = large and where != "plain"
    assert ("IMAGE_DISPLAY_SCALE_INVALID" in result["errors"]) is invalid
    if not invalid:
        assert result["errors"] == []
