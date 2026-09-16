"""Visible auto-numbering and ambiguous type declarations cannot be ignored."""

import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "mode",
    ["direct", "style", "derived", "bullet", "disabled", "unused", "header_empty", "table_style"],
)
def test_auto_numbering_is_explicitly_unsupported(tmp_path: Path, mode: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0]
    if mode in {"style", "disabled"}:
        paragraph.style = "List Number"
    elif mode == "bullet":
        paragraph.style = "List Bullet"
    elif mode in {"derived", "unused"}:
        style = doc.styles.add_style("DerivedNumbering", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["List Number"]
        if mode == "derived":
            paragraph.style = style
    elif mode == "header_empty":
        doc.sections[0].header.paragraphs[0].style = "List Bullet"
    elif mode == "table_style":
        style = doc.styles.add_style("TableNumbering", WD_STYLE_TYPE.TABLE)
        props = OxmlElement("w:pPr")
        numbering, number = OxmlElement("w:numPr"), OxmlElement("w:numId")
        number.set(qn("w:val"), "1")
        numbering.append(number)
        props.append(numbering)
        style._element.append(props)
        doc.tables[0].style = style
    if mode in {"direct", "disabled"}:
        numbering, number = OxmlElement("w:numPr"), OxmlElement("w:numId")
        number.set(qn("w:val"), "0" if mode == "disabled" else "1")
        numbering.append(number)
        paragraph._p.get_or_add_pPr().append(numbering)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    if mode in {"disabled", "unused"}:
        assert result["structure_status"] == "PASS"
    else:
        code = "UNSUPPORTED_VISIBLE_STORY" if mode == "header_empty" else "UNSUPPORTED_NUMBERING"
        assert code in result["errors"]
        assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize(
    "case",
    ["override_same", "override_conflict", "default_same", "default_conflict", "default_case"],
)
def test_duplicate_content_type_keys_are_rejected(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    root = etree.fromstring(files["[Content_Types].xml"])
    source = (
        next(n for n in root if n.get("PartName") == "/word/document.xml")
        if case.startswith("override")
        else next(n for n in root if n.get("Extension") == "png")
    )
    duplicate = copy.deepcopy(source)
    if case.endswith("conflict"):
        duplicate.set("ContentType", "text/plain")
    if case == "default_case":
        duplicate.set("Extension", "PNG")
    root.insert(0, duplicate)
    files["[Content_Types].xml"] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    assert "OPC_DUPLICATE_CONTENT_TYPE" in result["errors"]
    assert result["structure_status"] == "FAIL"
