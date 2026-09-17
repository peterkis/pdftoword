"""Hidden rows and enforced editing restrictions must fail acceptance."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("value", [None, "true", "1", "false", "0"])
def test_hidden_table_row(tmp_path: Path, value: str | None) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    hidden = OxmlElement("w:hidden")
    if value is not None:
        hidden.set(qn("w:val"), value)
    doc.tables[0].rows[0]._tr.get_or_add_trPr().append(hidden)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("HIDDEN_CONTENT" in result["errors"]) is (value not in {"false", "0"})


@pytest.mark.parametrize("value", [None, "true", "1", "on", "false", "0"])
def test_enforced_document_protection(tmp_path: Path, value: str | None) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    protection = OxmlElement("w:documentProtection")
    protection.set(qn("w:edit"), "readOnly")
    if value is not None:
        protection.set(qn("w:enforcement"), value)
    doc.settings._element.append(protection)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_DOCUMENT_PROTECTION" in result["errors"]) is (
        value in {"true", "1", "on"}
    )


@pytest.mark.parametrize("case", ["style", "inherited", "conditional", "off"])
def test_table_style_row_hidden(tmp_path: Path, case: str) -> None:
    from docx.enum.style import WD_STYLE_TYPE

    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    style = doc.styles.add_style("HiddenRowStyle", WD_STYLE_TYPE.TABLE)
    props, hidden = OxmlElement("w:trPr"), OxmlElement("w:hidden")
    hidden.set(qn("w:val"), "false" if case == "off" else "true")
    props.append(hidden)
    if case == "conditional":
        conditional = OxmlElement("w:tblStylePr")
        conditional.set(qn("w:type"), "firstRow")
        conditional.append(props)
        style._element.append(conditional)
    else:
        style._element.append(props)
    if case == "inherited":
        derived = doc.styles.add_style("DerivedHiddenRowStyle", WD_STYLE_TYPE.TABLE)
        derived.base_style = style
        style = derived
    doc.tables[0].style = style
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_VISIBILITY_STYLE" in result["errors"]) is (case != "off")
