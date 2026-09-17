"""Every package part, used table style and predicted edge needs distinct validation."""

import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "fault", ["missing_rels", "wrong_rels", "styles", "header", "footer", "undeclared_part"]
)
def test_all_parts_and_used_word_parts_need_valid_types(tmp_path: Path, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    if fault in {"header", "footer"}:
        doc = Document(str(path))
        getattr(doc.sections[0], fault).paragraphs[0].text = ""
        doc.save(str(path))
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    types = etree.fromstring(files["[Content_Types].xml"])
    if fault in {"missing_rels", "wrong_rels"}:
        declaration = next(c for c in types if c.get("Extension") == "rels")
        if fault == "missing_rels":
            types.remove(declaration)
        else:
            declaration.set("ContentType", "application/xml")
    elif fault == "undeclared_part":
        files["custom/data.unknown"] = b"synthetic package part"
    else:
        name = "/word/styles.xml" if fault == "styles" else "/word/" + fault + "1.xml"
        next(c for c in types if c.get("PartName") == name).set("ContentType", "application/xml")
    files["[Content_Types].xml"] = etree.tostring(types)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    assert any(c.startswith("OPC_") for c in result["errors"])
    assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize("mode", ["whole", "inherited", "override", "unused", "conditional"])
def test_used_table_style_visibility_is_not_ignored(tmp_path: Path, mode: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    style = doc.styles.add_style("TableHidden", WD_STYLE_TYPE.TABLE)
    if mode == "conditional":
        condition, props = OxmlElement("w:tblStylePr"), OxmlElement("w:rPr")
        condition.set(qn("w:type"), "firstRow")
        props.append(OxmlElement("w:vanish"))
        condition.append(props)
        style._element.append(condition)
    else:
        style.font.hidden = True
    if mode == "inherited":
        derived = doc.styles.add_style("DerivedTableHidden", WD_STYLE_TYPE.TABLE)
        derived.base_style = style
        style = derived
    if mode != "unused":
        doc.tables[0].style = style
    if mode == "override":
        for cell in doc.tables[0].rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.font.hidden = False
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["structure_status"] == ("PASS" if mode in {"override", "unused"} else "FAIL")
    if mode not in {"override", "unused"}:
        assert any(
            c in result["errors"] for c in ["HIDDEN_CONTENT", "UNSUPPORTED_VISIBILITY_STYLE"]
        )


@pytest.mark.parametrize("copies", [1, 2, 3])
def test_each_expected_relation_can_only_match_once(tmp_path: Path, copies: int) -> None:
    path, truth, sources = specimen(tmp_path)
    sources["relations"] = [copy.deepcopy(sources["relations"][0]) for _ in range(copies)]
    result = evaluate(path, truth, sources)
    metric = result["metrics"]["relation_precision"]
    assert metric["correct_count"] == 1
    assert metric["value"] == 1 / copies
    assert result["metrics"]["relation_coverage"]["correct_count"] == 1
    assert result["structure_status"] == ("PASS" if copies == 1 else "FAIL")


def test_content_type_must_have_valid_type_and_subtype(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    files["custom/data.bin"] = b"synthetic extra part"
    types = etree.fromstring(files["[Content_Types].xml"])
    entry = etree.SubElement(
        types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override"
    )
    entry.set("PartName", "/custom/data.bin")
    entry.set("ContentType", "/")
    files["[Content_Types].xml"] = etree.tostring(types)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    assert "OPC_PART_CONTENT_TYPE_INVALID" in result["errors"]
