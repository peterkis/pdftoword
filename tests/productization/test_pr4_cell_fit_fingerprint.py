"""Cell fitting cannot hide content, and runtime contracts belong to provenance."""

import shutil
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

import product_acceptance
from acceptance.metrics import evaluate, semantic_hash


@pytest.mark.parametrize("where", ["cell", "style"])
@pytest.mark.parametrize("enabled", [True, False])
def test_cell_fit_text(tmp_path: Path, where: str, enabled: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
    if where == "style":
        style = doc.styles.add_style("FittedCells", WD_STYLE_TYPE.TABLE)
        props = OxmlElement("w:tcPr")
        style._element.append(props)
        doc.tables[0].style = style
    fit = OxmlElement("w:tcFitText")
    fit.set(qn("w:val"), "true" if enabled else "false")
    props.append(fit)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is enabled
    assert result["structure_status"] == ("FAIL" if enabled else "PASS")


@pytest.mark.parametrize("changed", ["specs/layout-ir.schema.json", "uv.lock"])
def test_conversion_fingerprint_covers_runtime_contracts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, changed: str
) -> None:
    source = tmp_path / "source"
    for folder in ["prototypes/docx_output", "scripts/acceptance"]:
        shutil.copytree(product_acceptance.ROOT / folder, source / folder)
    for relative in [
        "scripts/product_acceptance.py",
        "scripts/docx_demo.py",
        "specs/layout-ir.schema.json",
        "uv.lock",
    ]:
        target = source / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(product_acceptance.ROOT / relative, target)
    monkeypatch.setattr(product_acceptance, "ROOT", source)
    first = product_acceptance.conversion_source_files()
    assert changed in first
    target = source / changed
    target.write_bytes(target.read_bytes() + b"\n")
    second = product_acceptance.conversion_source_files()
    assert semantic_hash(first) != semantic_hash(second)
