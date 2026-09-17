"""Content locks and unsupported extension effects must not pass acceptance."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "lock", [None, "unlocked", "sdtLocked", "contentLocked", "sdtContentLocked"]
)
def test_content_control_lock(tmp_path: Path, lock: str | None) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0]._p
    parent = paragraph.getparent()
    position = parent.index(paragraph)
    control, props, content = (
        OxmlElement("w:sdt"),
        OxmlElement("w:sdtPr"),
        OxmlElement("w:sdtContent"),
    )
    if lock is not None:
        node = OxmlElement("w:lock")
        node.set(qn("w:val"), lock)
        props.append(node)
    control.append(props)
    content.append(paragraph)
    control.append(content)
    parent.insert(position, control)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    locked = lock in {"contentLocked", "sdtContentLocked"}
    assert ("UNSUPPORTED_CONTENT_LOCK" in result["errors"]) is locked
    assert result["structure_status"] == ("FAIL" if locked else "PASS")


@pytest.mark.parametrize("where", ["text", "math", "style", "default"])
@pytest.mark.parametrize("effect", ["textFill", "textOutline"])
def test_extended_text_rendering(tmp_path: Path, where: str, effect: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if where == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif where == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    extended = etree.SubElement(
        props, "{http://schemas.microsoft.com/office/word/2010/wordml}" + effect
    )
    extended.append(OxmlElement("a:noFill"))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNSUPPORTED_VISIBILITY_STYLE" in result["errors"]
    assert result["structure_status"] == "FAIL"
