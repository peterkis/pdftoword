"""Uncertain edges and merged-cell payloads retain their distinct semantics."""

import copy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("case", ["mixed", "duplicate", "unknown", "absent", "unbound"])
def test_uncertain_edge_predictions_are_claimed_once(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    unit = copy.deepcopy(next(u for u in truth["units"] if u["kind"] == "figure_edge"))
    unit.update(unit_id="uncertain-edge", status="uncertain")
    unit["reference"]["to"] = "f"
    truth["units"].append(unit)
    prediction = dict(sources["relations"][0], to="f")
    if case != "absent":
        sources["relations"].append(prediction)
    if case == "duplicate":
        sources["relations"].append(copy.deepcopy(prediction))
    elif case == "unknown":
        sources["relations"].append(dict(prediction, to="b"))
    elif case == "unbound":
        sources["relations"][-1]["to"] = "nonexistent"
    result = evaluate(path, truth, sources)
    metric = result["metrics"]["relation_precision"]
    bad = case in {"duplicate", "unknown", "unbound"}
    excluded = case in {"mixed", "duplicate", "unknown"}
    assert metric["scored_count"] == (2 if bad else 1)
    assert metric["correct_count"] == 1
    assert metric["uncertain_count"] == int(excluded)
    assert metric["not_scored_count"] == int(excluded)
    assert metric["value"] == (0.5 if bad else 1)
    assert (result["structure_status"] == "FAIL") is bad


@pytest.mark.parametrize("payload", ["empty", "tab", "br", "cr", "math", "image", "space"])
def test_vertical_merge_continuation_payload_is_not_discarded(tmp_path: Path, payload: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    table = doc.tables[0]
    row = table.add_row()
    continuation = row.cells[0]
    row.cells[1].text = "56"
    for cell, value in [(table.cell(0, 0), "restart"), (continuation, "continue")]:
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), value)
        cell._tc.get_or_add_tcPr().append(merge)
    paragraph = continuation.paragraphs[0]
    if payload in {"tab", "br", "cr"}:
        paragraph.add_run()._r.append(OxmlElement("w:" + payload))
    elif payload == "space":
        paragraph.add_run(" ")
    elif payload == "math":
        paragraph._p.append(copy.deepcopy(doc.paragraphs[2]._p[-1]))
    elif payload == "image":
        paragraph.add_run().add_picture(str(tmp_path / "figure.png"))
    reference = next(u for u in truth["units"] if u["kind"] == "table")["reference"]
    reference["rows"] = 2
    reference["cells"][0]["rowspan"] = 2
    reference["cells"].append({"row": 1, "col": 1, "rowspan": 1, "colspan": 1, "text": "56"})
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("INVALID_TABLE_MERGE" in result["errors"]) is (payload != "empty")
    assert result["structure_status"] == ("PASS" if payload == "empty" else "FAIL")
