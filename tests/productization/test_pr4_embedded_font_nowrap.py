"""Unverified embedded fonts and nowrap growth cannot satisfy output claims."""

import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("tag", ["embedRegular", "embedBold", "embedItalic", "embedBoldItalic"])
@pytest.mark.parametrize("used", [True, False])
def test_used_embedded_font_is_rejected(tmp_path: Path, tag: str, used: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    doc.paragraphs[0].runs[0].font.name = "CustomEmbedded" if used else "Arial"
    doc.save(str(path))
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    root = etree.fromstring(files["word/fontTable.xml"])
    font, embed = OxmlElement("w:font"), OxmlElement("w:" + tag)
    font.set(qn("w:name"), "CustomEmbedded")
    embed.set(qn("r:id"), "embeddedFont")
    embed.set(qn("w:fontKey"), "{11111111-1111-1111-1111-111111111111}")
    font.append(embed)
    root.append(font)
    files["word/fontTable.xml"] = etree.tostring(root)
    rels = etree.Element(
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationships"
    )
    etree.SubElement(
        rels,
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
        Id="embeddedFont",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
        Target="fonts/custom.odttf",
    )
    files["word/_rels/fontTable.xml.rels"] = etree.tostring(rels)
    files["word/fonts/custom.odttf"] = b"unparsed synthetic font payload"
    types = etree.fromstring(files["[Content_Types].xml"])
    etree.SubElement(
        types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        PartName="/word/fonts/custom.odttf",
        ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont",
    )
    files["[Content_Types].xml"] = etree.tostring(types)
    with zipfile.ZipFile(path, "w") as archive:
        for n, data in files.items():
            archive.writestr(n, data)
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_FONT_MAPPING" in result["errors"]) is used
    assert result["structure_status"] == ("FAIL" if used else "PASS")


@pytest.mark.parametrize("where", ["cell", "style"])
@pytest.mark.parametrize("enabled", [True, False])
def test_no_wrap_cells(tmp_path: Path, where: str, enabled: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
    if where == "style":
        style = doc.styles.add_style("NoWrapCells", WD_STYLE_TYPE.TABLE)
        props = OxmlElement("w:tcPr")
        style._element.append(props)
        doc.tables[0].style = style
    node = OxmlElement("w:noWrap")
    node.set(qn("w:val"), "true" if enabled else "false")
    props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is enabled
    assert result["structure_status"] == ("FAIL" if enabled else "PASS")
