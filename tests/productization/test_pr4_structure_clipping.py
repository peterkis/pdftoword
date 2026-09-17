"""Reject clipped text/table formats and malformed WordML payload ancestry."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "tag,value,valid", [("w", "1", False), ("w", "100", True), ("fitText", "1", False)]
)
@pytest.mark.parametrize("where", ["direct", "math", "style", "default"])
def test_horizontal_text_scaling(
    tmp_path: Path, tag: str, value: str, valid: bool, where: str
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if where == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif where == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    node = OxmlElement("w:" + tag)
    node.set(qn("w:val"), value)
    props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is (not valid)


@pytest.mark.parametrize("rule", ["exact", "atLeast", "auto"])
def test_row_height_clipping(tmp_path: Path, rule: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    node = OxmlElement("w:trHeight")
    node.set(qn("w:val"), "1")
    node.set(qn("w:hRule"), rule)
    doc.tables[0].rows[0]._tr.get_or_add_trPr().append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_ROW_HEIGHT" in result["errors"]) is (rule == "exact")


@pytest.mark.parametrize("case", ["sectPr", "pPr", "run", "text", "math", "drawing"])
def test_payload_parent_structure(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0]
    if case == "sectPr":
        doc.sections[0]._sectPr.append(paragraph._p)
    elif case == "pPr":
        doc.paragraphs[1]._p.get_or_add_pPr().append(paragraph._p)
    elif case == "run":
        paragraph._p.get_or_add_pPr().append(paragraph.runs[0]._r)
    elif case == "text":
        run = paragraph.runs[0]._r
        run.get_or_add_rPr().append(run.find(qn("w:t")))
    elif case == "math":
        p = doc.paragraphs[2]._p
        p.get_or_add_pPr().append(p.xpath(".//m:oMath")[0])
    else:
        run = doc.paragraphs[-1].runs[0]._r
        run.get_or_add_rPr().append(run.find(qn("w:drawing")))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert any(
        c in result["errors"] for c in {"INVALID_WORD_STRUCTURE", "INVALID_DRAWING_CONTAINER"}
    )
    assert result["structure_status"] == "FAIL"
