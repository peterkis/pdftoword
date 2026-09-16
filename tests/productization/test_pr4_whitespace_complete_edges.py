"""Invisible whitespace and unclaimed complete-set edges cannot pass scoring."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("hidden", [True, False])
def test_hidden_space_is_content(tmp_path: Path, hidden: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0]
    paragraph.runs[0].text = "A"
    paragraph.add_run(" ").font.hidden = hidden
    paragraph.add_run("B")
    next(u for u in truth["units"] if u["unit_id"] == "a")["reference"]["text"] = "A B"
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("HIDDEN_CONTENT" in result["errors"]) is hidden
    assert result["structure_status"] == ("FAIL" if hidden else "PASS")


@pytest.mark.parametrize("case", ["matched", "duplicate", "unrelated", "no_prediction"])
def test_complete_uncertain_only_edges_claim_only_matching_predictions(
    tmp_path: Path, case: str
) -> None:
    path, truth, sources = specimen(tmp_path)
    truth["coverage_complete"] = True
    next(u for u in truth["units"] if u["kind"] == "figure_edge")["status"] = "uncertain"
    if case == "duplicate":
        sources["relations"].append(dict(sources["relations"][0]))
    elif case == "unrelated":
        sources["relations"].append(dict(sources["relations"][0], to="f"))
    elif case == "no_prediction":
        sources["relations"] = []
    result = evaluate(path, truth, sources)
    precision = result["metrics"]["relation_precision"]
    bad = case in {"duplicate", "unrelated"}
    assert precision["scored_count"] == int(bad)
    assert precision["correct_count"] == 0
    assert precision["not_scored_count"] == int(case != "no_prediction")
    assert precision["status"] == ("FAIL" if bad else "NOT_SCORED")
    assert (result["structure_status"] == "FAIL") is bad
