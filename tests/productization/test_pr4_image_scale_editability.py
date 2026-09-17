"""Extreme image shrinking and document locks cannot produce usable-output claims."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("case", ["one_emu", "tiny", "squashed", "normal"])
def test_actual_image_scale(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[-1]._p
    if case != "normal":
        width, height = {"one_emu": (1, 1), "tiny": (25400, 25400), "squashed": (914400, 25400)}[
            case
        ]
        for extent in paragraph.xpath(".//wp:extent | .//a:xfrm/a:ext"):
            extent.set("cx", str(width))
            extent.set("cy", str(height))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("IMAGE_DISPLAY_SCALE_INVALID" in result["errors"]) is (case != "normal")
    assert result["structure_status"] == ("PASS" if case == "normal" else "FAIL")


@pytest.mark.parametrize("kind", ["write", "document", "content"])
@pytest.mark.parametrize("has_units", [True, False])
def test_lock_errors_fail_editability(tmp_path: Path, kind: str, has_units: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    if kind == "content":
        paragraph = doc.paragraphs[0]._p
        parent, index = paragraph.getparent(), paragraph.getparent().index(paragraph)
        sdt, props, content = (
            OxmlElement("w:sdt"),
            OxmlElement("w:sdtPr"),
            OxmlElement("w:sdtContent"),
        )
        lock = OxmlElement("w:lock")
        lock.set(qn("w:val"), "contentLocked")
        props.append(lock)
        sdt.append(props)
        content.append(paragraph)
        sdt.append(content)
        parent.insert(index, sdt)
    else:
        node = OxmlElement("w:writeProtection" if kind == "write" else "w:documentProtection")
        node.set(qn("w:enforcement"), "true")
        doc.settings._element.append(node)
    doc.save(str(path))
    if not has_units:
        truth["units"] = []
    result = evaluate(path, truth, sources)
    assert result["editability_status"] == "FAIL"
