"""One-to-one formula allocation and unambiguous relationships are required."""

import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "case",
    [
        "correct",
        "wrong",
        "missing",
        "duplicate",
        "extra",
        "same",
        "same_missing",
        "unsupported",
        "unsupported_missing",
        "uncertain",
        "uncertain_extra",
    ],
)
def test_formula_nodes_are_matched_individually_without_reuse(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    next(a for a in truth["anchors"] if a["id"] == "f")["bbox"] = [0, 30, 50, 50]
    truth["anchors"].append({"id": "f2", "page": 1, "bbox": [50, 30, 100, 50]})
    ref = "x=1" if case in {"same", "same_missing"} else "y=2"
    if case.startswith("unsupported"):
        ref = r"\begin{matrix}y\end{matrix}"
    truth["units"].append(
        {
            "unit_id": "f2",
            "page": 1,
            "kind": "formula",
            "status": "uncertain" if case.startswith("uncertain") else "confirmed",
            "reference": {"source_anchor_id": "f2", "text": ref},
        }
    )
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[2]
    template = paragraph._p[-1]
    if case not in {"missing", "same_missing", "unsupported_missing"}:
        second = copy.deepcopy(template)
        second[0][-1].text = (
            "x=1" if case in {"same", "duplicate"} else "y=3" if case == "wrong" else "y=2"
        )
        paragraph._p.append(second)
    if case in {"extra", "uncertain_extra"}:
        extra = copy.deepcopy(template)
        extra[0][-1].text = "z=3"
        paragraph._p.append(extra)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    expected = 2 if case in {"correct", "same", "extra"} else 1
    assert result["metrics"]["formulas"]["correct_count"] == expected
    if case in {"correct", "same"}:
        assert result["content_status"] == result["structure_status"] == "PASS"
    elif case in {"unsupported", "uncertain"}:
        assert result["content_status"] == "REVIEW_REQUIRED"
        assert result["structure_status"] == "REVIEW_REQUIRED"
    else:
        assert "FAIL" in (result["content_status"], result["structure_status"])
    if case in {"extra", "uncertain_extra"}:
        assert "UNALIGNED_FORMULA" in result["errors"]
    if case == "unsupported_missing":
        assert "MISSING_FORMULA" in [f["code"] for f in result["failures"]]


@pytest.mark.parametrize("scope", ["body", "header", "footer"])
def test_position_tab_is_explicitly_unsupported(tmp_path: Path, scope: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = (
        doc.paragraphs[0] if scope == "body" else getattr(doc.sections[0], scope).paragraphs[0]
    )
    tab = OxmlElement("w:ptab")
    for key, value in [("alignment", "left"), ("relativeTo", "margin"), ("leader", "dot")]:
        tab.set(qn("w:" + key), value)
    paragraph.add_run()._r.append(tab)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert (
        "UNSUPPORTED_BODY_CONTENT" if scope == "body" else "UNSUPPORTED_VISIBLE_STORY"
    ) in result["errors"]
    assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize(
    "fault",
    ["document_duplicate", "root_duplicate", "invalid_id", "root_namespace", "child_namespace"],
)
def test_relationship_ids_must_be_unique_valid_names(tmp_path: Path, fault: str) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    name = "_rels/.rels" if fault == "root_duplicate" else "word/_rels/document.xml.rels"
    rels = etree.fromstring(files[name])
    relation = next(
        r
        for r in rels
        if r.get("Type", "").endswith(
            "/metadata/core-properties" if fault == "root_duplicate" else "/image"
        )
    )
    if fault == "root_namespace":
        rels.tag = "{urn:not-opc}Relationships"
    elif fault == "child_namespace":
        relation.tag = "{urn:not-opc}Relationship"
    elif fault == "invalid_id":
        relation.set("Id", "not an XML ID")
    else:
        rels.append(copy.deepcopy(relation))
    files[name] = etree.tostring(rels)
    with zipfile.ZipFile(path, "w") as archive:
        for part, data in files.items():
            archive.writestr(part, data)
    result = evaluate(path, truth, sources)
    assert (
        "OPC_INVALID_DECLARATION"
        if fault.endswith("namespace")
        else "OPC_RELATIONSHIP_ID_INVALID"
        if fault == "invalid_id"
        else "OPC_DUPLICATE_RELATIONSHIP_ID"
    ) in result["errors"]
    assert result["structure_status"] == "FAIL"


def test_relationship_xml_comments_are_not_relationship_entries(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for name in ["_rels/.rels", "word/_rels/document.xml.rels"]:
        root = etree.fromstring(files[name])
        root.insert(0, etree.Comment("ordinary XML comment"))
        files[name] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    result = evaluate(path, truth, sources)
    assert result["errors"] == []
    assert result["structure_status"] == "PASS"
