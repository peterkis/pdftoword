"""Unsupported backgrounds cannot make invisible text satisfy acceptance."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize(
    "case", ["direct", "math", "paragraph", "character", "inherited", "default", "table",
             "conditional", "highlight", "unused", "nil"]
)
def test_rendering_background_requires_explicit_support(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    target = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if case == "math":
        run = doc.paragraphs[2]._p.xpath(".//m:r")[0]
        target = OxmlElement("w:rPr")
        run.insert(0, target)
    elif case in {"paragraph", "character", "inherited", "unused", "table", "conditional"}:
        kind = WD_STYLE_TYPE.CHARACTER if case == "character" else (
            WD_STYLE_TYPE.TABLE if case in {"table", "conditional"} else WD_STYLE_TYPE.PARAGRAPH
        )
        style = doc.styles.add_style("BackgroundStyle", kind)
        target = style._element.get_or_add_rPr()
        if case == "inherited":
            derived = doc.styles.add_style("DerivedBackground", kind)
            derived.base_style = style
            style = derived
        if case == "character":
            doc.paragraphs[0].runs[0].style = style
        elif case in {"table", "conditional"}:
            doc.tables[0].style = style
            if case == "conditional":
                condition = OxmlElement("w:tblStylePr")
                condition.set(qn("w:type"), "firstRow")
                condition.append(target)
                style._element.append(condition)
        elif case != "unused":
            doc.paragraphs[0].style = style
    elif case == "default":
        target = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000" if case == "nil" else "FFFFFF")
    target.append(color)
    shading = OxmlElement("w:highlight" if case == "highlight" else "w:shd")
    shading.set(
        qn("w:val"), "white" if case == "highlight" else "nil" if case == "nil" else "clear"
    )
    if case != "highlight":
        shading.set(qn("w:fill"), "FFFFFF")
    target.append(shading)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    rejected = case not in {"unused", "nil"}
    assert ("UNSUPPORTED_VISIBILITY_STYLE" in result["errors"]) is rejected
    assert result["structure_status"] == ("FAIL" if rejected else "PASS")
