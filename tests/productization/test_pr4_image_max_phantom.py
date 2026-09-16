"""Oversized images and phantom formula payloads cannot count as visible evidence."""

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("fallback", [True, False])
@pytest.mark.parametrize("oversize", [True, False])
def test_image_page_upper_bound(tmp_path: Path, fallback: bool, oversize: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    sources["blocks"][-1]["images"][0]["fallback"] = fallback
    if oversize:
        doc: Any = Document(str(path))
        for extent in doc.paragraphs[-1]._p.xpath(".//wp:extent | .//a:xfrm/a:ext"):
            extent.set("cx", "91440000")
            extent.set("cy", "91440000")
        doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("IMAGE_DISPLAY_SCALE_INVALID" in result["errors"]) is oversize
    assert result["structure_status"] == ("FAIL" if oversize else "PASS")


@pytest.mark.parametrize("show", ["false", "true", "default", "plain"])
def test_phantom_controls_are_explicitly_unsupported(tmp_path: Path, show: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    math = doc.paragraphs[2]._p.xpath(".//m:oMath")[0]
    if show != "plain":
        phantom, props, content = (
            OxmlElement("m:phant"),
            OxmlElement("m:phantPr"),
            OxmlElement("m:e"),
        )
        if show != "default":
            node = OxmlElement("m:show")
            node.set(qn("m:val"), show)
            props.append(node)
        content.extend(list(math))
        phantom.append(props)
        phantom.append(content)
        math.append(phantom)
    next(u for u in truth["units"] if u["unit_id"] == "f")["reference"]["text"] = (
        r"\begin{matrix}x\end{matrix}"
    )
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_OMML_VISIBILITY" in result["errors"]) is (show != "plain")
    assert result["structure_status"] == ("REVIEW_REQUIRED" if show == "plain" else "FAIL")
