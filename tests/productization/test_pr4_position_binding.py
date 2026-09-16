"""Unverified positions and data-bound caches cannot pass acceptance."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("where", ["direct", "math", "style", "default"])
@pytest.mark.parametrize("value", ["31680", "0"])
def test_position_property(tmp_path: Path, where: str, value: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if where == "math":
        props = OxmlElement("w:rPr")
        doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
    elif where == "style":
        props = doc.styles["Normal"]._element.get_or_add_rPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
    position = OxmlElement("w:position")
    position.set(qn("w:val"), value)
    props.append(position)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is (value != "0")


@pytest.mark.parametrize("bound", [True, False])
def test_data_bound_cache_is_not_authoritative(tmp_path: Path, bound: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[0]._p
    parent, index = paragraph.getparent(), paragraph.getparent().index(paragraph)
    control, props = OxmlElement("w:sdt"), OxmlElement("w:sdtPr")
    if bound:
        binding = OxmlElement("w:dataBinding")
        binding.set(qn("w:storeItemID"), "{11111111-1111-1111-1111-111111111111}")
        binding.set(qn("w:xpath"), "/root/value")
        props.append(binding)
    content = OxmlElement("w:sdtContent")
    content.append(paragraph)
    control.append(props)
    control.append(content)
    parent.insert(index, control)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_DATA_BINDING" in result["errors"]) is bound
    assert result["structure_status"] == ("FAIL" if bound else "PASS")
