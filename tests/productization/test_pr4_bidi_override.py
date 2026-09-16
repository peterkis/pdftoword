"""Reject unsupported bdo containers and invalid bdo property placements."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("case", ["container", "direct", "math", "style", "default", "plain"])
def test_bidi_override_is_explicitly_rejected(tmp_path: Path, case: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    node = OxmlElement("w:bdo")
    node.set(qn("w:val"), "rtl")
    if case == "container":
        run = doc.paragraphs[0].runs[0]._r
        parent, index = run.getparent(), run.getparent().index(run)
        node.append(run)
        parent.insert(index, node)
    elif case != "plain":
        props = doc.paragraphs[0].runs[0]._r.get_or_add_rPr()
        if case == "math":
            props = OxmlElement("w:rPr")
            doc.paragraphs[2]._p.xpath(".//m:r")[0].insert(0, props)
        elif case == "style":
            props = doc.styles["Normal"]._element.get_or_add_rPr()
        elif case == "default":
            props = doc.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")[0]
        props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_BIDI_OVERRIDE" in result["errors"]) is (case != "plain")
    assert result["structure_status"] == ("PASS" if case == "plain" else "FAIL")
