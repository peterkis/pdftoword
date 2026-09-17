"""Bound paragraph indents and reject unverified floating-table placement."""
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("where", ["direct", "math", "style", "default"])
@pytest.mark.parametrize("indent", ["0", "720", "31680"])
def test_visible_indent_bound(tmp_path: Path, where: str, indent: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    props = doc.paragraphs[2 if where == "math" else 0]._p.get_or_add_pPr()
    if where == "style":
        props = doc.styles["Normal"]._element.get_or_add_pPr()
    elif where == "default":
        props = doc.styles._element.xpath("./w:docDefaults/w:pPrDefault/w:pPr")[0]
    node = OxmlElement("w:ind")
    node.set(qn("w:left"), indent)
    props.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is (indent == "31680")
    assert result["structure_status"] == ("FAIL" if indent == "31680" else "PASS")


@pytest.mark.parametrize("floating", [True, False])
def test_floating_table_is_explicitly_unsupported(tmp_path: Path, floating: bool) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    if floating:
        node = OxmlElement("w:tblpPr")
        node.set(qn("w:horzAnchor"), "page")
        node.set(qn("w:vertAnchor"), "page")
        node.set(qn("w:tblpX"), "31680")
        node.set(qn("w:tblpY"), "31680")
        doc.tables[0]._tbl.tblPr.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is floating
    assert result["structure_status"] == ("FAIL" if floating else "PASS")


@pytest.mark.parametrize("indent", ["0", "31680"])
def test_nonfloating_table_indent_bound(tmp_path: Path, indent: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    node = OxmlElement("w:tblInd")
    node.set(qn("w:type"), "dxa")
    node.set(qn("w:w"), indent)
    doc.tables[0]._tbl.tblPr.append(node)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNSUPPORTED_TEXT_POSITION" in result["errors"]) is (indent != "0")
