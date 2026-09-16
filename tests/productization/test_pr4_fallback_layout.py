"""Visibility, verified structural fallback and layout grids are separate checks."""

import copy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("mode", ["direct", "character", "paragraph", "override"])
def test_math_run_visibility_is_checked_before_normalization(tmp_path: Path, mode: str) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    paragraph = doc.paragraphs[2]
    run = paragraph._p[-1][0]
    props = OxmlElement("w:rPr")
    if mode == "direct":
        props.append(OxmlElement("w:vanish"))
    else:
        style = doc.styles.add_style(
            "MathHidden",
            WD_STYLE_TYPE.PARAGRAPH
            if mode in {"paragraph", "override"}
            else WD_STYLE_TYPE.CHARACTER,
        )
        style.font.hidden = True
        if mode == "character":
            ref = OxmlElement("w:rStyle")
            ref.set(qn("w:val"), style.style_id)
            props.append(ref)
        else:
            paragraph.style = style
        if mode == "override":
            vanish = OxmlElement("w:vanish")
            vanish.set(qn("w:val"), "0")
            props.append(vanish)
    run.insert(0, props)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("HIDDEN_CONTENT" in result["errors"]) is (mode != "override")
    assert result["structure_status"] == ("PASS" if mode == "override" else "FAIL")


@pytest.mark.parametrize("kind", ["formula", "table"])
@pytest.mark.parametrize("retain_wrong_structure", [False, True])
def test_verified_image_replacement_is_reviewable_but_wrong_structure_still_fails(
    tmp_path: Path, kind: str, retain_wrong_structure: bool
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    bid = "f" if kind == "formula" else "t"
    block = next(b for b in sources["blocks"] if b["block_id"] == bid)
    if kind == "formula":
        paragraph = doc.paragraphs[2]
        if retain_wrong_structure:
            paragraph._p[-1][0][-1].text = "x=2"
        else:
            paragraph._p.remove(paragraph._p[-1])
    else:
        paragraph = doc.tables[0].cell(0, 0).paragraphs[0]
        if retain_wrong_structure:
            paragraph.runs[0].text = "99"
        else:
            saved = copy.deepcopy(paragraph._p)
            for run in list(saved.findall(qn("w:r"))):
                saved.remove(run)
            table = doc.tables[0]._tbl
            table.getparent().replace(table, saved)
            paragraph = doc.paragraphs[3]
    paragraph.add_run().add_picture(str(tmp_path / "figure.png"))
    image = copy.deepcopy(sources["blocks"][-1]["images"][0])
    image.update(bbox=block["bbox"], fallback=True)
    block["images"] = [image]
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert result["structure_status"] == ("FAIL" if retain_wrong_structure else "REVIEW_REQUIRED")
    if not retain_wrong_structure:
        assert result["content_status"] == "REVIEW_REQUIRED"
        assert result["editability_status"] == "FAIL"


@pytest.mark.parametrize("shape", ["bordered", "empty_row", "empty_cell", "valid"])
def test_implicit_layout_wrapper_cannot_add_visible_or_empty_grid(
    tmp_path: Path, shape: str
) -> None:
    path, truth, sources = specimen(tmp_path)
    doc: Any = Document(str(path))
    original = doc.paragraphs[0]._p
    table = doc.add_table(
        rows=2 if shape == "empty_row" else 1, cols=2 if shape == "empty_cell" else 1
    )
    if shape == "bordered":
        table.style = "Table Grid"
    else:
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            child = OxmlElement("w:" + edge)
            child.set(qn("w:val"), "nil")
            borders.append(child)
        table._tbl.tblPr.append(borders)
    doc._element.body.insert(0, table._tbl)
    empty = table.cell(0, 0).paragraphs[0]._p
    empty.getparent().replace(empty, original)
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("UNEXPECTED_LAYOUT_GRID" in result["errors"]) is (shape != "valid")
    assert result["structure_status"] == ("PASS" if shape == "valid" else "FAIL")


def test_explicit_layout_reference_does_not_allow_an_extra_empty_row(tmp_path: Path) -> None:
    path, truth, sources = specimen(tmp_path)
    next(u for u in truth["units"] if u["kind"] == "table")["reference"]["is_data_table"] = False
    doc = Document(str(path))
    doc.tables[0].add_row()
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNEXPECTED_LAYOUT_GRID" in result["errors"]
    assert result["structure_status"] == "FAIL"
