"""Every actual drawing occurrence needs its own source image binding."""

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from tests.productization.test_acceptance_harness import specimen

from acceptance.metrics import evaluate


@pytest.mark.parametrize("placement", ["new_paragraph", "same_paragraph", "different_image"])
def test_extra_drawing_cannot_hide_behind_existing_image_hash(
    tmp_path: Path, placement: str
) -> None:
    from PIL import Image

    path, truth, sources = specimen(tmp_path)
    doc = Document(str(path))
    paragraph = doc.paragraphs[-1] if placement == "same_paragraph" else doc.add_paragraph()
    image = tmp_path / "figure.png"
    if placement == "different_image":
        image = tmp_path / "extra.png"
        Image.new("RGB", (20, 20), "red").save(image)
    paragraph.add_run().add_picture(str(image), width=Inches(1))
    doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert "UNALIGNED_IMAGE" in result["errors"]
    assert result["structure_status"] == "FAIL"


@pytest.mark.parametrize("second_occurrence_present", [False, True])
def test_each_declared_repeated_image_requires_a_distinct_drawing(
    tmp_path: Path, second_occurrence_present: bool
) -> None:
    import copy

    path, truth, sources = specimen(tmp_path)
    sources["blocks"][-1]["images"].append(copy.deepcopy(sources["blocks"][-1]["images"][0]))
    if second_occurrence_present:
        doc = Document(str(path))
        doc.paragraphs[-1].add_run().add_picture(str(tmp_path / "figure.png"), width=Inches(1))
        doc.save(str(path))
    result = evaluate(path, truth, sources)
    assert ("MISSING_OR_REPLACED_IMAGE" in result["errors"]) is not second_occurrence_present
    assert result["structure_status"] == ("PASS" if second_occurrence_present else "FAIL")
